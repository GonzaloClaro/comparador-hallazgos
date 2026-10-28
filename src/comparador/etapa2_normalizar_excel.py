"""Etapa 2: Normalización del Excel regulatorio a hallazgos canónicos.

Sin LLM. El schema del Excel es fijo (definido por la SEC).
"""

import re
import unicodedata
from collections.abc import Sequence
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from comparador.schemas import (
    Calificacion,
    Evidencia,
    FuenteEvidencia,
    Hallazgo,
    MetadatosInforme,
)

# Mínimo estructural: sin esto no hay hallazgo comparable.
COLUMNAS_REQUERIDAS = {
    "Paño(s)",
    "Relé",
    "Clasificación del Hallazgo",
    "Descripción del Hallazgo",
}

# Presentes en el formato SEC completo; algunas auditoras no las envían
# (Ensaut/Epsilon 2026-07-07). Cuando faltan, el campo queda sin dato y no
# se compara — nunca se inventa.
COLUMNAS_OPCIONALES = {
    "Acción Correctiva del Hallazgo",
    "Calificación",
}

# Variantes de header observadas en archivos reales.
ALIAS_COLUMNAS = {
    "Paño": "Paño(s)",                                # COORDINADO B/Theta
    "Equipo": "Relé",                                 # Ensaut/Epsilon
    "Clasificación de Hallazgo": "Clasificación del Hallazgo",  # Ensaut/Epsilon
}

MAPA_CALIFICACION: dict[str, Calificacion] = {
    "menor": Calificacion.MENOR,
    "relevante": Calificacion.RELEVANTE,
    "crítico": Calificacion.CRITICO,
    "critico": Calificacion.CRITICO,
    "criticidad menor": Calificacion.MENOR,
    "criticidad relevante": Calificacion.RELEVANTE,
    "crit. relevante": Calificacion.RELEVANTE,
    "crit. menor": Calificacion.MENOR,
    "hallazgo menor": Calificacion.MENOR,
    "hallazgo relevante": Calificacion.RELEVANTE,
    "hallazgo crítico": Calificacion.CRITICO,
}

# Código de paño ("BT1", "CT1A", "E10"). La Etapa 3 usa la misma función de
# separación (split_guion_entre_codigos) para que las tuplas de ambos lados calcen.
_CODIGO_PANO = r"[A-Za-z]+\d+[A-Za-z]?"

CAMPOS_METADATOS = {
    "Coordinado": "coordinado",
    "Subestación": "subestacion",
    "Informe de Auditoría": "informe_id",
    "Informe de Auditoria": "informe_id",
    "Fecha de Auditoría": "fecha_auditoria",
    "Fecha de Auditoria": "fecha_auditoria",
    # El Excel de Ensaut usa la etiqueta genérica "Fecha" (foto COORDINADO B 2026-07-07).
    "Fecha": "fecha_auditoria",
    "Versión de Informe de Auditoría": "version",
    "Version de Informe de Auditoria": "version",
}


class ExcelNormalizationError(Exception):
    pass


def _celda(row: "pd.Series[str]", col: str) -> str:
    """Valor de celda como string limpio; '' si está vacía o es NaN."""
    valor = row.get(col, "")
    if pd.isna(valor):
        return ""
    texto = str(valor).strip()
    return "" if texto == "nan" else texto


def _encontrar_fila_header(df_raw: pd.DataFrame) -> int:
    """Busca la fila que contiene los headers de la tabla de hallazgos.

    Aplica ALIAS_COLUMNAS a las celdas candidatas: un Excel con "Equipo" en
    vez de "Relé" también debe detectar su fila de headers.
    """
    for idx in range(min(30, len(df_raw))):
        row_values = {
            ALIAS_COLUMNAS.get(texto, texto)
            for texto in (str(v).strip() for v in df_raw.iloc[idx] if pd.notna(v))
        }
        if "Relé" in row_values and "Calificación" in row_values:
            return idx
        if "Descripción del Hallazgo" in row_values:
            return idx
    raise ExcelNormalizationError(
        "No se encontró la fila de headers de la tabla de hallazgos en las primeras 30 filas."
    )


def _extraer_metadatos(df_raw: pd.DataFrame, header_row: int) -> MetadatosInforme:
    """Extrae metadatos del informe de las filas superiores al header."""
    valores: dict[str, str] = {}

    for idx in range(header_row):
        for col_idx in range(min(5, df_raw.shape[1])):
            celda = df_raw.iloc[idx, col_idx]
            if pd.isna(celda):
                continue
            # Match exacto (con ":" final tolerado): startswith colisionaba —
            # "Fecha de Informe de Auditoría y Tabla de Hallazgos" empieza
            # con "Fecha" pero es otro campo.
            texto = str(celda).strip().rstrip(":").strip()
            campo_schema = CAMPOS_METADATOS.get(texto)
            if campo_schema and campo_schema not in valores:
                # El valor es la primera celda no vacía a la derecha del label
                # (las celdas intermedias pueden venir merged/vacías).
                for col_valor in range(col_idx + 1, df_raw.shape[1]):
                    valor = df_raw.iloc[idx, col_valor]
                    if pd.notna(valor) and str(valor).strip():
                        valores[campo_schema] = str(valor).strip()
                        break

    return MetadatosInforme(
        coordinado=valores.get("coordinado", ""),
        subestacion=valores.get("subestacion", ""),
        informe_id=valores.get("informe_id", ""),
        fecha_auditoria=valores.get("fecha_auditoria"),
        version=valores.get("version"),
    )


def _normalizar_calificacion(valor_raw: str) -> Calificacion:
    """Mapea variantes textuales al enum canónico. Falla si no se reconoce."""
    normalizado = valor_raw.strip().lower()
    if normalizado in MAPA_CALIFICACION:
        return MAPA_CALIFICACION[normalizado]
    raise ExcelNormalizationError(
        f"Calificación no reconocida: '{valor_raw}'. "
        f"Valores válidos: {list(MAPA_CALIFICACION.keys())}"
    )


def _normaliza(texto: str) -> str:
    """casefold + sin acentos + espacios colapsados, para comparar patrones."""
    sin_acentos = "".join(
        c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c)
    )
    return " ".join(sin_acentos.casefold().split())


def _es_transversal(rele: str, patrones_normalizados: Sequence[str]) -> bool:
    """La fila no pertenece a un relé específico (ej. "Análisis de coordinación").

    Match por contains sobre el valor normalizado: las variantes reales
    ("Análisis de coordinación de protecciones") deben seguir calzando.
    Los patrones vienen del catálogo por empresa (filas_transversales).

    Un relé sin ningún carácter alfanumérico ("-", "—", "/") no identifica
    equipo alguno: es transversal para toda empresa, sin depender del
    catálogo (Beta filas 221-222, Eta filas 51-52).
    """
    if not any(c.isalnum() for c in rele):
        return True
    rele_norm = _normaliza(rele)
    return any(patron in rele_norm for patron in patrones_normalizados)


def _clave_rele(texto: str) -> str:
    """Misma normalización que la clave de tupla de la Etapa 4."""
    return texto.replace("-", "").replace(" ", "").replace("_", "").casefold()


def _quitar_rele_del_pano(pano: str, rele: str) -> str:
    """'ET2 (SEL 387E)' con relé 'SEL 387E' -> 'ET2' (Excel Alfa).

    Solo si el paréntesis repite el relé de la fila: otro contenido entre
    paréntesis puede ser parte del nombre del paño y se conserva.
    """
    m = re.fullmatch(r"(.+?)\s*\(([^()]*)\)", pano)
    if m and _clave_rele(m.group(2)) == _clave_rele(rele):
        return m.group(1).strip()
    return pano


def split_guion_entre_codigos(pano: str) -> list[str]:
    """'BT1-CT1' -> ['BT1', 'CT1'] (Zeta, Gamma). La Etapa 3 la reusa.

    Solo cuando todos los lados del guion son códigos de paño: "E-1",
    "DT1 – P.NEUTRO" o un nombre de relé quedan intactos.
    """
    partes = re.split(r"\s*[-–—]\s*", pano.strip())
    if len(partes) > 1 and all(re.fullmatch(_CODIGO_PANO, p) for p in partes):
        return partes
    return [pano]


def _split_panos(panos_raw: str, rele: str) -> list[str]:
    """Separa paños múltiples. Ej: 'BT1, ET1, E1' -> ['BT1', 'ET1', 'E1'].

    La barra y el guion entre códigos también separan: "BT1/ET1" y
    "BT1-CT1" se desdoblan igual que hace la Etapa 3 con las secciones
    compuestas del PDF. Las tuplas (paño, relé) de ambos lados deben
    calzar en el matching.
    """
    partes = [p for p in re.split(r"\s*[,;/]\s*", panos_raw.strip()) if p]
    return [
        codigo
        for parte in partes
        for codigo in split_guion_entre_codigos(_quitar_rele_del_pano(parte, rele))
    ]


def _hoja_de_hallazgos(hojas: Sequence[str]) -> str:
    """La hoja con "hallazgo" en el nombre; si no hay ninguna, la primera.

    Con dos o más hojas de hallazgos no hay cómo saber cuál es la vigente:
    elegir una en silencio dejaría la otra fuera de la comparación.
    """
    de_hallazgos = [h for h in hojas if "hallazgo" in h.lower()]
    if len(de_hallazgos) > 1:
        raise ExcelNormalizationError(
            f"El Excel tiene {len(de_hallazgos)} hojas de hallazgos ({', '.join(de_hallazgos)}): "
            "deja una sola (la vigente) y vuelve a subirlo."
        )
    return de_hallazgos[0] if de_hallazgos else hojas[0]


def _columnas_en_hoja(ws: Worksheet, fila_header: int, nombres: Sequence[str]) -> dict[str, int]:
    """Columna (1-based) de cada encabezado en la fila del header según openpyxl.

    Es además la guarda de `fila_excel`: si pandas y openpyxl no ven el header
    en la misma fila, el número de fila que se muestra estaría corrido.
    """
    encontrados: dict[str, int] = {}
    for celda in ws[fila_header]:
        if celda.value is None:
            continue
        texto = str(celda.value).strip()
        nombre = ALIAS_COLUMNAS.get(texto, texto)
        if nombre in nombres and nombre not in encontrados:
            encontrados[nombre] = celda.column
    faltan = [n for n in nombres if n not in encontrados]
    if faltan:
        raise ExcelNormalizationError(
            f"Los encabezados {faltan} no están en la fila {fila_header} de la hoja "
            f"«{ws.title}» al leerla celda por celda: los números de fila no serían confiables."
        )
    return encontrados


def _valores_combinados(
    ws: Worksheet, columnas: Sequence[int], fila_header: int
) -> dict[tuple[int, int], str]:
    """(fila, columna) → valor visible de las celdas de una combinada vertical.

    En el Excel una celda combinada muestra su valor en todas sus filas, pero
    al leerla solo la primera lo trae. Paño y relé se heredan SOLO así: una
    celda vacía que no está combinada es un dato faltante (fila parcialmente
    llena → fallar), no una repetición de la fila anterior. Solo cuentan los
    rangos de una sola columna (la de Paño o la de Relé) bajo el encabezado:
    una combinada horizontal (Paño+Relé, Ítem+Paño) copiaría el valor de una
    columna en otra.
    """
    valores: dict[tuple[int, int], str] = {}
    for rango in ws.merged_cells.ranges:
        if rango.min_col != rango.max_col or rango.min_col not in columnas:
            continue
        if rango.min_row <= fila_header:
            continue
        valor = ws.cell(rango.min_row, rango.min_col).value
        texto = "" if valor is None else str(valor).strip()
        for fila in range(rango.min_row, rango.max_row + 1):
            valores[(fila, rango.min_col)] = texto
    return valores


def _visible(
    leido: str, fila: int, columna: int, combinadas: dict[tuple[int, int], str], nombre: str
) -> str:
    """El valor que ve la revisora en la celda: el de su combinada, si es parte de una.

    pandas lee el contenido guardado de cada celda, también el que queda oculto
    bajo una combinada (p. ej. un archivo guardado conservando el contenido de
    las celdas combinadas): si difiere de lo visible, no se adivina cuál vale.
    """
    if (fila, columna) not in combinadas:
        return leido
    visible = combinadas[(fila, columna)]
    if leido and leido != visible:
        raise ExcelNormalizationError(
            f"La celda de {nombre} de la fila {fila} es parte de una celda combinada que "
            f"muestra «{visible}», pero guarda «{leido}» oculto: descombina y repite el valor."
        )
    return visible


def _texto_plano(valor: object) -> str:
    return " ".join(str(valor).split()) if valor is not None else ""


def normalizar_excel(
    ruta: Path,
    patrones_transversales: Sequence[str] = (),
) -> tuple[list[Hallazgo], MetadatosInforme]:
    """Normaliza un Excel regulatorio a una lista de Hallazgos canónicos.

    Args:
        ruta: archivo Excel regulatorio.
        patrones_transversales: patrones del catálogo de la empresa que
            identifican filas sin relé específico (van a estado transversal).

    Returns:
        Tupla de (hallazgos, metadatos).

    Raises:
        ExcelNormalizationError: si el Excel no tiene la estructura esperada.
    """
    # Los Excel reales traen varias hojas (ej: "Definiciones" + "Tabla de
    # Hallazgos"). Se usa la hoja de hallazgos; si no existe, la primera.
    hojas = [str(h) for h in pd.ExcelFile(ruta).sheet_names]
    hoja = _hoja_de_hallazgos(hojas)

    df_raw = pd.read_excel(ruta, sheet_name=hoja, header=None, dtype=str)

    header_row = _encontrar_fila_header(df_raw)
    metadatos = _extraer_metadatos(df_raw, header_row)

    df = pd.read_excel(ruta, sheet_name=hoja, header=header_row, dtype=str)
    df.columns = df.columns.str.strip()
    df = df.rename(columns=ALIAS_COLUMNAS)

    faltantes = COLUMNAS_REQUERIDAS - set(df.columns)
    if faltantes:
        raise ExcelNormalizationError(
            f"Columnas faltantes en el Excel: {faltantes}. "
            f"Columnas encontradas: {list(df.columns)}"
        )

    tiene_calificacion = "Calificación" in df.columns
    tiene_accion = "Acción Correctiva del Hallazgo" in df.columns

    # El formato regulatorio completo DEBE traer las opcionales (el equipo revisor
    # 2026-07-16). Si faltan, se procesa igual pero queda registrado para
    # que la UI lo advierta — pedir a la auditora el Excel completo.
    metadatos.columnas_ausentes = sorted(COLUMNAS_OPCIONALES - set(df.columns))

    patrones_norm = [_normaliza(p) for p in patrones_transversales if p.strip()]

    # La misma hoja celda por celda: combinadas, filas ocultas y la guarda de
    # que fila_excel es la fila que ve la revisora al abrir el archivo.
    ws = load_workbook(ruta, data_only=True)[hoja]
    columnas = _columnas_en_hoja(
        ws, header_row + 1, ["Paño(s)", "Relé", "Descripción del Hallazgo"]
    )
    # Celdas combinadas (variante Ensaut): paño y relé se muestran en todas las
    # filas del grupo, pero solo la primera los trae.
    combinadas = _valores_combinados(ws, [columnas["Paño(s)"], columnas["Relé"]], header_row + 1)

    hallazgos: list[Hallazgo] = []

    for idx, row in df.iterrows():
        fila_excel = int(str(idx)) + header_row + 2
        panos_raw = _visible(_celda(row, "Paño(s)"), fila_excel, columnas["Paño(s)"],
                             combinadas, "Paño(s)")
        rele = _visible(_celda(row, "Relé"), fila_excel, columnas["Relé"], combinadas, "Relé")
        clasificacion = _celda(row, "Clasificación del Hallazgo")
        descripcion = _celda(row, "Descripción del Hallazgo")
        accion = _celda(row, "Acción Correctiva del Hallazgo") if tiene_accion else None
        calificacion_raw = _celda(row, "Calificación") if tiene_calificacion else ""

        # Fila sin contenido propio (paño/relé pueden venir de una celda
        # combinada): no es un hallazgo, se ignora.
        if not any([clasificacion, descripcion, accion or "", calificacion_raw]):
            continue

        # La descripción leída celda por celda debe ser la de pandas: si no, la
        # fila que se muestra no es la del hallazgo.
        # (Solo si pandas la leyó: "N/A" o "None" las lee vacías y entonces
        # corresponde el error de campo obligatorio vacío, no este.)
        celda_desc = _texto_plano(ws.cell(fila_excel, columnas["Descripción del Hallazgo"]).value)
        if descripcion and celda_desc != _texto_plano(descripcion):
            raise ExcelNormalizationError(
                f"La fila {fila_excel} de la hoja «{hoja}» no tiene la descripción que se leyó "
                "para ella: los números de fila no serían confiables."
            )

        # Fila parcialmente llena: fallar ruidosamente. Descartarla en silencio
        # haría desaparecer un hallazgo de la comparación sin aviso al revisor.
        obligatorios = {
            "Paño(s)": panos_raw,
            "Relé": rele,
            "Descripción del Hallazgo": descripcion,
        }
        if tiene_calificacion:
            obligatorios["Calificación"] = calificacion_raw
        vacios = [col for col, valor in obligatorios.items() if not valor]
        if vacios:
            herencia = (
                " (Paño y Relé solo heredan el valor de arriba si la celda está combinada "
                "con ella en vertical; las demás columnas nunca lo heredan)"
                if {"Paño(s)", "Relé"} & set(vacios) else ""
            )
            raise ExcelNormalizationError(
                f"Fila {fila_excel} del Excel tiene campos obligatorios vacíos: {vacios}{herencia}."
            )
        dimension = ws.row_dimensions.get(fila_excel)
        fila_oculta = bool(dimension is not None and dimension.hidden)

        calificacion = _normalizar_calificacion(calificacion_raw) if tiene_calificacion else None
        transversal = _es_transversal(rele, patrones_norm)

        # Una fila transversal es UNA entrada de revisión manual: desdoblarla
        # por paño generaría N copias idénticas de la misma fila_excel.
        panos = [panos_raw.strip()] if transversal else _split_panos(panos_raw, rele)
        es_multi = len(panos) > 1

        texto_crudo = f"{clasificacion}: {descripcion}"
        if accion:
            texto_crudo += f" | Acción: {accion}"

        for pano in panos:
            hallazgos.append(
                Hallazgo(
                    pano=pano,
                    rele=rele,
                    clasificacion=clasificacion,
                    descripcion=descripcion,
                    accion_correctiva=accion,
                    calificacion=calificacion,
                    es_transversal=transversal,
                    evidencia=Evidencia(
                        fuente=FuenteEvidencia.EXCEL,
                        fila_excel=fila_excel,
                        es_multi_pano_origen=es_multi,
                        texto_crudo=texto_crudo,
                        hoja_excel=hoja,
                        fila_oculta=fila_oculta,
                    ),
                )
            )

    return hallazgos, metadatos
