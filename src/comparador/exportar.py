"""Exportación de ReporteInforme a Excel (.xlsx) y CSV."""

from __future__ import annotations

import csv
import io
from collections import Counter
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from comparador import feedback as feedback_mod
from comparador.citas import etiqueta_verificacion
from comparador.evaluar_checklist import clasificar
from comparador.feedback import ETIQUETA_MOTIVO
from comparador.schemas import (
    ChecklistInforme,
    EstadoHallazgo,
    FeedbackRun,
    MotivoFeedback,
    ReporteInforme,
    ResultadoChecklist,
)

COLUMNAS = [
    "Paño", "Relé", "Clasificación", "Estado", "Tipo Match", "Confianza",
    "Descripción PDF", "Descripción Excel",
    "Acción Correctiva PDF", "Acción Correctiva Excel",
    "Calificación PDF", "Calificación Excel",
    "Justificación", "Página PDF", "Fila Excel", "Campos Diferentes",
    "Equivalencias de redacción",
    "Cita descripción PDF", "Cita acción PDF", "Hoja Excel", "Fila Excel oculta",
    # 0.9.0 (PD71): el motivo de `emparejamiento_dudoso` (asignación global). Desde
    # PD72 (decisión del usuario 2026-10-26) la columna está siempre, vacía si ningún
    # par lo trae; el semantico_bajo ya lo dice «Tipo Match».
    "Emparejamiento dudoso",
    # 0.10.0 (PD60): otros relés de la sección que nombra el hallazgo del PDF. Siempre
    # presente, como la anterior; vacía con la lista None (reporte < 0.10.0) o [].
    "Otros relés nombrados (PDF)",
]

FILLS_ESTADO = {
    EstadoHallazgo.COINCIDE: PatternFill(start_color="C6EFCE", fill_type="solid"),
    EstadoHallazgo.DIFIERE: PatternFill(start_color="FFEB9C", fill_type="solid"),
    EstadoHallazgo.FALTA_EN_EXCEL: PatternFill(start_color="FFC7CE", fill_type="solid"),
    EstadoHallazgo.EXTRA_EN_EXCEL: PatternFill(start_color="FFD8B1", fill_type="solid"),
    # Gris, como "no verificable" del checklist: revisión manual.
    EstadoHallazgo.TRANSVERSAL: PatternFill(start_color="D9D9D9", fill_type="solid"),
}


def _fila_match(r: Any) -> list[Any]:
    h_pdf = r.hallazgo_pdf
    h_excel = r.hallazgo_excel

    # El detalle (similitud/umbral, lado vacío) va entre paréntesis; separador
    # ";" porque el detalle contiene comas.
    campos_diff = "; ".join(
        d.campo.value + (f" ({d.detalle})" if d.detalle else "")
        for d in r.diferencias
    )
    equivalencias = _texto_equivalencias(r)

    return [
        r.pano,
        r.rele,
        h_pdf.clasificacion if h_pdf else (h_excel.clasificacion if h_excel else ""),
        r.estado.value,
        r.tipo_match.value,
        round(r.confianza, 2),
        h_pdf.descripcion if h_pdf else "",
        h_excel.descripcion if h_excel else "",
        (h_pdf.accion_correctiva or "") if h_pdf else "",
        (h_excel.accion_correctiva or "") if h_excel else "",
        h_pdf.calificacion.value if h_pdf and h_pdf.calificacion else "",
        h_excel.calificacion.value if h_excel and h_excel.calificacion else "",
        r.justificacion,
        h_pdf.evidencia.pagina_pdf if h_pdf and h_pdf.evidencia.pagina_pdf else "",
        h_excel.evidencia.fila_excel if h_excel and h_excel.evidencia.fila_excel else "",
        campos_diff,
        equivalencias,
        # Estado de la verificación de cada cita contra el PDF (schema 0.7.0;
        # "sin verificar" en reportes anteriores).
        etiqueta_verificacion(h_pdf.evidencia.verificacion_descripcion) if h_pdf else "",
        _cita_accion(h_pdf),
        (h_excel.evidencia.hoja_excel or "") if h_excel else "",
        # None = no se revisó (reporte < 0.7.0): vacío, nunca "no".
        _si_no(h_excel.evidencia.fila_oculta) if h_excel else "",
        r.emparejamiento_dudoso or "",
        "; ".join(h_pdf.evidencia.otros_reles_nombrados or []) if h_pdf else "",
    ]


def _si_no(valor: bool | None) -> str:
    return "" if valor is None else ("sí" if valor else "no")


def _cita_accion(h_pdf: Any) -> str:
    if not h_pdf or not h_pdf.accion_correctiva:
        return ""
    return etiqueta_verificacion(h_pdf.evidencia.verificacion_accion)


def _texto_equivalencias(r: Any) -> str:
    """Campos que el juez de redacción declaró mismo contenido pese a la baja
    similitud: viajan en la descarga para que la decisión automática no se
    pierda fuera del dashboard. Separador " | ": la justificación puede
    traer comas y punto y coma."""
    return " | ".join(
        f"{e.campo.value} (similitud {e.similitud:.2f}; juez de redacción: {e.justificacion})"
        for e in r.equivalencias_redaccion
    )


def exportar_xlsx(reporte: ReporteInforme) -> bytes:
    """Genera un Excel con formato condicional desde un ReporteInforme."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Comparación Hallazgos"

    header_font = Font(bold=True)
    for col_idx, col_name in enumerate(COLUMNAS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font

    for row_idx, resultado in enumerate(reporte.resultados, 2):
        fila = _fila_match(resultado)
        fill = FILLS_ESTADO.get(resultado.estado)

        for col_idx, valor in enumerate(fila, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            if fill:
                cell.fill = fill

    for col_idx in range(1, len(COLUMNAS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 20

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    _agregar_hoja_resumen(wb, reporte)
    if reporte.checklist:
        _agregar_hoja_checklist(wb, reporte)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _agregar_hoja_resumen(wb: Workbook, reporte: ReporteInforme) -> None:
    """Hoja que identifica el archivo: estos xlsx se acumulan y reenvían en el
    flujo regulatorio — sin esto, el export queda huérfano de contexto."""
    ws = wb.create_sheet("Resumen", 0)
    r = reporte.resumen
    filas = [
        ("Informe", reporte.informe_id),
        ("Coordinado", reporte.coordinado),
        ("Subestación", reporte.subestacion),
        ("Empresa auditora", reporte.empresa_auditora.value),
        ("Fecha de auditoría", reporte.fecha_auditoria or ""),
        ("Procesado", reporte.fecha_procesamiento.isoformat(sep=" ", timespec="minutes")),
        ("Modelo de extracción", reporte.modelo_extraccion),
        ("", ""),
        ("Total de hallazgos", r.total),
        ("Coinciden", r.coinciden),
        ("Coinciden con redacción distinta (de las anteriores)",
         r.coinciden_con_redaccion_distinta),
        ("Difieren", r.difieren),
        ("Faltan en Excel", r.faltan_en_excel),
        ("Extra en Excel (faltan en PDF)", r.extra_en_excel),
        ("Transversales (revisión manual)", r.transversales),
        ("Requieren revisión", r.requieren_revision),
    ]
    if r.filas_excel_ocultas is not None:   # None = reporte < 0.7.0, no se revisó
        filas.append(("Filas del Excel ocultas (filtro activo)", r.filas_excel_ocultas))
    if reporte.checklist:
        filas.append(("Checklist", f"{reporte.checklist.cumplen}/{reporte.checklist.total}"))
        if reporte.checklist.bajados_sin_cita_verificada is not None:
            filas.append(("Checklist: «cumple» bajados a no verificable sin cita verificada",
                          reporte.checklist.bajados_sin_cita_verificada))
        filas += _filas_lectura_por_partes(reporte.checklist)

    for row_idx, (etiqueta, valor) in enumerate(filas, 1):
        celda = ws.cell(row=row_idx, column=1, value=etiqueta)
        celda.font = Font(bold=True)
        ws.cell(row=row_idx, column=2, value=valor)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40


def _filas_lectura_por_partes(cl: ChecklistInforme) -> list[tuple[str, Any]]:
    """Filas del Resumen de un informe leído por partes (M08, schema 0.8.0);
    ninguna si se leyó entero o el reporte es < 0.8.0."""
    if cl.tramos_leidos is None:
        return []
    filas: list[tuple[str, Any]] = [
        ("Checklist: informe leído por partes (págs de cada parte)",
         ", ".join(f"{a}-{b}" for a, b in cl.tramos_leidos)),
        ("Checklist: «cumple» vistos solo en parte del informe (con aviso)",
         cl.marcados_un_tramo),
    ]
    if cl.bajados_un_tramo is not None:   # None = CHECKLIST_UN_TRAMO=marcar
        filas.append(("Checklist: «cumple» bajados a no verificable por lectura en partes",
                      cl.bajados_un_tramo))
    return filas


FILLS_CHECKLIST = {
    ResultadoChecklist.CUMPLE: PatternFill(start_color="C6EFCE", fill_type="solid"),
    ResultadoChecklist.NO_CUMPLE: PatternFill(start_color="FFC7CE", fill_type="solid"),
    ResultadoChecklist.NO_VERIFICABLE: PatternFill(start_color="D9D9D9", fill_type="solid"),
}

TEXTO_RESULTADO = {
    ResultadoChecklist.CUMPLE: "Cumple",
    ResultadoChecklist.NO_CUMPLE: "No cumple",
    ResultadoChecklist.NO_VERIFICABLE: "No es posible verificar",
}


def _agregar_hoja_checklist(wb: Workbook, reporte: ReporteInforme) -> None:
    assert reporte.checklist is not None
    ws = wb.create_sheet("Checklist")

    encabezados = ["N°", "Punto", "Resultado", "Páginas", "Referencia / Observación",
                   "Cita del modelo", "Página de la cita", "Verificación de la cita",
                   "Resultado del modelo", "Motivo del ajuste", "Lectura por partes"]
    for col_idx, nombre in enumerate(encabezados, 1):
        celda = ws.cell(row=1, column=col_idx, value=nombre)
        celda.font = Font(bold=True)

    for row_idx, item in enumerate(reporte.checklist.items, 2):
        fill = FILLS_CHECKLIST.get(item.resultado)
        valores = [
            item.numero,
            item.enunciado,
            TEXTO_RESULTADO[item.resultado],
            item.paginas or "",
            item.referencia or "",
            " | ".join(c.texto for c in item.citas),
            ", ".join(str(c.pagina) for c in item.citas if c.pagina),
            " | ".join(etiqueta_verificacion(c.verificacion) for c in item.citas),
            TEXTO_RESULTADO[item.resultado_modelo] if item.resultado_modelo else "",
            item.motivo_ajuste or "",
            item.aviso_tramos or "",
        ]
        for col_idx, valor in enumerate(valores, 1):
            celda = ws.cell(row=row_idx, column=col_idx, value=valor)
            if fill:
                celda.fill = fill

    ws.column_dimensions["B"].width = 90
    for letra in ("C", "D", "E", "G", "H", "I"):
        ws.column_dimensions[letra].width = 24
    for letra in ("F", "J", "K"):
        ws.column_dimensions[letra].width = 60


def exportar_feedback_guardado(reporte: ReporteInforme, run_id: str) -> str:
    """El .md con el feedback tal como está guardado (lo que bajan las descargas)."""
    return exportar_feedback_md(reporte, feedback_mod.cargar_feedback(run_id), run_id)


def exportar_feedback_md(reporte: ReporteInforme, feedback: FeedbackRun, run_id: str) -> str:
    """Markdown con resultados + respuestas del equipo revisor, como insumo para
    medir e iterar el sistema. Lo que más importa va primero, en la sección de
    contradicciones: los falsos cumple (el error caro), después los falsos no
    cumple y los resultados del comparador marcados ✗ con su motivo."""
    v_res = feedback.veredictos_resultados
    v_cl = feedback.veredictos_checklist
    respuestas = feedback.respuestas_checklist
    motivos_vigentes = feedback_mod.motivos_vigentes(feedback)
    r = reporte.resumen
    contradicciones: dict[str, list[str]] = {"falso_cumple": [], "falso_no_cumple": []}

    encabezado = [
        f"# Feedback de revisión — {run_id}",
        "",
        f"- Informe: {reporte.informe_id} | {reporte.coordinado} / {reporte.subestacion}",
        f"- Auditora: {reporte.empresa_auditora.value} | Modelo: {reporte.modelo_extraccion}",
        f"- Resumen del sistema: total={r.total} coinciden={r.coinciden} "
        f"(redacción distinta={r.coinciden_con_redaccion_distinta}) "
        f"difieren={r.difieren} faltan={r.faltan_en_excel} extra={r.extra_en_excel} "
        f"transversales={r.transversales}",
        "",
    ]
    lineas: list[str] = []

    def marca(v: str | None) -> str:
        return {"correcto": "✓ correcto", "incorrecto": "✗ INCORRECTO"}.get(v or "", "sin revisar")

    if reporte.checklist:
        cl = reporte.checklist
        lineas += [f"## Checklist ({cl.cumplen}/{cl.total} según el sistema)", ""]
        for item in cl.items:
            clave = str(item.numero)
            respuesta = respuestas.get(clave)
            contraste = _contraste_checklist(item.resultado, respuesta)
            if contraste in contradicciones:
                en_parte = (" (visto solo en parte del informe)"
                            if contraste == "falso_cumple" and item.aviso_tramos else "")
                contradicciones[contraste].append(
                    f"- checklist punto {item.numero}: {_ETIQUETA_CONTRASTE[contraste]}{en_parte} "
                    f"(sistema: {item.resultado.value}; revisora: {respuesta}) — {item.enunciado}")
            anterior = f" · feedback anterior: {marca(v_cl[clave])}" if clave in v_cl else ""
            lineas.append(
                f"- {item.numero}. Sistema: {item.resultado.value} · Revisora: "
                f"{respuesta or '—'} → {_ETIQUETA_CONTRASTE[contraste]}{anterior}"
                f" — {item.enunciado}"
                + (f" | obs: {item.referencia}" if item.referencia else "")
            )
            for c in item.citas:
                lineas.append(f"  - cita: «{c.texto}» ({etiqueta_verificacion(c.verificacion)})")
            if item.motivo_ajuste:
                lineas.append(f"  - bajado a {item.resultado.value}: {item.motivo_ajuste}")
            if item.aviso_tramos:
                lineas.append(f"  - lectura por partes: {item.aviso_tramos}")
        lineas.append("")

    lineas += ["## Resultados del comparador", ""]
    for idx, m in enumerate(reporte.resultados):
        desc = (m.hallazgo_pdf or m.hallazgo_excel)
        redaccion = " · redacción distinta" if m.equivalencias_redaccion else ""
        lineas.append(
            f"- #{idx} [{m.pano} | {m.rele}] {m.estado.value}{redaccion} "
            f"({m.tipo_match.value}, conf={m.confianza:.2f}) {marca(v_res.get(str(idx)))}"
            f" — {(desc.descripcion if desc else '')[:90]}"
        )
        motivo = _texto_motivo(motivos_vigentes, str(idx))
        if motivo:
            lineas.append(f"  - motivo del ✗: {motivo}")
        # Un emparejamiento dudoso de confianza alta (asignación global, PD71) no
        # se ve en el tipo de match: va su motivo.
        if m.emparejamiento_dudoso is not None:
            lineas.append(f"  - emparejamiento dudoso: {m.emparejamiento_dudoso}")
        # La viñeta copiada entera a cada relé (PD60) explica muchos ✗ de un «difiere».
        if m.hallazgo_pdf and m.hallazgo_pdf.evidencia.otros_reles_nombrados:
            lineas.append("  - la viñeta del PDF también nombra a: "
                          + ", ".join(m.hallazgo_pdf.evidencia.otros_reles_nombrados))
        # Una cita del PDF que no se pudo respaldar (o una fila oculta del Excel)
        # explica muchos ✗: va al lado del resultado.
        if m.hallazgo_pdf:
            ev = m.hallazgo_pdf.evidencia
            for campo, v in (("descripción", ev.verificacion_descripcion),
                             ("acción", ev.verificacion_accion)):
                if v is not None and not v.respaldada:
                    lineas.append(f"  - cita de la {campo} del PDF: {etiqueta_verificacion(v)}")
        if m.hallazgo_excel and m.hallazgo_excel.evidencia.fila_oculta:
            lineas.append(f"  - fila {m.hallazgo_excel.evidencia.fila_excel} del Excel oculta "
                          "(filtro activo)")
        # La equivalencia es una decisión automática: va completa, para que
        # un ✗ de la revisora sobre este resultado se pueda diagnosticar.
        for e in m.equivalencias_redaccion:
            lineas.append(
                f"  - equivalencia de redacción en {e.campo.value} (similitud "
                f"{e.similitud:.2f}): PDF «{e.valor_pdf}» · Excel «{e.valor_excel}» "
                f"— juez: {e.justificacion}"
            )

    incorrectos = [
        *contradicciones["falso_cumple"],
        *contradicciones["falso_no_cumple"],
        *(f"- resultado #{i}" + (f": {mot}" if (mot := _texto_motivo(motivos_vigentes, i)) else "")
          for i in sorted(v_res, key=int) if v_res[i] == "incorrecto"),
        *(f"- checklist punto {n} (feedback anterior ✗)"
          for n in sorted(v_cl, key=int) if v_cl[n] == "incorrecto"),
    ]
    prioridad = ["## Contradicciones y marcados INCORRECTOS (prioridad para mejorar)", ""]
    prioridad += incorrectos or ["(ninguno)"]
    motivos = Counter(f.motivo for f in motivos_vigentes.values() if f.motivo)
    if motivos:
        prioridad += ["", "## Motivos de los ✗ del comparador", ""]
        prioridad += [f"- {ETIQUETA_MOTIVO[motivo]}: {n}" for motivo, n in motivos.most_common()]

    lineas += ["", "## Comentarios de la revisora", "", feedback.comentarios or "(sin comentarios)"]
    if feedback.actualizado:
        lineas += ["", f"_Última actualización del feedback: {feedback.actualizado.isoformat()}_"]
    return "\n".join([*encabezado, *prioridad, "", *lineas]) + "\n"


_ETIQUETA_CONTRASTE = {
    "acierto": "coincide",
    "falso_cumple": "FALSO CUMPLE",
    "falso_no_cumple": "FALSO NO CUMPLE",
    "no_verificable": "el sistema no lo verificó",
    "no_aplica": "no aplica",
    "sin_respuesta": "sin respuesta",
}


def _contraste_checklist(resultado: ResultadoChecklist, respuesta: str | None) -> str:
    """Respuesta de la revisora contra el sistema. Solo SI/NO se clasifican:
    `clasificar` trata todo lo que no es SI como NO, y N/A no es un NO."""
    if respuesta in ("SI", "NO"):
        return clasificar(resultado, respuesta)
    return "no_aplica" if respuesta == "N/A" else "sin_respuesta"


def _texto_motivo(motivos: dict[str, MotivoFeedback], clave: str) -> str:
    m = motivos.get(clave)
    if m is None or (m.motivo is None and not m.detalle):
        return ""
    etiqueta = ETIQUETA_MOTIVO[m.motivo] if m.motivo else "sin motivo"
    return etiqueta + (f" — {m.detalle}" if m.detalle else "")


def exportar_csv(reporte: ReporteInforme) -> str:
    """Genera un CSV como string desde un ReporteInforme.

    Con BOM UTF-8: Excel en Windows (el entorno de la revisora) asume cp1252
    sin él y rompe tildes y eñes.
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(COLUMNAS)

    for resultado in reporte.resultados:
        writer.writerow(_fila_match(resultado))

    return "\ufeff" + output.getvalue()
