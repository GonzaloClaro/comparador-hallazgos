"""Etapa 2 (determinista, sin LLM) sobre los 7 Excel reales.

Conteos exactos verificados contra los Excel originales con el código actual
(2026-10-17): "hallazgos" son los Hallazgo normalizados (una fila con varios
paños se desdobla en uno por paño; una transversal no se desdobla) y
"filas" las filas distintas del Excel con contenido.
"""

from __future__ import annotations

from typing import Any

import pytest

from comparador.catalogo import cargar_hints_empresa
from comparador.etapa2_normalizar_excel import normalizar_excel
from comparador.schemas import Calificacion, EmpresaAuditora, Hallazgo
from tests.golden.datos import (
    SUBESTACIONES,
    casos,
    casos_informes,
    inventario,
    requiere_informes,
    requiere_mediciones,
    sin_campos_070,
)

pytestmark = [pytest.mark.golden, requiere_informes]

SIN_COLUMNAS: list[str] = []
SIN_ACCION_NI_CALIFICACION = ["Acción Correctiva del Hallazgo", "Calificación"]

# subestación: (empresa, hallazgos, filas, filas transversales, columnas
# ausentes, hallazgos por calificación {Menor, Relevante, None}).
ESPERADO: dict[str, tuple[EmpresaAuditora, int, int, list[int], list[str], dict[Any, int]]] = {
    "Epsilon": (EmpresaAuditora.ENSAUT, 23, 22, [35, 36, 37], SIN_ACCION_NI_CALIFICACION,
                {None: 23}),
    "Añañuca": (EmpresaAuditora.RELIABLE_ENERGY, 20, 20, [], SIN_COLUMNAS,
                {Calificacion.MENOR: 6, Calificacion.RELEVANTE: 14}),
    "Zeta": (EmpresaAuditora.RELIABLE_ENERGY, 17, 14, [], SIN_COLUMNAS,
                {Calificacion.MENOR: 3, Calificacion.RELEVANTE: 14}),
    "Alfa": (EmpresaAuditora.RELIABLE_ENERGY, 93, 93, [], SIN_COLUMNAS,
                 {Calificacion.MENOR: 43, Calificacion.RELEVANTE: 50}),
    "Beta": (EmpresaAuditora.ENSAUT, 223, 201, [221, 222, 223], SIN_COLUMNAS,
               {Calificacion.MENOR: 133, Calificacion.RELEVANTE: 90}),
    "Eta": (EmpresaAuditora.ENSAUT, 43, 30, [51, 52], SIN_COLUMNAS,
                 {Calificacion.MENOR: 23, Calificacion.RELEVANTE: 20}),
    "Gamma": (EmpresaAuditora.RELIABLE_ENERGY, 32, 26, [], SIN_COLUMNAS,
                  {Calificacion.MENOR: 11, Calificacion.RELEVANTE: 21}),
}


def _normalizar(sub: str) -> list[Hallazgo]:
    caso = casos_informes()[sub]
    hints = cargar_hints_empresa(caso["empresa"])
    hallazgos, metadatos = normalizar_excel(caso["excel"], hints.filas_transversales)
    assert metadatos.columnas_ausentes == ESPERADO[sub][4]
    return hallazgos


def test_estan_los_siete_informes_medidos() -> None:
    # Una carpeta nueva en informes_reales/ no entra (golden limitados al inventario).
    assert set(SUBESTACIONES) <= set(casos_informes())


@pytest.mark.parametrize("sub", SUBESTACIONES)
def test_conteos_exactos(sub: str) -> None:
    empresa, n_hallazgos, n_filas, transversales, _, por_calificacion = ESPERADO[sub]
    assert casos_informes()[sub]["empresa"] == empresa

    hallazgos = _normalizar(sub)

    assert len(hallazgos) == n_hallazgos
    assert len({h.evidencia.fila_excel for h in hallazgos}) == n_filas
    assert [h.evidencia.fila_excel for h in hallazgos if h.es_transversal] == transversales
    conteo: dict[Any, int] = {}
    for h in hallazgos:
        conteo[h.calificacion] = conteo.get(h.calificacion, 0) + 1
    assert conteo == por_calificacion
    # Una fila transversal es una sola entrada, aunque nombre varios paños.
    assert not any(h.es_transversal and h.evidencia.es_multi_pano_origen for h in hallazgos)


CASOS_B2 = casos("b2")


@requiere_mediciones
@pytest.mark.parametrize("sub", SUBESTACIONES)
def test_igual_a_lo_medido_en_b2(sub: str) -> None:
    """Mismo resultado que guardó cada run de B2 (los 3 modelos) y que el
    inventario de la medición."""
    hallazgos = [h.model_dump(mode="json") for h in _normalizar(sub)]

    runs = [c for c in CASOS_B2 if c.subestacion == sub]
    assert len(runs) == 3
    for caso in runs:
        assert sin_campos_070(caso.json("etapa2_excel_normalizado.json")) == sin_campos_070(
            hallazgos
        ), caso.id
        if caso.tiene("metadata.json"):
            meta = caso.json("metadata.json")
            assert meta["total_hallazgos_excel"] == len(hallazgos)
            assert meta["total_transversales_excel"] == sum(h["es_transversal"] for h in hallazgos)
    fila = inventario()[sub]
    assert fila["filas_excel_normalizadas"] == len(hallazgos)
    assert fila["filas_excel_transversales"] == sum(h["es_transversal"] for h in hallazgos)


# Hoja de hallazgos y filas ocultas (filtro activo) de cada Excel real,
# medidas el 2026-10-18 abriendo cada archivo con openpyxl.
HOJA_Y_OCULTAS = {
    "Epsilon": ("Tabla de Hallazgos", 0),
    "Añañuca": ("Tabla de Hallazgos", 0),
    "Zeta": ("Tabla de Hallazgos", 0),
    "Alfa": ("Tabla de Hallazgos", 0),
    "Beta": ("Tabla de Hallazgos", 0),
    # Eta trae un autofiltro activo: 22 de sus 30 filas no se ven al abrirlo.
    "Eta": ("Tabla de Hallazgos (2)", 29),
    "Gamma": ("Tabla de Hallazgos", 0),
}


def _plano(valor: Any) -> str:
    return " ".join(str(valor).split()) if valor is not None else ""


@pytest.mark.parametrize("sub", SUBESTACIONES)
def test_cada_referencia_apunta_a_su_celda(sub: str) -> None:
    """Hoja + fila llevan a la celda real: la descripción y el relé que se ven ahí
    (resolviendo celdas combinadas) son los del hallazgo. Lectura independiente
    de la Etapa 2: openpyxl celda por celda."""
    from openpyxl import load_workbook

    hallazgos = _normalizar(sub)
    hoja, ocultas = HOJA_Y_OCULTAS[sub]
    assert {h.evidencia.hoja_excel for h in hallazgos} == {hoja}
    assert sum(h.evidencia.fila_oculta for h in hallazgos) == ocultas

    ws = load_workbook(casos_informes()[sub]["excel"], data_only=True)[hoja]
    fila_header, col = next(
        (f, {_plano(c.value): c.column for c in ws[f] if c.value is not None})
        for f in range(1, 31)
        if any(_plano(c.value) == "Descripción del Hallazgo" for c in ws[f])
    )
    col_rele = col.get("Relé") or col["Equipo"]

    def visible(fila: int, columna: int) -> str:
        for rango in ws.merged_cells.ranges:
            if rango.min_row <= fila <= rango.max_row and rango.min_col <= columna <= rango.max_col:
                return _plano(ws.cell(rango.min_row, rango.min_col).value)
        return _plano(ws.cell(fila, columna).value)

    for h in hallazgos:
        fila = h.evidencia.fila_excel
        assert fila is not None and fila > fila_header
        assert _plano(ws.cell(fila, col["Descripción del Hallazgo"]).value) == _plano(
            h.descripcion
        ), (sub, fila)
        assert visible(fila, col_rele) == _plano(h.rele), (sub, fila)
        oculta = ws.row_dimensions.get(fila)
        assert h.evidencia.fila_oculta == bool(oculta is not None and oculta.hidden), (sub, fila)
