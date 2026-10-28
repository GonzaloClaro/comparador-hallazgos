"""Tests de exportación: CSV con BOM y xlsx con hoja Resumen."""

import io

from openpyxl import load_workbook

from comparador.etapa5_consolidar import consolidar_reporte
from comparador.exportar import exportar_csv, exportar_feedback_md, exportar_xlsx
from comparador.schemas import (
    CampoDiferente,
    Diferencia,
    EmpresaAuditora,
    EquivalenciaRedaccion,
    EstadoHallazgo,
    FeedbackRun,
    MatchResult,
    MetadatosInforme,
    ReporteInforme,
    TipoMatch,
)


def _reporte() -> object:
    resultados = [
        MatchResult(
            pano="C1", rele="SEL 351R",
            estado=EstadoHallazgo.DIFIERE, tipo_match=TipoMatch.UNICIDAD,
            confianza=1.0, justificacion="Única 'Verificación DUF' en la tupla.",
            hallazgo_pdf=None, hallazgo_excel=None,
            diferencias=[Diferencia(
                campo=CampoDiferente.DESCRIPCION,
                valor_pdf="versión PDF", valor_excel="versión Excel",
                detalle="Similitud semántica 0.40, bajo el umbral de coincidencia 0.85",
            )],
        ),
    ]
    metadatos = MetadatosInforme(
        coordinado="COORDINADO A", subestacion="Iota", informe_id="INF-001",
    )
    return consolidar_reporte(resultados, metadatos, EmpresaAuditora.ENSAUT)


class TestCsv:
    def test_empieza_con_bom_utf8(self):
        # Sin BOM, Excel en Windows asume cp1252 y rompe tildes y eñes.
        csv = exportar_csv(_reporte())
        assert csv.startswith("\ufeff")

    def test_campos_diferentes_incluyen_detalle(self):
        csv = exportar_csv(_reporte())
        assert "descripcion (Similitud semántica 0.40" in csv


class TestXlsx:
    def test_hoja_resumen_identifica_el_archivo(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(_reporte())))
        assert wb.sheetnames[0] == "Resumen"
        valores = {ws[0]: ws[1] for ws in wb["Resumen"].iter_rows(values_only=True)}
        assert valores["Informe"] == "INF-001"
        assert valores["Coordinado"] == "COORDINADO A"
        assert valores["Total de hallazgos"] == 1

    def test_tabla_con_freeze_y_autofiltro(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(_reporte())))
        ws = wb["Comparación Hallazgos"]
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref is not None


class TestTransversalesExport:
    def _reporte_con_transversal(self) -> object:
        resultados = [
            MatchResult(
                pano="BT2", rele="Análisis de coordinación",
                estado=EstadoHallazgo.TRANSVERSAL, tipo_match=TipoMatch.SIN_MATCH,
                confianza=0.0, justificacion="Fila transversal, revisión manual.",
                hallazgo_pdf=None, hallazgo_excel=None,
            ),
        ]
        metadatos = MetadatosInforme(
            coordinado="COORDINADO A", subestacion="Iota", informe_id="INF-001",
        )
        return consolidar_reporte(resultados, metadatos, EmpresaAuditora.ENSAUT)

    def test_fila_transversal_exporta_con_estado_y_color(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte_con_transversal())))
        ws = wb["Comparación Hallazgos"]
        fila = list(ws.iter_rows(min_row=2, max_row=2))[0]
        estado_col = [c.value for c in list(ws.iter_rows(min_row=1, max_row=1))[0]].index("Estado")
        assert fila[estado_col].value == "transversal"
        assert fila[0].fill.start_color.rgb.endswith("D9D9D9")

    def test_hoja_resumen_incluye_transversales(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte_con_transversal())))
        valores = {ws[0]: ws[1] for ws in wb["Resumen"].iter_rows(values_only=True)}
        assert valores["Transversales (revisión manual)"] == 1


class TestEquivalenciasRedaccionExport:
    """La decisión automática del juez de redacción no se pierde en la descarga."""

    def _reporte_con_equivalencia(self) -> ReporteInforme:
        resultados = [
            MatchResult(
                pano="C1", rele="COOPER FORM 6",
                estado=EstadoHallazgo.COINCIDE, tipo_match=TipoMatch.UNICIDAD,
                confianza=1.0, justificacion="Única en la tupla.",
                hallazgo_pdf=None, hallazgo_excel=None,
                equivalencias_redaccion=[EquivalenciaRedaccion(
                    campo=CampoDiferente.DESCRIPCION,
                    valor_pdf="Sincronización horaria externa. No cumple.",
                    valor_excel="Equipo no posee sincronización horaria vía GPS.",
                    similitud=0.8012,
                    justificacion="Mismo incumplimiento; el Excel precisa que es GPS.",
                )],
            ),
        ]
        metadatos = MetadatosInforme(
            coordinado="COORDINADO A", subestacion="Zeta", informe_id="INF-002",
        )
        return consolidar_reporte(resultados, metadatos, EmpresaAuditora.RELIABLE_ENERGY)

    def test_xlsx_columna_de_equivalencias(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte_con_equivalencia())))
        ws = wb["Comparación Hallazgos"]
        encabezados = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        fila = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
        valor = fila[encabezados.index("Equivalencias de redacción")]
        assert valor == (
            "descripcion (similitud 0.80; juez de redacción: Mismo incumplimiento; "
            "el Excel precisa que es GPS.)"
        )
        assert fila[encabezados.index("Estado")] == "coincide"

    def test_hoja_resumen_cuenta_las_coincidencias_con_redaccion_distinta(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte_con_equivalencia())))
        valores = {ws[0]: ws[1] for ws in wb["Resumen"].iter_rows(values_only=True)}
        assert valores["Coinciden"] == 1
        assert valores["Coinciden con redacción distinta (de las anteriores)"] == 1

    def test_csv_incluye_la_equivalencia(self):
        assert "juez de redacción: Mismo incumplimiento" in exportar_csv(
            self._reporte_con_equivalencia()
        )

    def test_feedback_md_muestra_la_equivalencia_completa(self):
        md = exportar_feedback_md(
            self._reporte_con_equivalencia(), FeedbackRun(), "run-x",
        )
        assert "(redacción distinta=1)" in md
        assert "coincide · redacción distinta (unicidad" in md
        assert (
            "  - equivalencia de redacción en descripcion (similitud 0.80): "
            "PDF «Sincronización horaria externa. No cumple.» · Excel «Equipo no "
            "posee sincronización horaria vía GPS.» — juez: Mismo incumplimiento; "
            "el Excel precisa que es GPS."
        ) in md


class TestReferenciasEnLaDescarga:
    """Schema 0.7.0: el estado de cada cita, la hoja y la fila oculta viajan en
    la descarga, para que la referencia no dependa del dashboard."""

    def _reporte(self) -> ReporteInforme:
        from comparador.schemas import (
            EstadoCita,
            Evidencia,
            FuenteEvidencia,
            Hallazgo,
            VerificacionCita,
        )

        pdf = Hallazgo(
            pano="C2", rele="SEL 351", clasificacion="Ajustes",
            descripcion="Ajuste de la función 50: No Cumplir", accion_correctiva="Ajustar",
            evidencia=Evidencia(
                fuente=FuenteEvidencia.PDF, pagina_pdf=14, texto_crudo="x",
                verificacion_descripcion=VerificacionCita(
                    estado=EstadoCita.NO_ENCONTRADA, pagina_citada=14),
                verificacion_accion=VerificacionCita(estado=EstadoCita.VERIFICADA, pagina=14),
            ),
        )
        excel = Hallazgo(
            pano="C2", rele="SEL 351", clasificacion="Ajustes",
            descripcion="Ajuste de la función 50: No Cumple", accion_correctiva="Ajustar",
            evidencia=Evidencia(fuente=FuenteEvidencia.EXCEL, fila_excel=27, texto_crudo="x",
                                hoja_excel="Tabla de Hallazgos (2)", fila_oculta=True),
        )
        resultados = [MatchResult(
            pano="C2", rele="SEL 351", estado=EstadoHallazgo.COINCIDE,
            tipo_match=TipoMatch.UNICIDAD, confianza=1.0, justificacion="única",
            hallazgo_pdf=pdf, hallazgo_excel=excel,
        )]
        metadatos = MetadatosInforme(coordinado="c", subestacion="s", informe_id="i")
        return consolidar_reporte(resultados, metadatos, EmpresaAuditora.ENSAUT)

    def test_xlsx_y_csv_traen_las_columnas_de_referencia(self):
        r = self._reporte()
        ws = load_workbook(io.BytesIO(exportar_xlsx(r)))["Comparación Hallazgos"]
        filas = list(ws.iter_rows(values_only=True))
        fila = dict(zip(filas[0], filas[1], strict=True))
        assert fila["Cita descripción PDF"] == "no aparece en la p. 14"
        assert fila["Cita acción PDF"] == "verificada en la p. 14"
        assert (fila["Hoja Excel"], fila["Fila Excel"], fila["Fila Excel oculta"]) == (
            "Tabla de Hallazgos (2)", 27, "sí")
        assert "no aparece en la p. 14" in exportar_csv(r)

    def test_resumen_cuenta_las_filas_ocultas(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte())))
        valores = {f[0]: f[1] for f in wb["Resumen"].iter_rows(values_only=True)}
        assert valores["Filas del Excel ocultas (filtro activo)"] == 1

    def test_feedback_md_explica_la_cita_y_la_fila_oculta(self):
        md = exportar_feedback_md(self._reporte(), FeedbackRun(), "run")
        assert "cita de la descripción del PDF: no aparece en la p. 14" in md
        assert "fila 27 del Excel oculta" in md
        assert "cita de la acción" not in md   # la acción sí está respaldada

    def test_reporte_anterior_sin_verificacion(self):
        csv = exportar_csv(_reporte())   # sin hallazgos: columnas vacías, sin errores
        assert csv.splitlines()[0].endswith(
            "Cita descripción PDF,Cita acción PDF,Hoja Excel,Fila Excel oculta,"
            "Emparejamiento dudoso,Otros relés nombrados (PDF)")


MOTIVO = ("Revisar a qué hallazgo corresponde esta fila. El Excel dice que no cumple, pero en "
          "la misma pág. 26 hay otro hallazgo del PDF que dice que cumple (descripción "
          "«Comparación con ECAP adjunto en INFOTÉCNICA: Cumple») y que también podría ser el "
          "de esta fila. Si es ese, el Excel invirtió el veredicto.")


class TestEmparejamientoDudosoEnLaDescarga:
    """Schema 0.9.0 (PD71): el par con `emparejamiento_dudoso` va en la descarga
    como un semantico_bajo: una columna con su motivo y el conteo «Requieren
    revisión» del Resumen. La columna aparece siempre, vacía cuando ningún par trae
    el campo (PD72, decisión del usuario 2026-10-26; antes, solo si algún par lo
    traía). Desde PD60 (schema 0.10.0) la sigue «Otros relés nombrados (PDF)»."""

    def _reporte(self, motivo: str | None = MOTIVO) -> ReporteInforme:
        def par(tipo: TipoMatch, confianza: float, motivo: str | None) -> MatchResult:
            return MatchResult(
                pano="E7", rele="SEL 451", estado=EstadoHallazgo.COINCIDE, tipo_match=tipo,
                confianza=confianza, justificacion="j", hallazgo_pdf=None,
                hallazgo_excel=None, emparejamiento_dudoso=motivo)

        resultados = [par(TipoMatch.SEMANTICO_ALTO, 0.8, motivo),
                      par(TipoMatch.SEMANTICO_BAJO, 0.62, None),
                      par(TipoMatch.SEMANTICO_ALTO, 0.9, None),
                      par(TipoMatch.UNICIDAD, 1.0, None)]
        metadatos = MetadatosInforme(coordinado="c", subestacion="s", informe_id="i")
        return consolidar_reporte(resultados, metadatos, EmpresaAuditora.RELIABLE_ENERGY)

    def test_xlsx_y_csv_con_la_columna_penultima(self):
        # PD60 suma «Otros relés nombrados (PDF)» después: la columna queda penúltima.
        r = self._reporte()
        ws = load_workbook(io.BytesIO(exportar_xlsx(r)))["Comparación Hallazgos"]
        filas = list(ws.iter_rows(values_only=True))
        assert filas[0][-2] == "Emparejamiento dudoso"
        # Solo el motivo: el semantico_bajo ya lo dice la columna «Tipo Match».
        por_confianza = {fila[5]: fila[-2] for fila in filas[1:]}
        assert por_confianza == {0.8: MOTIVO, 0.62: None, 0.9: None, 1.0: None}
        assert MOTIVO in exportar_csv(r)

    def test_sin_el_campo_la_columna_esta_y_va_vacia(self):
        # Sin pares dudosos (la asignación apagada, o ninguna fila disputada), también
        # con un semantico_bajo: la columna sigue después de «Fila Excel oculta», y va
        # vacía en todas las filas (el semantico_bajo ya lo dice «Tipo Match»). PD72:
        # antes no aparecía y las columnas eran las 21 del 0.8.0; PD60 suma la 23.
        r = self._reporte(motivo=None)
        ws = load_workbook(io.BytesIO(exportar_xlsx(r)))["Comparación Hallazgos"]
        filas = list(ws.iter_rows(values_only=True))
        assert (len(filas[0]), filas[0][-3:-1]) == (23, ("Fila Excel oculta",
                                                         "Emparejamiento dudoso"))
        assert ws.max_column == 23
        assert [fila[-2] for fila in filas[1:]] == [None] * 4
        csv = exportar_csv(r)
        lineas = csv.splitlines()
        assert lineas[0].endswith(",Hoja Excel,Fila Excel oculta,Emparejamiento dudoso,"
                                  "Otros relés nombrados (PDF)")
        assert all(linea.count(",") == lineas[0].count(",") for linea in lineas[1:])
        assert all(linea.endswith(",,") for linea in lineas[1:])
        assert "Confianza baja" not in csv
        assert "emparejamiento dudoso" not in exportar_feedback_md(r, FeedbackRun(), "run")

    def test_resumen_cuenta_el_par_dudoso_en_requieren_revision(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte())))
        valores = {f[0]: f[1] for f in wb["Resumen"].iter_rows(values_only=True)}
        assert valores["Requieren revisión"] == 2

    def test_feedback_md_trae_el_motivo(self):
        md = exportar_feedback_md(self._reporte(), FeedbackRun(), "run")
        assert md.count("emparejamiento dudoso") == 1
        assert f"  - emparejamiento dudoso: {MOTIVO}" in md


class TestChecklistConCitaEnLaDescarga:
    def _reporte(self) -> ReporteInforme:
        from comparador.schemas import (
            ChecklistInforme,
            CitaChecklist,
            EstadoCita,
            ItemChecklist,
            ResultadoChecklist,
            VerificacionCita,
        )

        checklist = ChecklistInforme(items=[
            ItemChecklist(numero=6, enunciado="DUF", resultado=ResultadoChecklist.NO_VERIFICABLE,
                          resultado_modelo=ResultadoChecklist.CUMPLE,
                          motivo_ajuste="cita no verificada",
                          citas=[CitaChecklist(pagina=12, texto="Diagrama unilineal",
                                               verificacion=VerificacionCita(
                                                   estado=EstadoCita.NO_ENCONTRADA,
                                                   pagina_citada=12))]),
        ], cumplen=0, total=38, bajados_sin_cita_verificada=1)
        metadatos = MetadatosInforme(coordinado="c", subestacion="s", informe_id="i")
        return consolidar_reporte([], metadatos, EmpresaAuditora.ENSAUT, checklist)

    def test_hoja_checklist_y_resumen(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte())))
        filas = list(wb["Checklist"].iter_rows(values_only=True))
        fila = dict(zip(filas[0], filas[1], strict=True))
        assert fila["Resultado"] == "No es posible verificar"
        assert (fila["Cita del modelo"], fila["Página de la cita"]) == ("Diagrama unilineal", "12")
        assert fila["Verificación de la cita"] == "no aparece en la p. 12"
        assert (fila["Resultado del modelo"], fila["Motivo del ajuste"]) == (
            "Cumple", "cita no verificada")
        resumen = {f[0]: f[1] for f in wb["Resumen"].iter_rows(values_only=True)}
        assert resumen["Checklist: «cumple» bajados a no verificable sin cita verificada"] == 1

    def test_feedback_md(self):
        md = exportar_feedback_md(self._reporte(), FeedbackRun(), "run")
        assert "cita: «Diagrama unilineal» (no aparece en la p. 12)" in md
        assert "bajado a no_verificable: cita no verificada" in md


def test_reporte_anterior_no_afirma_filas_visibles():
    """Un reporte < 0.7.0 no revisó las filas ocultas: la descarga no dice "no" ni 0."""
    from comparador.schemas import Evidencia, FuenteEvidencia, Hallazgo

    excel = Hallazgo(pano="C1", rele="R", clasificacion="x", descripcion="d",
                     evidencia=Evidencia(fuente=FuenteEvidencia.EXCEL, fila_excel=7,
                                         texto_crudo="d"))
    datos = consolidar_reporte([MatchResult(
        pano="C1", rele="R", estado=EstadoHallazgo.EXTRA_EN_EXCEL, tipo_match=TipoMatch.SIN_MATCH,
        confianza=0.0, justificacion="x", hallazgo_pdf=None, hallazgo_excel=excel,
    )], MetadatosInforme(coordinado="c", subestacion="s", informe_id="i"),
        EmpresaAuditora.ENSAUT).model_dump(mode="json")
    from tests.golden.datos import sin_campos_070

    viejo = ReporteInforme.model_validate(sin_campos_070(datos))
    wb = load_workbook(io.BytesIO(exportar_xlsx(viejo)))
    filas = list(wb["Comparación Hallazgos"].iter_rows(values_only=True))
    fila = dict(zip(filas[0], filas[1], strict=True))
    assert fila["Fila Excel oculta"] in (None, "")
    resumen = {f[0] for f in wb["Resumen"].iter_rows(values_only=True)}
    assert "Filas del Excel ocultas (filtro activo)" not in resumen


class TestFeedbackDelEquipoEnLaDescarga:
    """Respuestas SI/NO/N/A del checklist contra el sistema y motivos de los ✗ (F6g)."""

    def _reporte(self) -> ReporteInforme:
        from comparador.schemas import ChecklistInforme, ItemChecklist, ResultadoChecklist

        cumple, no, nv = (ResultadoChecklist.CUMPLE, ResultadoChecklist.NO_CUMPLE,
                          ResultadoChecklist.NO_VERIFICABLE)
        items = [ItemChecklist(numero=n, enunciado=f"Punto {n}", resultado=res)
                 for n, res in enumerate([cumple, cumple, no, nv, cumple, no], start=1)]
        reporte = _reporte()
        assert isinstance(reporte, ReporteInforme)
        return reporte.model_copy(update={
            "checklist": ChecklistInforme(items=items, cumplen=3, total=38)})

    def _feedback(self) -> FeedbackRun:
        from comparador.schemas import MotivoFeedback, MotivoIncorrecto

        return FeedbackRun(
            # 1 acierto · 2 falso cumple · 3 falso no cumple · 4 no verificado
            # · 5 N/A · 6 sin respuesta
            respuestas_checklist={"1": "SI", "2": "NO", "3": "SI", "4": "SI", "5": "N/A"},
            veredictos_resultados={"0": "incorrecto"},
            motivos_resultados={"0": MotivoFeedback(
                motivo=MotivoIncorrecto.MISMA_INFORMACION, detalle="solo cambia la redacción")},
        )

    def test_cada_punto_contrasta_al_sistema_con_la_revisora(self):
        md = exportar_feedback_md(self._reporte(), self._feedback(), "run")
        assert "- 1. Sistema: cumple · Revisora: SI → coincide — Punto 1" in md
        assert "- 2. Sistema: cumple · Revisora: NO → FALSO CUMPLE — Punto 2" in md
        assert "- 3. Sistema: no_cumple · Revisora: SI → FALSO NO CUMPLE — Punto 3" in md
        assert ("- 4. Sistema: no_verificable · Revisora: SI → el sistema no lo verificó"
                in md)
        # N/A no es un NO: no cuenta como falso cumple.
        assert "- 5. Sistema: cumple · Revisora: N/A → no aplica — Punto 5" in md
        assert "- 6. Sistema: no_cumple · Revisora: — → sin respuesta — Punto 6" in md

    def test_contradicciones_primero_los_falsos_cumple(self):
        md = exportar_feedback_md(self._reporte(), self._feedback(), "run")
        seccion = md.split("## Contradicciones y marcados INCORRECTOS")[1].split("\n## ")[0]
        lineas = [linea for linea in seccion.splitlines() if linea.startswith("- ")]
        assert lineas == [
            "- checklist punto 2: FALSO CUMPLE (sistema: cumple; revisora: NO) — Punto 2",
            "- checklist punto 3: FALSO NO CUMPLE (sistema: no_cumple; revisora: SI) — Punto 3",
            "- resultado #0: Es la misma información, con otra redacción — solo cambia la "
            "redacción",
        ]

    def test_motivos_junto_al_resultado_y_contados(self):
        md = exportar_feedback_md(self._reporte(), self._feedback(), "run")
        assert ("  - motivo del ✗: Es la misma información, con otra redacción — solo cambia "
                "la redacción") in md
        assert "## Motivos de los ✗ del comparador\n\n- Es la misma información, con otra " \
               "redacción: 1" in md

    def test_sin_respuestas_no_hay_contradicciones(self):
        md = exportar_feedback_md(self._reporte(), FeedbackRun(), "run")
        assert "FALSO" not in md
        assert "(ninguno)" in md
        assert "## Motivos de los ✗" not in md

    def test_contradicciones_van_antes_que_el_detalle(self):
        md = exportar_feedback_md(self._reporte(), self._feedback(), "run")
        assert (md.index("## Contradicciones") < md.index("## Motivos de los ✗")
                < md.index("## Checklist") < md.index("## Resultados del comparador")
                < md.index("## Comentarios"))

    def test_un_motivo_huerfano_no_se_muestra_ni_se_cuenta(self):
        # Un ✗ que pasó a ✓ en la UI vanilla (o en un feedback viejo) puede
        # conservar su motivo: no es un ✗ y no cuenta para la métrica.
        fb = self._feedback().model_copy(update={"veredictos_resultados": {"0": "correcto"}})
        md = exportar_feedback_md(self._reporte(), fb, "run")
        assert "motivo del ✗" not in md
        assert "## Motivos de los ✗" not in md

    def test_la_descarga_lee_lo_guardado(self, tmp_path, monkeypatch):
        from comparador import feedback
        from comparador.config import settings
        from comparador.exportar import exportar_feedback_guardado

        monkeypatch.setattr(settings, "runs_dir", tmp_path / "runs")
        feedback.guardar_feedback("run", self._feedback())
        md = exportar_feedback_guardado(self._reporte(), "run")
        assert md == exportar_feedback_md(self._reporte(), feedback.cargar_feedback("run"), "run")
        assert "FALSO CUMPLE" in md

    def test_feedback_anterior_del_checklist_se_conserva(self):
        fb = FeedbackRun(veredictos_checklist={"2": "incorrecto"})
        md = exportar_feedback_md(self._reporte(), fb, "run")
        assert "→ sin respuesta · feedback anterior: ✗ INCORRECTO — Punto 2" in md
        assert "- checklist punto 2 (feedback anterior ✗)" in md


class TestLecturaPorPartesEnLaDescarga:
    """M08 (schema 0.8.0) en las descargas: filas del Resumen solo si el informe se
    leyó por partes, columna K de la hoja Checklist y la línea del .md por punto."""

    PARTES = [(1, 228), (229, 456)]

    def _reporte(self, *, bajar: bool = False) -> ReporteInforme:
        from comparador.etapa0_checklist import aviso_un_tramo, motivo_un_tramo
        from comparador.schemas import (
            CausaAjuste,
            ChecklistInforme,
            ItemChecklist,
            ResultadoChecklist,
        )

        siete: dict[str, object] = {"resultado": ResultadoChecklist.CUMPLE}
        if bajar:
            siete = {"resultado": ResultadoChecklist.NO_VERIFICABLE,
                     "resultado_modelo": ResultadoChecklist.CUMPLE,
                     "motivo_ajuste": motivo_un_tramo([2], self.PARTES),
                     "causa_ajuste": CausaAjuste.UN_TRAMO}
        items = [
            ItemChecklist(numero=1, enunciado="Portada", resultado=ResultadoChecklist.CUMPLE,
                          tramos_respaldo=[1, 2]),
            ItemChecklist(numero=3, enunciado="DUF", resultado=ResultadoChecklist.CUMPLE,
                          tramos_respaldo=[1], aviso_tramos=aviso_un_tramo([1], self.PARTES)),
            ItemChecklist(numero=7, enunciado="Planilla", tramos_respaldo=[2],
                          aviso_tramos=aviso_un_tramo([2], self.PARTES), **siete),
        ]
        checklist = ChecklistInforme(
            items=items, cumplen=2 if bajar else 3, total=38, tramos_leidos=self.PARTES,
            marcados_un_tramo=1 if bajar else 2, bajados_un_tramo=1 if bajar else None)
        metadatos = MetadatosInforme(coordinado="c", subestacion="s", informe_id="i")
        return consolidar_reporte([], metadatos, EmpresaAuditora.ENSAUT, checklist)

    def test_hoja_checklist_columna_k_al_final(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte())))
        filas = list(wb["Checklist"].iter_rows(values_only=True))
        assert filas[0][10] == "Lectura por partes" and len(filas[0]) == 11
        assert [f[10] for f in filas[1:]] == [
            None,
            "Visto en 1 de 2 partes del informe (págs. 1-228); las págs. 229-456 no lo "
            "respaldan: revísalas antes de darlo por cumplido.",
            "Visto en 1 de 2 partes del informe (págs. 229-456); las págs. 1-228 no lo "
            "respaldan: revísalas antes de darlo por cumplido.",
        ]

    def test_resumen_en_modo_marcar(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte())))
        resumen = {f[0]: f[1] for f in wb["Resumen"].iter_rows(values_only=True)}
        assert resumen["Checklist: informe leído por partes (págs de cada parte)"] == (
            "1-228, 229-456")
        assert resumen["Checklist: «cumple» vistos solo en parte del informe (con aviso)"] == 2
        assert "Checklist: «cumple» bajados a no verificable por lectura en partes" \
            not in resumen

    def test_resumen_en_modo_bajar(self):
        wb = load_workbook(io.BytesIO(exportar_xlsx(self._reporte(bajar=True))))
        resumen = {f[0]: f[1] for f in wb["Resumen"].iter_rows(values_only=True)}
        assert resumen["Checklist: «cumple» vistos solo en parte del informe (con aviso)"] == 1
        assert resumen["Checklist: «cumple» bajados a no verificable por lectura en partes"] == 1

    def test_feedback_md_con_la_lectura_por_partes(self):
        fb = FeedbackRun(respuestas_checklist={"1": "NO", "3": "NO", "7": "SI"})
        md = exportar_feedback_md(self._reporte(), fb, "run")
        assert ("  - lectura por partes: Visto en 1 de 2 partes del informe (págs. 1-228); "
                "las págs. 229-456 no lo respaldan: revísalas antes de darlo por cumplido."
                in md)
        assert md.count("  - lectura por partes: ") == 2
        seccion = md.split("## Contradicciones y marcados INCORRECTOS")[1].split("\n## ")[0]
        assert ("- checklist punto 3: FALSO CUMPLE (visto solo en parte del informe) "
                "(sistema: cumple; revisora: NO) — DUF") in seccion
        # El 1 lo respaldan las dos partes: falso cumple sin el sufijo.
        assert ("- checklist punto 1: FALSO CUMPLE (sistema: cumple; revisora: NO) — Portada"
                in seccion)

    def test_feedback_md_en_modo_bajar(self):
        md = exportar_feedback_md(self._reporte(bajar=True), FeedbackRun(), "run")
        assert ("  - bajado a no_verificable: El informe se leyó en partes por su longitud; "
                "este punto exige revisar todo el informe: revisa las páginas 1-228.") in md
        assert md.count("  - lectura por partes: ") == 2

    def test_reporte_leido_entero_o_anterior_sin_lo_nuevo(self):
        """Un checklist leído entero (o un reporte < 0.8.0) no suma filas al Resumen
        ni líneas al .md; la columna K queda vacía."""
        reporte = TestChecklistConCitaEnLaDescarga()._reporte()
        wb = load_workbook(io.BytesIO(exportar_xlsx(reporte)))
        resumen = {f[0] for f in wb["Resumen"].iter_rows(values_only=True)}
        assert not any("por partes" in str(etiqueta) for etiqueta in resumen)
        filas = list(wb["Checklist"].iter_rows(values_only=True))
        assert filas[0][10] == "Lectura por partes"
        assert all(f[10] is None for f in filas[1:])
        fb = FeedbackRun(respuestas_checklist={"6": "NO"})
        md = exportar_feedback_md(reporte, fb, "run")
        assert "lectura por partes" not in md and "visto solo en parte" not in md


class TestOtrosRelesNombradosEnLaDescarga:
    """Schema 0.10.0 (PD60, decisión del usuario 2026-10-26, D4a): los otros relés de
    la sección que nombra el hallazgo del PDF van en una columna fija, siempre, al
    final (vacía con la lista None o []), y en una línea del .md."""

    ZT, MD = "INGEPAC EF-ZT", "INGEPAC EF-MD"

    def _reporte(self, *listas: list[str] | None) -> ReporteInforme:
        from comparador.schemas import Evidencia, FuenteEvidencia, Hallazgo

        def pdf(otros: list[str] | None) -> Hallazgo:
            return Hallazgo(
                pano="BT2", rele="INGEPAC EF-TD", clasificacion="Documentación Técnica",
                descripcion="El ECAP presenta diferencias", accion_correctiva="Actualizar",
                evidencia=Evidencia(fuente=FuenteEvidencia.PDF, pagina_pdf=16,
                                    texto_crudo="x", otros_reles_nombrados=otros))

        resultados = [MatchResult(
            pano="BT2", rele="INGEPAC EF-TD", estado=EstadoHallazgo.FALTA_EN_EXCEL,
            tipo_match=TipoMatch.SIN_MATCH, confianza=0.0, justificacion="sin par",
            hallazgo_pdf=pdf(otros), hallazgo_excel=None) for otros in listas]
        metadatos = MetadatosInforme(coordinado="c", subestacion="s", informe_id="i")
        return consolidar_reporte(resultados, metadatos, EmpresaAuditora.RELIABLE_ENERGY)

    def test_columna_al_final_con_los_nombres(self):
        r = self._reporte([self.ZT, self.MD], [], None)
        ws = load_workbook(io.BytesIO(exportar_xlsx(r)))["Comparación Hallazgos"]
        filas = list(ws.iter_rows(values_only=True))
        assert filas[0][-1] == "Otros relés nombrados (PDF)"
        assert [f[-1] for f in filas[1:]] == ["INGEPAC EF-ZT; INGEPAC EF-MD", None, None]
        assert "INGEPAC EF-ZT; INGEPAC EF-MD" in exportar_csv(r)

    def test_sin_hallazgos_la_columna_esta(self):
        r = _reporte()   # sin hallazgos: la columna está y va vacía
        csv = exportar_csv(r)
        assert csv.splitlines()[0].endswith(",Otros relés nombrados (PDF)")

    def test_feedback_md_trae_la_linea(self):
        md = exportar_feedback_md(self._reporte([self.ZT, self.MD], [], None),
                                  FeedbackRun(), "run")
        assert "  - la viñeta del PDF también nombra a: INGEPAC EF-ZT, INGEPAC EF-MD" in md
        assert md.count("también nombra a") == 1
