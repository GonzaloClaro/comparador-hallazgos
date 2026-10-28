from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from comparador.etapa2_normalizar_excel import (
    ExcelNormalizationError,
    normalizar_excel,
)
from comparador.schemas import Calificacion, FuenteEvidencia


def _crear_excel_basico(
    path: Path,
    filas: list[list[str]],
    metadatos: dict[str, str] | None = None,
    headers: list[str] | None = None,
) -> Path:
    """Crea un Excel sintético que imita la estructura real de la SEC."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None

    meta = metadatos or {
        "Coordinado": "COORDINADO A",
        "Subestación": "Iota",
        "Informe de Auditoría": "RELE-MAR24-PPP-000",
        "Fecha de Auditoría": "12-03-2024 a 22-03-2024",
    }

    row_num = 1
    for clave, valor in meta.items():
        ws.cell(row=row_num, column=1, value=clave)
        ws.cell(row=row_num, column=2, value=valor)
        row_num += 1

    row_num += 1  # fila en blanco

    headers = headers or [
        "Paño(s)", "Relé", "ID Infotécnica",
        "Clasificación del Hallazgo", "Descripción del Hallazgo",
        "Acción Correctiva del Hallazgo", "Calificación",
    ]
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=row_num, column=col_idx, value=header)

    for fila in filas:
        row_num += 1
        for col_idx, valor in enumerate(fila, 1):
            ws.cell(row=row_num, column=col_idx, value=valor)

    filepath = path / "test.xlsx"
    wb.save(filepath)
    return filepath


class TestNormalizarExcelBasico:
    def test_un_hallazgo_simple(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "No se encuentran actualizadas las relaciones de transformación en DUF.",
             "Se debe actualizar la relación de transformación del DUF.",
             "Menor"],
        ])

        hallazgos, meta = normalizar_excel(excel)

        assert len(hallazgos) == 1
        h = hallazgos[0]
        assert h.pano == "C1"
        assert h.rele == "SEL 351R"
        assert h.clasificacion == "Verificación DUF"
        assert h.calificacion == Calificacion.MENOR
        assert h.evidencia.fuente == FuenteEvidencia.EXCEL
        assert h.evidencia.fila_excel is not None
        assert h.evidencia.es_multi_pano_origen is False

    def test_multiples_hallazgos(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo 1", "Acción 1", "Menor"],
            ["C1", "SEL 351R", "7253", "Verificación ECAP",
             "Hallazgo 2", "Acción 2", "Relevante"],
            ["C2", "SEL 351R", "7254", "Verificación DUF",
             "Hallazgo 3", "Acción 3", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert len(hallazgos) == 3
        assert hallazgos[1].calificacion == Calificacion.RELEVANTE
        assert hallazgos[2].pano == "C2"

    def test_metadatos_extraidos(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        _, meta = normalizar_excel(excel)
        assert meta.coordinado == "COORDINADO A"
        assert meta.subestacion == "Iota"
        assert meta.informe_id == "RELE-MAR24-PPP-000"


class TestDesdoblamiento:
    def test_multi_pano_coma(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["BT1, ET1, E1", "ABB RET670", "13390", "Verificación ECAP",
             "Hallazgo transversal", "Acción transversal", "Relevante"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert len(hallazgos) == 3
        assert {h.pano for h in hallazgos} == {"BT1", "ET1", "E1"}
        for h in hallazgos:
            assert h.evidencia.es_multi_pano_origen is True
            assert h.descripcion == "Hallazgo transversal"

    def test_multi_pano_punto_y_coma(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["BT1; ET1", "ABB RET670", "13390", "Verificación ECAP",
             "Hallazgo transversal", "Acción transversal", "Relevante"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert {h.pano for h in hallazgos} == {"BT1", "ET1"}

    def test_pano_con_barra_se_desdobla(self, tmp_path: Path):
        # "BT1/ET1" se desdobla igual que en la Etapa 3 (decisión 2026-07-07).
        excel = _crear_excel_basico(tmp_path, [
            ["BT1/ET1", "ABB RET670", "13390", "Verificación ECAP",
             "Hallazgo", "Acción", "Relevante"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert {h.pano for h in hallazgos} == {"BT1", "ET1"}
        assert all(h.evidencia.es_multi_pano_origen for h in hallazgos)

    def test_guion_entre_codigos_se_desdobla(self, tmp_path: Path):
        # Zeta / Gamma: diferencial "BT1-CT1", igual que la Etapa 3.
        excel = _crear_excel_basico(tmp_path, [
            ["BT1-CT1", "SEL 387", "1", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert [h.pano for h in hallazgos] == ["BT1", "CT1"]
        assert all(h.evidencia.es_multi_pano_origen for h in hallazgos)

    @pytest.mark.parametrize("panos_raw", ["BT1 – CT1", "BT1—CT1"])
    def test_guion_largo_entre_codigos_se_desdobla(self, tmp_path: Path, panos_raw: str):
        excel = _crear_excel_basico(tmp_path, [
            [panos_raw, "SEL 387", "1", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert [h.pano for h in hallazgos] == ["BT1", "CT1"]

    def test_guion_y_coma_combinados(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["BT1-CT1, ET1", "SEL 387", "1", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert [h.pano for h in hallazgos] == ["BT1", "CT1", "ET1"]

    def test_separadores_mezclados_se_parten_todos(self, tmp_path: Path):
        # Antes solo se cortaba por el primer separador encontrado: "BT1/CT1, ET1"
        # dejaba el paño "BT1/CT1" entero y su tupla no calzaba nunca.
        excel = _crear_excel_basico(tmp_path, [
            ["BT1/CT1, ET1; E1", "SEL 387", "1", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert [h.pano for h in hallazgos] == ["BT1", "CT1", "ET1", "E1"]

    @pytest.mark.parametrize("compuesto", [
        "BT1-CT1", "BT1 – CT1", "A1-B1-C1", "E-1", "DT1 – P.NEUTRO", "CT1A-CT1B", "1-2",
    ])
    def test_etapa3_desdobla_igual_que_etapa2(self, compuesto: str):
        # Las tuplas (paño, relé) de ambos lados deben calzar: misma regla.
        from comparador.etapa2_normalizar_excel import split_guion_entre_codigos
        from comparador.etapa3_extraer_hallazgos import _split_pano_compuesto

        assert _split_pano_compuesto(compuesto) == split_guion_entre_codigos(compuesto)

    @pytest.mark.parametrize("pano", ["DT1 – P.NEUTRO", "E-1"])
    def test_guion_sin_dos_codigos_no_se_parte(self, tmp_path: Path, pano: str):
        # Eta: "DT1 – P.NEUTRO" es un solo paño.
        excel = _crear_excel_basico(tmp_path, [
            [pano, "GE 12IAC53A5A", "1", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert [h.pano for h in hallazgos] == [pano]
        assert hallazgos[0].evidencia.es_multi_pano_origen is False

    def test_pano_unico_no_desdobla(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert len(hallazgos) == 1
        assert hallazgos[0].evidencia.es_multi_pano_origen is False


class TestNormalizacionCalificacion:
    @pytest.mark.parametrize("variante,esperado", [
        ("Menor", Calificacion.MENOR),
        ("menor", Calificacion.MENOR),
        ("Relevante", Calificacion.RELEVANTE),
        ("relevante", Calificacion.RELEVANTE),
        ("Crítico", Calificacion.CRITICO),
        ("critico", Calificacion.CRITICO),
        ("Criticidad menor", Calificacion.MENOR),
        ("Criticidad relevante", Calificacion.RELEVANTE),
        ("Hallazgo menor", Calificacion.MENOR),
        ("Hallazgo relevante", Calificacion.RELEVANTE),
    ])
    def test_variantes_validas(self, tmp_path: Path, variante: str, esperado: Calificacion):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo", "Acción", variante],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert hallazgos[0].calificacion == esperado

    def test_calificacion_invalida_falla(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo", "Acción", "Desconocido"],
        ])

        with pytest.raises(ExcelNormalizationError, match="Calificación no reconocida"):
            normalizar_excel(excel)


class TestVariantesRealesDeFormato:
    """Variantes observadas en fotos de Excels reales (2026-07-07)."""

    def test_columna_pano_sin_ese(self, tmp_path: Path):
        # El Excel de Ensaut/COORDINADO B usa "Paño", no "Paño(s)".
        excel = _crear_excel_basico(
            tmp_path,
            [["BT1/ET1", "ABB RET670", "13390", "Verificación ECAP",
              "Hallazgo", "Acción", "Relevante"]],
            headers=[
                "Paño", "Relé", "ID Infotécnica (paño/sist.de Prot)",
                "Clasificación del Hallazgo", "Descripción del Hallazgo",
                "Acción Correctiva del Hallazgo", "Calificación",
            ],
        )

        hallazgos, _ = normalizar_excel(excel)
        assert {h.pano for h in hallazgos} == {"BT1", "ET1"}

    def test_metadatos_estilo_ensaut(self, tmp_path: Path):
        # Etiqueta "Fecha" genérica y sin fila de versión.
        excel = _crear_excel_basico(
            tmp_path,
            [["E1", "SEL 351S", "13393", "Verificación DUF",
              "Hallazgo", "Acción", "Menor"]],
            metadatos={
                "Fecha": "Monday, October 14, 2024",
                "Coordinado": "COORDINADO C -GRUPO COORDINADO B",
                "Subestación": "Theta",
                "Informe de Auditoría": "EN.000000.0_Informe Auditoría Técnica SE Theta_Rev.3",
            },
        )

        _, meta = normalizar_excel(excel)
        assert meta.coordinado == "COORDINADO C -GRUPO COORDINADO B"
        assert meta.fecha_auditoria == "Monday, October 14, 2024"
        assert meta.informe_id.startswith("EN.000000.0")

    def test_pano_con_su_rele_entre_parentesis(self, tmp_path: Path):
        # Excel Alfa: "ET2 (SEL 387E)" con relé "SEL 387E" es el paño ET2.
        excel = _crear_excel_basico(tmp_path, [
            ["ET2 (SEL 387E)", "SEL 387E", "1", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
            ["HT2 (SEL-387E)", "SEL 387E", "1", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert [h.pano for h in hallazgos] == ["ET2", "HT2"]
        assert not any(h.evidencia.es_multi_pano_origen for h in hallazgos)

    def test_parentesis_distinto_al_rele_no_se_toca(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["ET2 (otra cosa)", "SEL 387E", "1", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert hallazgos[0].pano == "ET2 (otra cosa)"

    def test_parentesis_se_quita_en_cada_pano_del_split(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["ET2 (SEL 387E), HT2 (SEL 387E)", "SEL 387E", "1", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert [h.pano for h in hallazgos] == ["ET2", "HT2"]

    def test_etiqueta_fecha_informe_no_pisa_fecha_auditoria(self, tmp_path: Path):
        # "Fecha de Informe de Auditoría y Tabla de Hallazgos" (Excel COORDINADO A) no
        # debe confundirse con la fecha de auditoría.
        excel = _crear_excel_basico(
            tmp_path,
            [["C1", "SEL 351R", "7253", "Verificación DUF",
              "Hallazgo", "Acción", "Menor"]],
            metadatos={
                "Coordinado": "COORDINADO A",
                "Subestación": "Iota",
                "Fecha de Auditoría": "12-03-2024 a 22-03-2024",
                "Informe de Auditoría": "RELE-MAR24-PPP-000",
                "Fecha de Informe de Auditoría y Tabla de Hallazgos": "Wednesday, April 29, 2026",
            },
        )

        _, meta = normalizar_excel(excel)
        assert meta.fecha_auditoria == "12-03-2024 a 22-03-2024"


class TestFormatoRealArchivos:
    """Estructuras observadas en las Tablas de Hallazgos reales (2026-07-07)."""

    def test_multi_hoja_usa_la_de_hallazgos(self, tmp_path: Path):
        # Los archivos reales traen hoja "Definiciones" primero.
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])
        wb = load_workbook(excel)
        original = wb.worksheets[0]
        original.title = "Tabla de Hallazgos"
        defs = wb.create_sheet("Definiciones", 0)  # queda como primera hoja
        defs.cell(row=1, column=1, value="Glosario: Descripción del Hallazgo")
        wb.save(excel)

        hallazgos, _ = normalizar_excel(excel)
        assert len(hallazgos) == 1

    def test_metadatos_con_celdas_intermedias_vacias(self, tmp_path: Path):
        # Real Reliable: label en col B, valor en col F (merged de por medio).
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.cell(row=1, column=2, value="Coordinado")
        ws.cell(row=1, column=6, value="COORDINADO A.")
        ws.cell(row=2, column=2, value="Subestación")
        ws.cell(row=2, column=6, value="S/E Añañuca")
        headers = ["Paño(s)", "Relé", "ID", "Clasificación del Hallazgo",
                   "Descripción del Hallazgo", "Acción Correctiva del Hallazgo",
                   "Calificación"]
        for j, hdr in enumerate(headers, 1):
            ws.cell(row=4, column=j, value=hdr)
        fila = ["C1", "SEL 351R", "1", "Verificación DUF", "H", "A", "Menor"]
        for j, v in enumerate(fila, 1):
            ws.cell(row=5, column=j, value=v)
        ruta = tmp_path / "meta_offset.xlsx"
        wb.save(ruta)

        _, meta = normalizar_excel(ruta)
        assert meta.coordinado == "COORDINADO A."
        assert meta.subestacion == "S/E Añañuca"


class TestValidacionEstructural:
    def test_columna_faltante_falla(self, tmp_path: Path):
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.cell(row=1, column=1, value="Relé")
        ws.cell(row=1, column=2, value="Calificación")
        filepath = tmp_path / "incompleto.xlsx"
        wb.save(filepath)

        with pytest.raises(ExcelNormalizationError, match="Columnas faltantes"):
            normalizar_excel(filepath)

    def test_filas_vacias_ignoradas(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
            ["", "", "", "", "", "", ""],
            ["C2", "SEL 351R", "7254", "Verificación ECAP",
             "Hallazgo 2", "Acción 2", "Relevante"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert len(hallazgos) == 2

    def test_fila_parcial_falla_ruidosamente(self, tmp_path: Path):
        # Una fila con descripción pero sin relé no puede descartarse en silencio.
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "", "7253", "Verificación DUF",
             "Hallazgo sin relé", "Acción", "Menor"],
        ])

        with pytest.raises(ExcelNormalizationError, match="obligatorios vacíos"):
            normalizar_excel(excel)

    def test_accion_correctiva_vacia_es_valida(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo", "", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert len(hallazgos) == 1
        assert hallazgos[0].accion_correctiva == ""


class TestTrazabilidad:
    def test_fila_excel_preservada(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo 1", "Acción 1", "Menor"],
            ["C2", "SEL 351R", "7254", "Verificación ECAP",
             "Hallazgo 2", "Acción 2", "Relevante"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert hallazgos[0].evidencia.fila_excel is not None
        assert hallazgos[1].evidencia.fila_excel is not None
        assert hallazgos[0].evidencia.fila_excel != hallazgos[1].evidencia.fila_excel

    def test_texto_crudo_contiene_descripcion(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Relaciones de transformación desactualizadas",
             "Actualizar DUF", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)
        assert "Relaciones de transformación desactualizadas" in hallazgos[0].evidencia.texto_crudo


class TestFilasTransversales:
    def test_fila_transversal_marcada(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo normal", "Acción", "Menor"],
            ["C1", "Análisis de coordinación", "", "Coord. de Protecciones",
             "Estudio transversal de coordinación", "Revisar", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(
            excel, patrones_transversales=["análisis de coordinación"]
        )

        assert [h.es_transversal for h in hallazgos] == [False, True]

    def test_transversal_insensible_a_mayusculas_y_acentos(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "ANALISIS DE COORDINACION", "", "Coord. de Protecciones",
             "Sin acentos y en mayúsculas", "Revisar", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(
            excel, patrones_transversales=["análisis de coordinación"]
        )

        assert hallazgos[0].es_transversal is True

    def test_transversal_no_desdobla_multi_pano(self, tmp_path: Path):
        # Una fila transversal = una entrada de revisión manual, aunque su
        # celda de paños liste varios.
        excel = _crear_excel_basico(tmp_path, [
            ["BT1, ET1, E1", "Análisis de coordinación", "", "Coord. de Protecciones",
             "Estudio transversal", "Revisar", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(
            excel, patrones_transversales=["análisis de coordinación"]
        )

        assert len(hallazgos) == 1
        assert hallazgos[0].pano == "BT1, ET1, E1"
        assert hallazgos[0].evidencia.es_multi_pano_origen is False

    def test_sin_patrones_nada_es_transversal(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "Análisis de coordinación", "", "Coord. de Protecciones",
             "Estudio transversal", "Revisar", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)

        assert hallazgos[0].es_transversal is False

    def test_patron_no_marca_rele_normal(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo normal", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(
            excel, patrones_transversales=["análisis de coordinación"]
        )

        assert hallazgos[0].es_transversal is False

    @pytest.mark.parametrize("rele", ["-", "—", "–", "/", "..."])
    def test_rele_sin_alfanumericos_es_transversal_sin_patrones(
        self, tmp_path: Path, rele: str
    ):
        # Beta filas 221-222 / Eta 51-52: relé "-", para cualquier empresa.
        excel = _crear_excel_basico(tmp_path, [
            ["BT2, ET2", rele, "", "Ver. Tiempo Máx. Actuación",
             "Tiempo teórico de despeje", "Revisar", "Relevante"],
        ])

        hallazgos, _ = normalizar_excel(excel)

        assert len(hallazgos) == 1
        assert hallazgos[0].es_transversal is True
        assert hallazgos[0].pano == "BT2, ET2"

    def test_rele_con_guion_no_es_transversal(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL-787", "", "Verificación DUF",
             "Hallazgo normal", "Acción", "Menor"],
        ])

        hallazgos, _ = normalizar_excel(excel)

        assert hallazgos[0].es_transversal is False


class TestColumnasAusentes:
    def test_columnas_opcionales_ausentes_quedan_registradas(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF", "Hallazgo"],
        ], headers=[
            "Paño(s)", "Relé", "ID Infotécnica",
            "Clasificación del Hallazgo", "Descripción del Hallazgo",
        ])

        _, meta = normalizar_excel(excel)

        assert meta.columnas_ausentes == [
            "Acción Correctiva del Hallazgo", "Calificación",
        ]

    def test_excel_completo_no_registra_ausencias(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            ["C1", "SEL 351R", "7253", "Verificación DUF",
             "Hallazgo", "Acción", "Menor"],
        ])

        _, meta = normalizar_excel(excel)

        assert meta.columnas_ausentes == []


def _editar(excel: Path, editar) -> Path:
    """Abre el Excel sintético, aplica `editar(wb, ws)` y lo guarda."""
    wb = load_workbook(excel)
    ws = wb.active
    assert ws is not None
    editar(wb, ws)
    wb.save(excel)
    return excel


FILA = ["C1", "SEL 351R", "7253", "Verificación DUF", "Hallazgo", "Acción", "Menor"]


class TestReferenciasExcel:
    """La referencia que ve la revisora (hoja + fila) apunta a la celda real
    (prioridad del usuario 2026-10-18: referencias confiables)."""

    def test_fila_excel_es_la_fila_que_se_ve_en_la_hoja(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            [*FILA[:4], "Hallazgo A", *FILA[5:]],
            [None] * 7,                      # fila vacía en medio de la tabla
            [*FILA[:4], "Hallazgo B", *FILA[5:]],
        ])
        ws = load_workbook(excel).active
        assert ws is not None
        fila_de = {ws.cell(f, 5).value: f for f in range(1, ws.max_row + 1)}

        hallazgos, _ = normalizar_excel(excel)

        assert [h.evidencia.fila_excel for h in hallazgos] == [
            fila_de["Hallazgo A"], fila_de["Hallazgo B"]
        ]

    def test_registra_la_hoja(self, tmp_path: Path):
        def con_definiciones(wb, ws):
            ws.title = "Tabla de Hallazgos"
            wb.create_sheet("Definiciones", 0)

        excel = _editar(_crear_excel_basico(tmp_path, [FILA]), con_definiciones)

        hallazgos, _ = normalizar_excel(excel)
        assert hallazgos[0].evidencia.hoja_excel == "Tabla de Hallazgos"

    def test_dos_hojas_de_hallazgos_fallan(self, tmp_path: Path):
        def dos_hojas(wb, ws):
            ws.title = "Tabla de Hallazgos"
            wb.copy_worksheet(ws).title = "Tabla de Hallazgos (2)"

        excel = _editar(_crear_excel_basico(tmp_path, [FILA]), dos_hojas)

        with pytest.raises(ExcelNormalizationError, match="2 hojas de hallazgos"):
            normalizar_excel(excel)

    def test_pano_y_rele_combinados_se_heredan(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            [*FILA[:4], "Hallazgo A", *FILA[5:]],
            [None, None, *FILA[2:4], "Hallazgo B", *FILA[5:]],
        ])
        ws0 = load_workbook(excel).active
        assert ws0 is not None
        primera = ws0.max_row - 1

        def combinar(wb, ws):
            ws.merge_cells(start_row=primera, end_row=primera + 1, start_column=1, end_column=1)
            ws.merge_cells(start_row=primera, end_row=primera + 1, start_column=2, end_column=2)

        hallazgos, _ = normalizar_excel(_editar(excel, combinar))
        assert [(h.pano, h.rele) for h in hallazgos] == [("C1", "SEL 351R"), ("C1", "SEL 351R")]

    def test_rele_vacio_sin_combinar_falla_en_vez_de_heredar(self, tmp_path: Path):
        # Antes un ffill le daba el relé de la fila anterior en silencio.
        excel = _crear_excel_basico(tmp_path, [
            [*FILA[:4], "Hallazgo A", *FILA[5:]],
            ["C1", None, *FILA[2:4], "Hallazgo B", *FILA[5:]],
        ])

        with pytest.raises(ExcelNormalizationError, match="combinada con ella en vertical"):
            normalizar_excel(excel)

    def test_fila_oculta_se_marca(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [
            [*FILA[:4], "Visible", *FILA[5:]],
            [*FILA[:4], "Oculta", *FILA[5:]],
        ])
        ws0 = load_workbook(excel).active
        assert ws0 is not None
        ultima = ws0.max_row

        def ocultar(wb, ws):
            ws.row_dimensions[ultima].hidden = True

        hallazgos, _ = normalizar_excel(_editar(excel, ocultar))
        assert {h.descripcion: h.evidencia.fila_oculta for h in hallazgos} == {
            "Visible": False, "Oculta": True,
        }

    def test_combinada_horizontal_no_copia_el_pano_al_rele(self, tmp_path: Path):
        # Revisión 2026-10-18: Paño+Relé combinados en una fila no inventan el relé.
        excel = _crear_excel_basico(tmp_path, [
            [*FILA[:4], "Hallazgo A", *FILA[5:]],
            ["E2", None, *FILA[2:4], "Hallazgo B", *FILA[5:]],
        ])
        ws0 = load_workbook(excel).active
        assert ws0 is not None
        fila = ws0.max_row

        def combinar(wb, ws):
            ws.merge_cells(start_row=fila, end_row=fila, start_column=1, end_column=2)

        with pytest.raises(ExcelNormalizationError, match="Relé"):
            normalizar_excel(_editar(excel, combinar))

    def test_contenido_oculto_bajo_una_combinada_falla(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        # pandas lee lo guardado bajo la combinada; la revisora ve el valor del ancla.
        excel = _crear_excel_basico(tmp_path, [
            [*FILA[:4], "Hallazgo A", *FILA[5:]],
            ["C1", None, *FILA[2:4], "Hallazgo B", *FILA[5:]],
        ])
        ws0 = load_workbook(excel).active
        assert ws0 is not None
        primera = ws0.max_row - 1

        def combinar(wb, ws):
            ws.merge_cells(start_row=primera, end_row=primera + 1, start_column=2, end_column=2)

        ruta = _editar(excel, combinar)
        import comparador.etapa2_normalizar_excel as e2

        real = e2.pd.read_excel

        def leer_con_oculto(*a, **k):  # un archivo que guardó contenido bajo la combinada
            df = real(*a, **k)
            if "Relé" in df.columns:
                df.loc[df["Descripción del Hallazgo"] == "Hallazgo B", "Relé"] = "SEL 751"
            return df

        monkeypatch.setattr(e2.pd, "read_excel", leer_con_oculto)
        with pytest.raises(ExcelNormalizationError, match="oculto"):
            normalizar_excel(ruta)

    def test_descripcion_na_da_el_error_de_campo_vacio(self, tmp_path: Path):
        excel = _crear_excel_basico(tmp_path, [[*FILA[:4], "N/A", *FILA[5:]]])
        with pytest.raises(ExcelNormalizationError, match="campos obligatorios vacíos"):
            normalizar_excel(excel)
