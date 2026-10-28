"""Integración de punta a punta del pipeline, sin APIs.

Corre `ejecutar_pipeline` con el código real de todas las etapas (checklist,
E1 con su post-proceso, E2 sobre un Excel en disco, E3 con recorte,
verificación de citas, reparación y reintento, E4 con unicidad, juez de
matching, polaridad y juez de redacción, E5 y persistencia). Solo se
reemplazan las fronteras con proveedores externos: `llamar_llm_estructurado`
de cada etapa, `calcular_embeddings` y `contar_tokens_pdf` (count_tokens). Un
guardia hace fallar el test si algo intenta crear un cliente real.

Los insumos se generan en tmp_path: un PDF con capa de texto (pymupdf) y un
Excel con el formato SEC (openpyxl). Los runs van a un runs_dir temporal.
"""

from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import json
import math
import re
import shutil
import sys
from collections.abc import Callable
from pathlib import Path
from typing import Any

import fitz
import numpy as np
import pytest
from fastapi.testclient import TestClient
from numpy.typing import NDArray
from openpyxl import Workbook
from pydantic import BaseModel
from pypdf import PdfReader

from comparador import etapa0_checklist as e0
from comparador import etapa1_descubrir_panos as e1
from comparador import etapa3_extraer_hallazgos as e3
from comparador import etapa4_matching as e4
from comparador import evaluar_lote, troceo
from comparador.config import Settings, settings
from comparador.etapa3_extraer_hallazgos import ExtraccionInvalidaError
from comparador.etapa4_matching import JuezInvalidoError
from comparador.exportar import exportar_csv, exportar_feedback_md, exportar_xlsx
from comparador.llm import client as llm_client
from comparador.llm import embeddings as llm_embeddings
from comparador.pipeline import ejecutar_pipeline
from comparador.schemas import (
    CampoDiferente,
    EmpresaAuditora,
    EstadoHallazgo,
    FeedbackRun,
    MatchResult,
    ReporteInforme,
    ResultadoChecklist,
    TipoMatch,
)

# --- Textos del escenario (sintéticos) ---

D_DUF = "No se encuentran actualizadas las relaciones de transformación en el DUF."
A_DUF = "Actualizar las relaciones de transformación del DUF."
SINCRO_PDF = (
    "Sincronización horaria externa. No cumple, ya que se incumple con el artículo "
    "3-24 de la NTSyCS."
)
SINCRO_EXCEL = (
    "Equipo no posee sincronización horaria vía GPS, por lo que incumple con el "
    "artículo 3-24 de la NTSyCS."
)
A_SINCRO = "Gestionar la sincronización horaria del equipo."
E7_PDF = "Comparación con ECAP adjunto en INFOTÉCNICA: Cumple"
E7_EXCEL = "Comparación con ECAP adjunto en INFOTÉCNICA: No cumple"
A_ECAP = "Actualizar el ECAP en Infotécnica."
AJ_PDF = "Los ajustes de la función 51 difieren del estudio de coordinación vigente."
AJ_EXCEL = "No se encuentra el registro de eventos del relé."
ACC_AJ_PDF = "Ajustar la función 51 según el estudio vigente."
ACC_AJ_EXCEL = "Habilitar el registro de eventos."
P1 = "La función 50BF no está habilitada."
A_P1 = "Habilitar la función 50BF."
P2 = "La función 79 tiene un tiempo muerto distinto al del estudio."
A_P2 = "Ajustar el tiempo muerto de la función 79."
D_387 = "La pendiente de la función 87T no coincide con el estudio."
A_387 = "Ajustar la pendiente de la función 87T."
Q_PDF = "No se realizó la prueba de inyección secundaria de la función 67N."
A_Q = "Realizar la prueba de inyección secundaria pendiente."
R5_EXCEL = "El DUF no muestra el interruptor 52C3."
A_R5 = "Actualizar el DUF con el interruptor 52C3."
TRANSV = "No se adjunta el estudio de coordinación de protecciones de la subestación."
A_TRANSV = "Adjuntar el estudio de coordinación."

# Checklist con cita (settings.cita_checklist): el texto que "copia el modelo".
PORTADA = "INFORME DE AUDITORÍA TÉCNICA DE PROTECCIONES"
ANEXO = "registro de pruebas de inyección secundaria del paño"
CITA_INVENTADA = "Todos los relés cuentan con sincronización horaria GPS vigente"

# Clasificaciones (lista SEC).
C_DUF = "Verificación DUF"
C_SINCRO = "Ver. Acceso, Registro y Sincr."
C_ECAP = "Verificación ECAP"
C_COORD = "Ver. Coord. de Protecciones"
C_OPER = "Prueba Oper. de Protección"
C_87T = "Prueba Estabilidad 87T"
C_PEND = "Prueba Pendiente"

# Pares con similitud fijada; el resto de los textos es ortogonal entre sí.
PARES_SIMILITUD = {(SINCRO_PDF, SINCRO_EXCEL): 0.80, (E7_PDF, E7_EXCEL): 0.97}
# Lo que el juez de matching empareja (descripción PDF → descripción Excel).
JUEZ_MATCHING = {P1: P1, P2: P2}
# Lo que el juez de redacción declara mismo contenido.
EQUIVALENTES = {(SINCRO_PDF, SINCRO_EXCEL)}

# Tokens falsos por llamada, distintos por etapa para detectar sumas cruzadas.
TOKENS = {
    "checklist": (900, 100),
    "e1": (150, 50),
    "e3": (25, 5),
    "juez": (5, 2),
    "redaccion": (40, 10),
}
TOKENS_EMBEDDINGS = 11

# Páginas físicas del PDF sintético: portada, C1 (2 págs), BT1-CT1, C3.
PAGINAS = [
    "INFORME DE AUDITORÍA TÉCNICA DE PROTECCIONES\nSubestación Prueba",
    f"Tabla Resumen Paño C1 - Relé SEL 751\n{D_DUF}\n{A_DUF}\n{SINCRO_PDF}",
    f"Tabla Resumen Paño C1 (continuación)\n{E7_PDF}",
    "Relé SEL 551C, Serial: 1, Paño BT1\nRelé SEL 387, Serial: 2, Paños BT1 y CT1\n"
    f"Tabla Resumen Paño BT1-CT1\n{P1}\n{P2}\n{D_387}",
    f"Tabla Resumen Paño C3 - Relé SEL 451\n{AJ_PDF}\n{Q_PDF}",
]

# Filas del Excel: paño(s), relé, clasificación, descripción, acción, calificación.
FILAS_EXCEL = [
    ["C1", "SEL 751", C_DUF, D_DUF, A_DUF, "Menor"],
    ["C1", "SEL 751", C_SINCRO, SINCRO_EXCEL, A_SINCRO, "Relevante"],
    ["C1", "SEL 751", C_ECAP, E7_EXCEL, A_ECAP, "Menor"],
    ["BT1", "SEL 551C", C_OPER, P1, A_P1, "Menor"],
    ["BT1", "SEL 551C", C_OPER, P2, A_P2, "Relevante"],
    ["BT1-CT1", "SEL 387", C_87T, D_387, A_387, "Menor"],
    ["C3", "SEL 451", C_COORD, AJ_EXCEL, ACC_AJ_EXCEL, "Menor"],
    ["C3", "SEL 451", C_DUF, R5_EXCEL, A_R5, "Menor"],
    ["BT1, C1", "Análisis de coordinación", C_COORD, TRANSV, A_TRANSV, "Menor"],
]


def _h(pano: str, rele: str, clasificacion: str, desc: str, accion: str, calificacion: str,
       pagina: int | None = 1) -> dict[str, Any]:
    return {
        "pano": pano, "rele": rele, "clasificacion": clasificacion, "descripcion": desc,
        "accion_correctiva": accion, "calificacion": calificacion,
        "evidencia": {"pagina_pdf": pagina, "texto_crudo": desc},
    }


# E3: respuestas por paño y por intento. El primer intento de BT1-CT1 trae un
# relé que no está en el mapa de la Etapa 1 → reintento con retroalimentación.
RESPUESTAS_E3: dict[str, list[list[dict[str, Any]]]] = {
    "C1": [[
        _h("C1", "SEL 751", C_DUF, D_DUF, A_DUF, "Menor"),
        _h("C1", "SEL 751", C_SINCRO, SINCRO_PDF, A_SINCRO, "Relevante"),
        # Cita la página 1 del recorte (pág. 2), pero el texto está en la 3:
        # la verificación de citas la corrige y deja la citada a la vista.
        _h("C1", "SEL 751", C_ECAP, E7_PDF, A_ECAP, "Menor"),
    ]],
    "BT1-CT1": [
        [_h("BT1", "SEL 999", C_OPER, P1, A_P1, "Menor")],
        [
            _h("BT1", "SEL 551C", C_OPER, P1, A_P1, "Menor"),
            _h("BT1", "SEL 551C", C_OPER, P2, A_P2, "Relevante"),
            _h("BT1-CT1", "SEL 387", C_87T, D_387, A_387, "Menor"),
        ],
    ],
    "C3": [[
        _h("C3", "SEL 451", C_COORD, AJ_PDF, ACC_AJ_PDF, "Relevante"),
        _h("C3", "SEL 451", C_PEND, Q_PDF, A_Q, "Menor"),
    ]],
}

# E1: lo que "devuelve el modelo", con un relé con el paño pegado y un
# sub-paño con prefijo: el post-proceso real de la Etapa 1 los repara.
RESPUESTA_E1 = {"panos": [
    {"nombre": "C1", "pagina_inicio": 2, "pagina_fin": 3, "reles": [
        {"nombre": "SEL 751", "evidencia_origen": "Tabla Resumen Paño C1", "sub_panos": []},
    ]},
    {"nombre": "BT1-CT1", "pagina_inicio": 4, "pagina_fin": 4, "reles": [
        {"nombre": "SEL 551C (Paño BT1)", "evidencia_origen": "Relé SEL 551C, Paño BT1",
         "sub_panos": []},
        {"nombre": "SEL 387", "evidencia_origen": "Relé SEL 387, Paños BT1 y CT1",
         "sub_panos": ["Paños BT1 y CT1"]},
    ]},
    {"nombre": "C3", "pagina_inicio": 5, "pagina_fin": 5, "reles": [
        {"nombre": "SEL 451", "evidencia_origen": "Tabla Resumen Paño C3", "sub_panos": []},
    ]},
]}


# --- Insumos en disco ---


def _crear_pdf(ruta: Path, paginas_extra: int = 0) -> Path:
    doc = fitz.open()
    for texto in [*PAGINAS, *(f"Anexo {i}: {ANEXO}" for i in range(1, paginas_extra + 1))]:
        pagina = doc.new_page()
        pagina.insert_textbox(fitz.Rect(40, 120, 560, 740), texto, fontsize=9)
    doc.save(ruta)
    return ruta


def _crear_excel(ruta: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tabla de Hallazgos"
    for fila, (clave, valor) in enumerate([
        ("Coordinado", "Coordinado de Prueba"),
        ("Subestación", "Subestación Prueba"),
        ("Informe de Auditoría", "AUD-PRUEBA-001"),
        ("Fecha de Auditoría", "01-01-2026"),
    ], start=1):
        ws.cell(row=fila, column=1, value=clave)
        ws.cell(row=fila, column=2, value=valor)
    headers = ["Paño(s)", "Relé", "ID Infotécnica", "Clasificación del Hallazgo",
               "Descripción del Hallazgo", "Acción Correctiva del Hallazgo", "Calificación"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=6, column=col, value=header)
    for n, (pano, rele, clasif, desc, accion, calif) in enumerate(FILAS_EXCEL, start=7):
        for col, valor in enumerate([pano, rele, "", clasif, desc, accion, calif], start=1):
            ws.cell(row=n, column=col, value=valor)
    wb.save(ruta)
    return ruta


def _sha256(ruta: Path) -> str:
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


# --- Fakes de las fronteras externas ---


def _meta(etapa: str, data: Any, kw: dict[str, Any]) -> dict[str, Any]:
    prompt, completion = TOKENS[etapa]
    return {
        "modelo": kw["modelo"],
        "tiempo_segundos": 0.01,
        "tokens_prompt": prompt,
        "tokens_completion": completion,
        "tokens_total": prompt + completion,
        "tokens_cache_lectura": 0,
        # Como llm.client: lo que se pidió de thinking/effort (la caché de jueces lo usa).
        "parametros_thinking": llm_client.parametros_thinking(kw["modelo"], settings.effort_llm),
        "system_prompt": kw["system_prompt"],
        "user_prompt_text": kw["user_prompt"],
        "pdf_incluido": kw.get("pdf_base64") is not None,
        "response_raw": data,
    }


class LLMFalso:
    """Un solo `llamar_llm_estructurado` falso para todas las etapas: despacha
    por el response_model que pide cada una y registra cada llamada."""

    def __init__(self, juez_redaccion_omite_ids: bool = False) -> None:
        self.llamadas: list[dict[str, Any]] = []
        self.intentos_e3: dict[str, int] = {}
        self.juez_redaccion_omite_ids = juez_redaccion_omite_ids

    def __call__(self, **kw: Any) -> tuple[BaseModel, dict[str, Any]]:
        self.llamadas.append(kw)
        modelo = kw["response_model"]
        etapa, data = {
            e0._ChecklistResponse: ("checklist", self._checklist),
            e0._ChecklistConCitaResponse: ("checklist", self._checklist),
            e1._ListaPanosResponse: ("e1", lambda _: RESPUESTA_E1),
            e3._ListaHallazgosResponse: ("e3", self._e3),
            e4._JuezResponse: ("juez", self._juez_matching),
            e4._JuezResponseJustificaPrimero: ("juez", self._juez_matching),
            e4._JuezRedaccionResponse: ("redaccion", self._juez_redaccion),
        }[modelo]
        respuesta = data(kw)
        return modelo.model_validate(respuesta), _meta(etapa, respuesta, kw)

    def de(self, response_model: type[BaseModel]) -> list[dict[str, Any]]:
        return [k for k in self.llamadas if k["response_model"] is response_model]

    @staticmethod
    def _checklist(kw: dict[str, Any]) -> dict[str, Any]:
        numeros = range(1, e0.TOTAL_PUNTOS + 1)
        con_cita = kw["response_model"] is e0._ChecklistConCitaResponse
        tramo = re.search(r"contiene las páginas (\d+) a (\d+)", kw["user_prompt"])
        if tramo is None:
            # Informe entero: todo cumple salvo el 5 (no cumple) y el 6. Con
            # cita, el 4 "cita" un texto que el PDF no tiene: baja a no verificable.
            resultado = {5: "no_cumple", 6: "no_verificable"}
            items = [
                {"numero": n, "resultado": resultado.get(n, "cumple"), "paginas": "1",
                 "referencia": "portada"}
                for n in numeros
            ]
            if con_cita:
                for item in items:
                    texto = CITA_INVENTADA if item["numero"] == 4 else PORTADA
                    item["cita"] = ({"pagina": 1, "texto": texto}
                                    if item["resultado"] == "cumple" else None)
            return {"items": items}
        # Por tramo: el 1 se ve en el primero; el 2 y el 3 en el segundo.
        # Las páginas son relativas al tramo (la etapa las pasa a absolutas).
        primero = tramo.group(1) == "1"
        por_tramo = {1: ("cumple", "1", PORTADA)} if primero else {
            2: ("cumple", "2", f"Anexo 2: {ANEXO}"), 3: ("no_cumple", "1", None),
        }
        items = []
        for n in numeros:
            resultado, paginas, cita = por_tramo.get(n, ("no_verificable", None, None))
            item: dict[str, Any] = {"numero": n, "resultado": resultado,
                                    "paginas": paginas or "", "referencia": None}
            if con_cita:
                item["cita"] = {"pagina": int(paginas), "texto": cita} if cita and paginas else None
            items.append(item)
        return {"items": items}

    def _e3(self, kw: dict[str, Any]) -> dict[str, Any]:
        m = re.search(r'Extrae todos los hallazgos del paño "([^"]+)"', kw["user_prompt"])
        assert m is not None
        pano = m.group(1)
        n = self.intentos_e3.get(pano, 0)
        self.intentos_e3[pano] = n + 1
        return {"hallazgos": RESPUESTAS_E3[pano][n]}

    @staticmethod
    def _juez_matching(kw: dict[str, Any]) -> dict[str, Any]:
        prompt = kw["user_prompt"]
        m = re.search(r'Hallazgo del PDF:\n  "(.*)"\n', prompt)
        assert m is not None
        candidatos = re.findall(r"^  \[(\d+)\] (.*)$", prompt, flags=re.MULTILINE)
        esperado = JUEZ_MATCHING.get(m.group(1))
        indice = next((int(i) for i, texto in candidatos if texto == esperado), None)
        return {
            "indice_candidato": indice,
            "confianza": 92 if indice is not None else 10,
            "justificacion": "mismo hallazgo" if indice is not None else "ningún candidato",
        }

    def _juez_redaccion(self, kw: dict[str, Any]) -> dict[str, Any]:
        items = json.loads(kw["user_prompt"].split("\n\n", 1)[1].rsplit("\n\n", 1)[0])
        veredictos = [
            {
                "id": it["id"],
                "justificacion": "mismo incumplimiento de sincronización"
                if (it["texto_pdf"], it["texto_excel"]) in EQUIVALENTES
                else "describen problemas distintos",
                "mismo_contenido": (it["texto_pdf"], it["texto_excel"]) in EQUIVALENTES,
            }
            for it in items
        ]
        return {"veredictos": veredictos[:-1] if self.juez_redaccion_omite_ids else veredictos}


def embeddings_falsos(textos: list[str]) -> tuple[NDArray[np.float64], dict[str, Any]]:
    """Vectores con la similitud de PARES_SIMILITUD; el resto, ortogonales."""
    base = 2 * len(PARES_SIMILITUD)
    dim = base + len(textos)
    fijos: dict[str, NDArray[np.float64]] = {}
    for n, ((a, b), sim) in enumerate(PARES_SIMILITUD.items()):
        va = np.zeros(dim)
        va[2 * n] = 1.0
        vb = np.zeros(dim)
        vb[2 * n], vb[2 * n + 1] = sim, math.sqrt(1 - sim * sim)
        fijos[a], fijos[b] = va, vb
    matriz = np.zeros((len(textos), dim))
    for i, texto in enumerate(textos):
        if texto in fijos:
            matriz[i] = fijos[texto]
        else:
            matriz[i, base + i] = 1.0
    return matriz, {"etapa": "embeddings", "modelo": "falso", "tokens_total": TOKENS_EMBEDDINGS}


def _paginas_del_pdf(pdf_base64: str) -> int:
    return len(PdfReader(io.BytesIO(base64.b64decode(pdf_base64))).pages)


class ContadorTokensFalso:
    """count_tokens falso: tokens proporcionales a las páginas del PDF."""

    def __init__(self, por_pagina: int) -> None:
        self.por_pagina = por_pagina
        self.llamadas = 0

    def __call__(self, *, modelo: str, system_prompt: str, user_prompt: str,
                 pdf_base64: str) -> int:
        self.llamadas += 1
        return _paginas_del_pdf(pdf_base64) * self.por_pagina


def _cliente_prohibido() -> Any:
    raise AssertionError("El test intentó crear un cliente real de un proveedor LLM")


def _solo_la_prueba_de_la_clave(textos: list[str]) -> tuple[NDArray[np.float64], dict[str, Any]]:
    """Embeddings que solo responden la prueba de la clave de OpenAI (PD66): E4
    debe usar los vectores guardados, sin pagar ninguno."""
    if textos != [llm_embeddings.PALABRA_PRUEBA]:
        raise AssertionError(f"E4 pidió {len(textos)} embeddings: debía usar los guardados")
    return embeddings_falsos(textos)


@pytest.fixture
def entorno(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> dict[str, Any]:
    """runs_dir temporal, credenciales de mentira, fakes en cada frontera y
    guardia contra clientes reales."""
    runs = tmp_path / "runs"
    monkeypatch.setattr(settings, "runs_dir", runs)
    monkeypatch.setattr(settings, "openai_api_key", "sk-test")
    monkeypatch.setattr(settings, "anthropic_api_key", "sk-ant-test")
    monkeypatch.setattr(settings, "modelo_extraccion", "claude-test")
    monkeypatch.setattr(settings, "modelo_extraccion_fallback", "")
    monkeypatch.setattr(settings, "modelo_juez_matching", "")
    monkeypatch.setattr(settings, "juez_redaccion", True)
    monkeypatch.setattr(settings, "threshold_descripcion", 0.85)
    monkeypatch.setattr(settings, "threshold_accion", 0.85)
    monkeypatch.setattr(settings, "top_k_candidatos", 5)
    monkeypatch.setattr(settings, "umbral_tokens_troceo", 900_000)
    monkeypatch.setattr(settings, "tokens_por_tramo", 550_000)
    monkeypatch.setattr(llm_client, "get_client", _cliente_prohibido)
    monkeypatch.setattr(llm_client, "get_client_anthropic", _cliente_prohibido)
    # embeddings.py importó get_client por nombre: su referencia es otra.
    monkeypatch.setattr(llm_embeddings, "get_client", _cliente_prohibido)

    llm = LLMFalso()
    for modulo in (e0, e1, e3, e4):
        monkeypatch.setattr(modulo, "llamar_llm_estructurado", llm)
    monkeypatch.setattr(e4, "calcular_embeddings", embeddings_falsos)
    contador = ContadorTokensFalso(por_pagina=1_000)
    monkeypatch.setattr(troceo, "contar_tokens_pdf", contador)

    insumos = tmp_path / "insumos"
    insumos.mkdir()
    return {"runs": runs, "llm": llm, "contador": contador, "insumos": insumos}


def _correr(entorno: dict[str, Any], paginas_extra: int = 0) -> tuple[ReporteInforme, Path,
                                                                       list[str]]:
    pdf = _crear_pdf(entorno["insumos"] / "informe.pdf", paginas_extra)
    excel = _crear_excel(entorno["insumos"] / "hallazgos.xlsx")
    etapas: list[str] = []
    reporte, run_dir = asyncio.run(ejecutar_pipeline(
        pdf, excel, EmpresaAuditora.ENSAUT, on_progreso=etapas.append,
    ))
    return reporte, run_dir, etapas


def _json(ruta: Path) -> Any:
    return json.loads(ruta.read_text(encoding="utf-8"))


def _nombres(carpeta: Path) -> list[str]:
    return sorted(p.name for p in carpeta.iterdir())


def _tokens_raw(carpeta: Path) -> int:
    return sum(_json(p).get("tokens_total", 0) for p in carpeta.glob("*.json"))


def _por_descripcion(reporte: ReporteInforme, descripcion: str) -> MatchResult:
    for r in reporte.resultados:
        h = r.hallazgo_pdf or r.hallazgo_excel
        if h is not None and h.descripcion == descripcion:
            return r
    raise AssertionError(f"sin resultado para {descripcion!r}")


class TestPipelinePuntaAPunta:
    def test_run_completo_persiste_cada_etapa(self, entorno: dict[str, Any]) -> None:
        reporte, run_dir, etapas = _correr(entorno)

        # El run vive en el runs_dir temporal, nunca en runs/ del repo.
        assert run_dir.parent == entorno["runs"]
        # El Excel (sin LLM) va primero: un Excel inválido falla antes de gastar.
        assert etapas == ["normalizar_excel", "verificaciones_paralelas", "descubrir_panos",
                          "extraccion_pdf", "matching", "consolidar"]
        meta_run = _json(run_dir / "metadata.json")
        assert set(meta_run["relojes_etapas"]) == {"e2", "checklist", "e1", "e3", "e4", "e5"}
        assert all(r["segundos"] >= 0 for r in meta_run["relojes_etapas"].values())
        version = meta_run["version_codigo"]
        assert len(version["sha256_codigo"]) == 16 and version["archivos_codigo"] > 0
        assert len(version["sha256_checklist_puntos.yaml"]) == 16

        # Invariante 20: inputs hasheados, JSON de cada etapa y raw de cada llamada.
        assert (run_dir / "input_pdf.hash").read_text() == _sha256(
            entorno["insumos"] / "informe.pdf")
        assert (run_dir / "input_excel.hash").read_text() == _sha256(
            entorno["insumos"] / "hallazgos.xlsx")
        for nombre in ("checklist.json", "etapa1_panos.json", "etapa2_excel_normalizado.json",
                       "etapa3_pdf_hallazgos.json", "etapa4_matches.json",
                       "reporte_final.json", "metadata.json"):
            assert (run_dir / nombre).is_file(), nombre
        assert not (run_dir / "troceo.json").exists()
        assert not (run_dir / "adecuacion.json").exists()
        assert _nombres(run_dir / "checklist_raw") == ["verificar_checklist.json"]
        assert _nombres(run_dir / "etapa1_raw") == ["descubrir_panos.json"]
        assert _nombres(run_dir / "etapa3_raw") == [
            "extraer_BT1-CT1.json", "extraer_BT1-CT1_intento1.json",
            "extraer_C1.json", "extraer_C3.json",
        ]
        # 3 llamadas al juez de matching: P1 y P2 (misma clasificación en
        # BT1/SEL 551C, sin unicidad) y el hallazgo de C3 sin par.
        assert _nombres(run_dir / "etapa4_llm_judge_raw") == [
            "embeddings.json", "judge_0.json", "judge_1.json", "judge_2.json",
        ]
        assert _nombres(run_dir / "etapa4_redaccion_raw") == ["redaccion_0.json"]

        # metadata.json: tokens por etapa = suma de sus raw (control de F2).
        meta = _json(run_dir / "metadata.json")
        assert meta["tokens_checklist"] == 1_000 == _tokens_raw(run_dir / "checklist_raw")
        assert meta["tokens_e1"] == 200 == _tokens_raw(run_dir / "etapa1_raw")
        assert meta["tokens_e3"] == 4 * 30 == _tokens_raw(run_dir / "etapa3_raw")
        assert meta["tokens_e4"] == 3 * 7
        assert meta["tokens_embeddings"] == TOKENS_EMBEDDINGS
        assert meta["tokens_e4"] + meta["tokens_embeddings"] == _tokens_raw(
            run_dir / "etapa4_llm_judge_raw")
        assert meta["tokens_e4_redaccion"] == 50 == _tokens_raw(run_dir / "etapa4_redaccion_raw")
        assert meta["schema_version"] == reporte.schema_version == "0.11.0"
        assert meta["checklist_un_tramo"] == "marcar"
        assert meta["empresa_auditora"] == "Ensaut"
        assert meta["total_panos"] == 3
        assert meta["total_hallazgos_pdf"] == 9  # SEL 387 se desdobla en BT1 y CT1
        assert meta["total_hallazgos_excel"] == 10  # fila BT1-CT1 desdoblada
        assert meta["total_transversales_excel"] == 1
        assert meta["juez_redaccion"] is True
        assert isinstance(meta["duracion_segundos"], float)

        # El reporte persistido valida contra el schema y es el devuelto.
        persistido = ReporteInforme.model_validate(_json(run_dir / "reporte_final.json"))
        assert persistido == reporte

    def test_resultados_del_matching(self, entorno: dict[str, Any]) -> None:
        reporte, _, _ = _correr(entorno)

        res = reporte.resumen
        assert (res.total, res.coinciden, res.difieren, res.faltan_en_excel,
                res.extra_en_excel, res.transversales, res.coinciden_con_redaccion_distinta,
                res.requieren_revision) == (11, 6, 2, 1, 1, 1, 1, 2)

        # Invariante 14: nunca cross-paño ni cross-relé.
        for r in reporte.resultados:
            if r.hallazgo_pdf and r.hallazgo_excel:
                assert e4._clave_tupla(r.hallazgo_pdf) == e4._clave_tupla(r.hallazgo_excel)

        # Paráfrasis bajo el umbral → juez de redacción: coincide con equivalencia.
        sincro = _por_descripcion(reporte, SINCRO_PDF)
        assert sincro.estado == EstadoHallazgo.COINCIDE
        assert [e.campo for e in sincro.equivalencias_redaccion] == [CampoDiferente.DESCRIPCION]
        assert sincro.equivalencias_redaccion[0].similitud == pytest.approx(0.80, abs=1e-4)

        # Polaridad opuesta sobre el umbral → difiere sin pasar por el juez.
        e7 = _por_descripcion(reporte, E7_PDF)
        assert e7.estado == EstadoHallazgo.DIFIERE
        assert [d.campo for d in e7.diferencias] == [CampoDiferente.DESCRIPCION]
        assert "Polaridad opuesta" in (e7.diferencias[0].detalle or "")

        # Invariante 16: descripción → acción correctiva → calificación.
        ajustes = _por_descripcion(reporte, AJ_PDF)
        assert ajustes.tipo_match == TipoMatch.UNICIDAD
        assert [d.campo for d in ajustes.diferencias] == [
            CampoDiferente.DESCRIPCION, CampoDiferente.ACCION_CORRECTIVA,
            CampoDiferente.CALIFICACION,
        ]
        assert "el juez de redacción concluyó que el contenido difiere" in (
            ajustes.diferencias[0].detalle or "")

        # Sin unicidad (misma clasificación dos veces) → embeddings + juez.
        for desc in (P1, P2):
            par = _por_descripcion(reporte, desc)
            assert (par.estado, par.tipo_match) == (EstadoHallazgo.COINCIDE,
                                                    TipoMatch.SEMANTICO_ALTO)
            assert par.confianza == pytest.approx(0.92)

        assert _por_descripcion(reporte, Q_PDF).estado == EstadoHallazgo.FALTA_EN_EXCEL
        assert _por_descripcion(reporte, R5_EXCEL).estado == EstadoHallazgo.EXTRA_EN_EXCEL
        transversal = _por_descripcion(reporte, TRANSV)
        assert transversal.estado == EstadoHallazgo.TRANSVERSAL
        assert transversal.hallazgo_pdf is None

        # El checklist del informe entero llega al reporte.
        assert reporte.checklist is not None
        # 36 "cumple" del modelo; el 4 baja a no verificable: su cita no está en el PDF.
        assert reporte.checklist.cumplen == 35
        assert reporte.checklist.bajados_sin_cita_verificada == 1
        assert reporte.checklist.items[4].resultado == ResultadoChecklist.NO_CUMPLE
        punto_4 = reporte.checklist.items[3]
        assert (punto_4.resultado, punto_4.resultado_modelo) == (
            ResultadoChecklist.NO_VERIFICABLE, ResultadoChecklist.CUMPLE)
        assert punto_4.motivo_ajuste and "no aparece en la p. 1" in punto_4.motivo_ajuste
        (cita_1,) = reporte.checklist.items[0].citas
        assert cita_1.verificacion and cita_1.verificacion.estado == "verificada"

    def test_e1_e3_post_proceso_reintento_y_paginas(self, entorno: dict[str, Any]) -> None:
        reporte, run_dir, _ = _correr(entorno)
        llm: LLMFalso = entorno["llm"]

        # E1: el relé con el paño pegado y el sub-paño con prefijo se reparan.
        panos = _json(run_dir / "etapa1_panos.json")
        reles = {r["nombre"]: r["sub_panos"] for p in panos for r in p["reles"]}
        assert reles == {"SEL 751": [], "SEL 551C": ["BT1"], "SEL 387": ["BT1-CT1"],
                         "SEL 451": []}
        assert len(_json(run_dir / "etapa1_raw" / "descubrir_panos.json")["reparaciones"]) == 2

        # E3: el prompt muestra los sub-paños; el reintento lleva la
        # retroalimentación y queda registrado en ambos raw.
        llamadas_bt1 = [k for k in llm.de(e3._ListaHallazgosResponse)
                        if 'paño "BT1-CT1"' in k["user_prompt"]]
        assert len(llamadas_bt1) == 2
        assert "SEL 551C (sub-paños: BT1)" in llamadas_bt1[0]["system_prompt"]
        assert "RETROALIMENTACIÓN" in llamadas_bt1[1]["user_prompt"]
        intento1 = _json(run_dir / "etapa3_raw" / "extraer_BT1-CT1_intento1.json")
        definitivo = _json(run_dir / "etapa3_raw" / "extraer_BT1-CT1.json")
        assert intento1["problemas"] == ["el relé 'SEL 999' no está en el mapa de la Etapa 1"]
        assert definitivo["reintento"] is True
        assert definitivo["problemas_primer_intento"] == intento1["problemas"]

        # Páginas absolutas y verificadas contra la capa de texto; desdoble multi-paño.
        hallazgos = _json(run_dir / "etapa3_pdf_hallazgos.json")
        paginas = {h["descripcion"]: h["evidencia"]["pagina_pdf"] for h in hallazgos}
        assert paginas[D_DUF] == 2
        assert paginas[E7_PDF] == 3  # corregida por la verificación de citas
        assert paginas[P1] == 4
        assert paginas[Q_PDF] == 5
        sel387 = [(h["pano"], h["evidencia"]["es_multi_pano_origen"])
                  for h in hallazgos if h["rele"] == "SEL 387"]
        assert sel387 == [("BT1", True), ("CT1", True)]
        # PD60 (0.10.0): cada hallazgo del PDF queda revisado ([] = no nombra a otro
        # relé de su sección) y la marca se persiste (invariante 20).
        assert all(h["evidencia"]["otros_reles_nombrados"] is not None for h in hallazgos)

        # Estado de cada cita, a la vista: la del modelo queda cuando se corrigió.
        ev = {h["descripcion"]: h["evidencia"] for h in hallazgos}
        assert ev[D_DUF]["verificacion_descripcion"] == {
            "estado": "verificada", "pagina": 2, "pagina_citada": None, "repetida": False}
        assert ev[D_DUF]["verificacion_accion"]["estado"] == "verificada"
        assert ev[E7_PDF]["verificacion_descripcion"] == {
            "estado": "en_otra_pagina", "pagina": 3, "pagina_citada": 2, "repetida": False}
        # La acción de la sincronización no está en el PDF: se marca, no se descarta.
        assert ev[SINCRO_PDF]["verificacion_accion"] == {
            "estado": "no_encontrada", "pagina": None, "pagina_citada": 2, "repetida": False}
        # El reporte final lleva las mismas verificaciones y la hoja del Excel.
        en_reporte = {r.hallazgo_pdf.descripcion: r.hallazgo_pdf.evidencia
                      for r in reporte.resultados if r.hallazgo_pdf}
        assert en_reporte[E7_PDF].verificacion_descripcion is not None
        assert {r.hallazgo_excel.evidencia.hoja_excel
                for r in reporte.resultados if r.hallazgo_excel} == {"Tabla de Hallazgos"}
        assert reporte.resumen.filas_excel_ocultas == 0

    def test_citas_se_verifican_contra_el_pdf_con_texto(self, entorno: dict[str, Any]) -> None:
        # Con un modelo no-Claude el PDF se adecúa (rasterizado, sin capa de
        # texto, misma paginación): las citas se verifican contra el original.
        pdf = _crear_pdf(entorno["insumos"] / "informe.pdf")
        adecuado = fitz.open()
        with fitz.open(pdf) as original:
            for pagina in original:
                imagen = adecuado.new_page(width=pagina.rect.width, height=pagina.rect.height)
                imagen.insert_image(imagen.rect, pixmap=pagina.get_pixmap(dpi=40))
        ruta_adecuado = entorno["insumos"] / "informe_adecuado.pdf"
        adecuado.save(ruta_adecuado)
        excel = _crear_excel(entorno["insumos"] / "hallazgos.xlsx")

        reporte, run_dir = asyncio.run(ejecutar_pipeline(
            ruta_adecuado, excel, EmpresaAuditora.ENSAUT, ruta_pdf_texto=pdf,
        ))

        assert (run_dir / "input_pdf_texto.hash").read_text().strip() == hashlib.sha256(
            pdf.read_bytes()).hexdigest()
        ev = {r.hallazgo_pdf.descripcion: r.hallazgo_pdf.evidencia
              for r in reporte.resultados if r.hallazgo_pdf}
        v = ev[D_DUF].verificacion_descripcion
        assert v is not None and v.estado == "verificada"

    def test_troceo_persiste_troceo_json(self, entorno: dict[str, Any]) -> None:
        # 10 páginas × 100k tokens = 1M > umbral de 900k → dos tramos de 5 págs.
        entorno["contador"].por_pagina = 100_000
        reporte, run_dir, _ = _correr(entorno, paginas_extra=5)

        assert _json(run_dir / "troceo.json") == {
            "checklist": {"total_paginas": 10, "tramos": [
                {"paginas": [1, 5], "tokens": 500_000},
                {"paginas": [6, 10], "tokens": 500_000},
            ]},
            "descubrir_panos": {
                "paginas_leidas": [1, 5], "total_paginas": 10,
                # C3 termina en la última página del tramo leído.
                "advertencia": "Los paños ['C3'] terminan en la última página leída (5); "
                               "su tabla resumen podría continuar fuera del tramo.",
            },
        }
        raw = _json(run_dir / "checklist_raw" / "verificar_checklist.json")
        assert len(raw["llamadas"]) == 2
        meta = _json(run_dir / "metadata.json")
        assert meta["tokens_checklist"] == raw["tokens_total"] == 2_000

        # Fusión de tramos con páginas absolutas.
        assert reporte.checklist is not None
        items = {i.numero: (i.resultado, i.paginas) for i in reporte.checklist.items}
        assert items[1] == (ResultadoChecklist.CUMPLE, "1")
        assert items[2] == (ResultadoChecklist.CUMPLE, "7")
        # La página de la cita también pasa a absoluta (posición 2 del tramo 6-10).
        (cita_2,) = reporte.checklist.items[1].citas
        assert cita_2.pagina == 7
        assert cita_2.verificacion and cita_2.verificacion.estado == "verificada"
        assert items[3] == (ResultadoChecklist.NO_CUMPLE, "6")
        assert items[4][0] == ResultadoChecklist.NO_VERIFICABLE
        assert reporte.checklist.cumplen == 2
        # M08 (por defecto «marcar»): el 1 y el 2 los respalda una sola parte.
        assert reporte.checklist.tramos_leidos == [(1, 5), (6, 10)]
        assert (reporte.checklist.marcados_un_tramo, reporte.checklist.bajados_un_tramo) == (
            2, None)
        assert reporte.checklist.items[0].aviso_tramos == (
            "Visto en 1 de 2 partes del informe (págs. 1-5); las págs. 6-10 no lo respaldan: "
            "revísalas antes de darlo por cumplido.")
        assert reporte.checklist.items[2].aviso_tramos is None   # no cumple
        # La comparación sigue completa aunque el checklist se haya troceado.
        assert reporte.resumen.total == 11

    def test_informe_que_no_cabe_falla_en_el_flujo_de_lote(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Mismo flujo que la UI (evaluar_lote._correr_caso): con un modelo no
        Claude, un PDF sobre el límite de páginas ya no se recorta: falla con un
        error claro antes de gastar en el modelo (decisión del usuario 2026-10-18)."""
        monkeypatch.setattr(settings, "modelo_extraccion", "gpt-test")
        monkeypatch.setattr("comparador.convertir.MAX_PAGINAS_OTROS", 5)
        pdf = _crear_pdf(entorno["insumos"] / "informe.pdf", paginas_extra=1)  # 6 págs
        excel = _crear_excel(entorno["insumos"] / "Tabla de Hallazgos.xlsx")
        caso = {"nombre": "prueba", "empresa": EmpresaAuditora.ENSAUT, "informe": pdf,
                "excel": excel, "gt": None}

        resultado = asyncio.run(evaluar_lote._correr_caso(caso))

        assert "PdfExcedeLimitesError" in resultado["error"]
        run_dir = entorno["runs"] / resultado["run_id"]
        assert _json(run_dir / "error.json")["tipo"] == "PdfExcedeLimitesError"
        assert not (run_dir / "reporte_final.json").exists()
        assert not list(run_dir.glob("*_adecuado.pdf"))
        assert (run_dir / "input_informe_original.hash").read_text() == _sha256(pdf)
        # Nada se gastó: ni modelo ni count_tokens.
        assert entorno["llm"].llamadas == []
        assert entorno["contador"].llamadas == 0

    def test_reporte_se_sirve_y_exporta(self, entorno: dict[str, Any]) -> None:
        """El run nuevo se lee desde el server legacy y se exporta."""
        from comparador.server import app

        reporte, run_dir, _ = _correr(entorno)
        cliente = TestClient(app)
        for ruta in ("reporte", "export/xlsx", "export/csv", "export/feedback"):
            respuesta = cliente.get(f"/api/runs/{run_dir.name}/{ruta}")
            assert respuesta.status_code == 200, ruta
        listado = cliente.get("/api/runs")
        assert listado.status_code == 200
        assert run_dir.name in listado.text

        csv = exportar_csv(reporte)
        assert "Equivalencias de redacción" in csv
        assert "mismo incumplimiento de sincronización" in csv
        assert "mismo incumplimiento de sincronización" in exportar_feedback_md(
            reporte, FeedbackRun(), run_dir.name)
        assert exportar_xlsx(reporte)[:2] == b"PK"


class TestFallosPersistenRaw:
    def test_e3_invalido_tras_reintento_persiste_ambos_intentos(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        malo = [_h("BT1", "SEL 999", C_OPER, P1, A_P1, "Menor")]
        monkeypatch.setitem(RESPUESTAS_E3, "BT1-CT1", [malo, malo])

        with pytest.raises(ExtraccionInvalidaError, match="SEL 999"):
            _correr(entorno)

        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()
        assert {"extraer_BT1-CT1.json", "extraer_BT1-CT1_intento1.json"} <= set(
            _nombres(run_dir / "etapa3_raw"))
        assert _json(run_dir / "etapa3_raw" / "extraer_BT1-CT1.json")["problemas"]

    def test_juez_de_redaccion_invalido_persiste_los_raw_pagados(
        self, entorno: dict[str, Any]
    ) -> None:
        entorno["llm"].juez_redaccion_omite_ids = True

        with pytest.raises(JuezInvalidoError, match="ids faltantes"):
            _correr(entorno)

        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()
        assert _nombres(run_dir / "etapa4_redaccion_raw") == ["redaccion_0.json"]
        assert "embeddings.json" in _nombres(run_dir / "etapa4_llm_judge_raw")
        assert "extraer_C1.json" in _nombres(run_dir / "etapa3_raw")


class TestFallosAntesYDuranteChecklistYE1:
    """1.30 y 1.31 (05 §1.3): nada se paga antes de validar el Excel; si el checklist o E1
    fallan, queda lo que la otra dejó y la respuesta pagada de la que falló, y no hay reporte."""

    def test_excel_invalido_falla_sin_ninguna_llamada_pagada(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from comparador.etapa2_normalizar_excel import ExcelNormalizationError

        def _excel_sin_tabla(ruta: Path) -> Path:
            wb = Workbook()
            assert wb.active is not None
            wb.active.title = "Tabla de Hallazgos"
            wb.active["A1"] = "sin tabla"
            wb.save(ruta)
            return ruta

        monkeypatch.setattr(sys.modules[__name__], "_crear_excel", _excel_sin_tabla)
        with pytest.raises(ExcelNormalizationError):
            _correr(entorno)
        assert entorno["llm"].llamadas == []
        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()

    def test_si_falla_e1_queda_el_checklist_y_la_respuesta_pagada_de_e1(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(sys.modules[__name__], "RESPUESTA_E1", {"panos": []})
        with pytest.raises(e1.DescubrimientoInvalidoError, match="ningún paño"):
            _correr(entorno)
        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()
        assert len(_json(run_dir / "checklist.json")["items"]) == e0.TOTAL_PUNTOS
        assert _nombres(run_dir / "checklist_raw") == ["verificar_checklist.json"]
        fallida = _json(run_dir / "etapa1_raw" / "descubrir_panos.json")
        assert fallida["llamada_fallida"] is True and fallida["response_raw"] == {"panos": []}
        assert not (run_dir / "etapa1_panos.json").exists()
        # El run fallido también dice con qué código y criterios corrió.
        assert len(_json(run_dir / "version.json")["sha256_codigo"]) == 16

    def test_criterios_del_checklist_invalidos_fallan_antes_de_gastar(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        roto = tmp_path / "puntos.yaml"
        roto.write_text("puntos:\n  - {numero: 1, enunciado: a, donde_buscar: b}\n",
                        encoding="utf-8")
        monkeypatch.setattr(e0, "RUTA_PUNTOS", roto)
        with pytest.raises(e0.ChecklistConfigError):
            _correr(entorno)
        assert entorno["llm"].llamadas == [] and entorno["contador"].llamadas == 0

    def test_si_falla_el_checklist_quedan_los_panos_y_la_respuesta_pagada(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(LLMFalso, "_checklist", staticmethod(
            lambda kw: {"items": [{"numero": 1, "cita": None, "resultado": "cumple"}]}))
        with pytest.raises(e0.ChecklistIncompletoError, match="no respondió los puntos"):
            _correr(entorno)
        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()
        assert not (run_dir / "checklist.json").exists()
        assert _json(run_dir / "checklist_raw" / "verificar_checklist.json")["llamada_fallida"]
        assert _json(run_dir / "etapa1_panos.json")
        assert _nombres(run_dir / "etapa1_raw") == ["descubrir_panos.json"]

    def test_si_fallan_los_dos_sale_el_error_del_checklist_con_nota_del_otro(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(sys.modules[__name__], "RESPUESTA_E1", {"panos": []})
        monkeypatch.setattr(LLMFalso, "_checklist", staticmethod(
            lambda kw: {"items": [{"numero": 1, "cita": None, "resultado": "cumple"}]}))
        with pytest.raises(e0.ChecklistIncompletoError) as info:
            _correr(entorno)
        assert any("DescubrimientoInvalidoError" in n for n in info.value.__notes__)

    def test_tramo_que_falla_guarda_los_tramos_ya_pagados(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # Dos tramos (como test_troceo_persiste_troceo_json); el segundo responde incompleto.
        entorno["contador"].por_pagina = 100_000
        original = LLMFalso._checklist

        def _segundo_incompleto(kw: dict[str, Any]) -> dict[str, Any]:
            respuesta = original(kw)
            if "contiene las páginas 6 a" in kw["user_prompt"]:
                respuesta["items"] = respuesta["items"][:3]
            return respuesta

        monkeypatch.setattr(LLMFalso, "_checklist", staticmethod(_segundo_incompleto))
        with pytest.raises(e0.ChecklistIncompletoError):
            _correr(entorno, paginas_extra=5)
        (run_dir,) = list(entorno["runs"].iterdir())
        raw = _json(run_dir / "checklist_raw" / "verificar_checklist.json")
        assert raw["troceo"]["total_paginas"] == 10
        assert [bool(m.get("llamada_fallida")) for m in raw["llamadas"]] == [False, True]
        assert not (run_dir / "checklist.json").exists()


class TestLlamadasFallidasEnE3YE4:
    """Revisión adversarial de F6e: el primer intento de un paño que falla en el cliente y un
    error de E4 que no viene de un juez también dejan guardado lo ya pagado."""

    def test_primer_intento_de_e3_que_falla_en_el_cliente_deja_su_raw(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        original = LLMFalso._e3

        def _e3(self: LLMFalso, kw: dict[str, Any]) -> dict[str, Any]:
            if 'del paño "C1"' in kw["user_prompt"]:
                raise llm_client.LlamadaFallidaError("respuesta que no valida", {
                    "modelo": kw["modelo"], "llamada_fallida": True,
                    "error": "respuesta que no valida", "tokens_prompt": 640,
                    "tokens_completion": 1, "usage_parcial": True,
                })
            return original(self, kw)

        monkeypatch.setattr(LLMFalso, "_e3", _e3)
        with pytest.raises(ExtraccionInvalidaError) as info:
            _correr(entorno)
        assert isinstance(info.value.__cause__, llm_client.LlamadaFallidaError)
        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()
        raw = _json(run_dir / "etapa3_raw" / "extraer_C1.json")
        assert (raw["llamada_fallida"], raw["tokens_prompt"], raw["pano"]) == (True, 640, "C1")

    def test_un_error_de_e4_que_no_viene_del_juez_deja_lo_pagado(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def _falla(*a: Any, **kw: Any) -> Any:
            raise ValueError("bug simulado en 4e")

        monkeypatch.setattr(e4, "_revisar_redaccion", _falla)
        with pytest.raises(ValueError, match="bug simulado"):
            _correr(entorno)
        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()
        nombres = _nombres(run_dir / "etapa4_llm_judge_raw")
        assert "embeddings.json" in nombres and any(n.startswith("judge_") for n in nombres)


class TestFallosDelJuezPersistenLoPagado:
    """1.32 (05 §1.3): si una llamada de un juez de E4 falla en el cliente, sale el error
    original (sin veredicto por defecto) y quedan los raw ya pagados, con la entrada fallida."""

    @staticmethod
    def _falla_en(monkeypatch: pytest.MonkeyPatch, metodo: str, n: int) -> None:
        original = getattr(LLMFalso, metodo)
        llamadas = {"n": 0}

        def _envuelto(self: LLMFalso, kw: dict[str, Any]) -> dict[str, Any]:
            llamadas["n"] += 1
            if llamadas["n"] == n:
                raise RuntimeError("Error de red simulado tras los reintentos")
            return original(self, kw) if metodo == "_juez_redaccion" else original(kw)

        monkeypatch.setattr(LLMFalso, metodo, _envuelto)

    def test_falla_del_juez_de_matching(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._falla_en(monkeypatch, "_juez_matching", 2)
        with pytest.raises(RuntimeError, match="Error de red simulado"):
            _correr(entorno)
        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()
        jueces = [_json(p) for p in sorted((run_dir / "etapa4_llm_judge_raw").glob("judge_*"))]
        assert [bool(j.get("llamada_fallida")) for j in jueces] == [False, True]
        assert "RuntimeError" in jueces[1]["error"] and jueces[1]["modelo"]
        assert jueces[1]["pano"] and jueces[1]["rele"]

    def test_falla_del_juez_de_redaccion(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._falla_en(monkeypatch, "_juez_redaccion", 1)
        with pytest.raises(RuntimeError, match="Error de red simulado"):
            _correr(entorno)
        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()
        fallida = _json(run_dir / "etapa4_redaccion_raw" / "redaccion_0.json")
        assert fallida["llamada_fallida"] is True and fallida["tipo"] == "juez_redaccion"
        assert fallida["lote"]
        assert "embeddings.json" in _nombres(run_dir / "etapa4_llm_judge_raw")


def _error_openai(tipo: str) -> Exception:
    """Los errores que da el SDK de OpenAI con la clave revocada o sin fondos."""
    import httpx
    import openai

    pedido = httpx.Request("POST", "https://api.openai.com/v1/embeddings")
    if tipo == "autenticacion":
        return openai.AuthenticationError(
            "Incorrect API key provided", response=httpx.Response(401, request=pedido),
            body={"code": "invalid_api_key"})
    return openai.RateLimitError(
        "You exceeded your current quota", response=httpx.Response(429, request=pedido),
        body={"code": "insufficient_quota", "type": "insufficient_quota"})


class TestPruebaClaveOpenAI:
    """PD66: todo run que pasa por E4 calcula el embedding de una palabra antes de
    la primera llamada pagada a Claude (la misma función que usará E4: la cubren
    los falsos de siempre). Si la clave no responde, falla sin gastar en Claude.
    El modo solo checklist y la repetición estricta de E4 no la hacen."""

    @staticmethod
    def _registrar(monkeypatch: pytest.MonkeyPatch, llm: LLMFalso,
                   falla: Exception | None = None) -> list[tuple[list[str], int]]:
        """Cada llamada a los embeddings con cuántas llamadas al modelo había antes."""
        vistos: list[tuple[list[str], int]] = []

        def _calcular(textos: list[str]) -> tuple[NDArray[np.float64], dict[str, Any]]:
            vistos.append((list(textos), len(llm.llamadas)))
            if falla is not None:
                raise falla
            return embeddings_falsos(textos)

        monkeypatch.setattr(e4, "calcular_embeddings", _calcular)
        return vistos

    def test_el_pipeline_la_hace_antes_de_claude_y_la_registra(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        vistos = self._registrar(monkeypatch, entorno["llm"])

        _, run_dir, _ = _correr(entorno)

        assert vistos[0] == ([llm_embeddings.PALABRA_PRUEBA], 0)   # ninguna llamada antes
        assert len(vistos) == 2 and vistos[1][0] != [llm_embeddings.PALABRA_PRUEBA]  # E4
        prueba = _json(run_dir / "metadata.json")["prueba_clave_openai"]
        assert {k: v for k, v in prueba.items() if k != "segundos"} == {
            "ok": True, "modelo": settings.modelo_embeddings, "modelo_servido": None,
            "tokens": TOKENS_EMBEDDINGS}
        # No es una llamada de E4: ni raw ni tokens de embeddings.
        assert _nombres(run_dir / "etapa4_llm_judge_raw").count("embeddings.json") == 1
        assert _json(run_dir / "metadata.json")["tokens_embeddings"] == TOKENS_EMBEDDINGS

    @pytest.mark.parametrize(("tipo", "causa"), [
        ("autenticacion", "autenticación rechazada (401)"),
        ("sin_fondos", "sin fondos o cuota agotada (429 insufficient_quota)"),
    ])
    def test_si_la_clave_no_responde_el_pipeline_no_llama_a_claude(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tipo: str, causa: str
    ) -> None:
        self._registrar(monkeypatch, entorno["llm"], _error_openai(tipo))

        with pytest.raises(llm_embeddings.ClaveOpenAIError) as info:
            _correr(entorno)

        assert "La clave de OpenAI no responde: ¿sin fondos o revocada?" in str(info.value)
        assert causa in str(info.value)
        assert entorno["llm"].llamadas == [] and entorno["contador"].llamadas == 0
        (run_dir,) = list(entorno["runs"].iterdir())
        for no_llego in ("checklist.json", "etapa1_panos.json", "metadata.json"):
            assert not (run_dir / no_llego).exists(), no_llego

    def test_en_el_flujo_de_lote_la_fila_lleva_el_error(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        self._registrar(monkeypatch, entorno["llm"], _error_openai("sin_fondos"))
        caso = {"nombre": "prueba", "empresa": EmpresaAuditora.ENSAUT, "gt": None,
                "informe": _crear_pdf(entorno["insumos"] / "informe.pdf"),
                "excel": _crear_excel(entorno["insumos"] / "Tabla de Hallazgos.xlsx")}

        resultado = asyncio.run(evaluar_lote._correr_caso(caso))

        assert "ClaveOpenAIError" in resultado["error"] and "no responde" in resultado["error"]
        assert entorno["llm"].llamadas == []
        error = _json(entorno["runs"] / resultado["run_id"] / "error.json")
        assert error["tipo"] == "ClaveOpenAIError"

    def test_la_cli_completa_no_crea_runs_si_la_clave_no_responde(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        caso = tmp_path / "informes" / "S-E Prueba (Auditor Ensaut)"
        caso.mkdir(parents=True)
        _crear_pdf(caso / "informe.pdf")
        _crear_excel(caso / "Tabla de Hallazgos.xlsx")
        salida = entorno["runs"].parent / "evaluacion_B3.json"
        vistos = self._registrar(monkeypatch, entorno["llm"], _error_openai("autenticacion"))
        monkeypatch.setattr(sys, "argv", [
            "evaluar_lote", "--modelos", "claude-test", "--salida", str(salida),
            "--base", str(tmp_path / "informes")])

        with pytest.raises(SystemExit) as info:
            evaluar_lote.main()

        assert isinstance(info.value.code, str) and "no responde" in info.value.code
        assert vistos == [([llm_embeddings.PALABRA_PRUEBA], 0)]
        assert entorno["llm"].llamadas == [] and not salida.exists()
        assert not entorno["runs"].exists() or not list(entorno["runs"].iterdir())

    def test_la_cli_completa_registra_la_prueba_en_la_salida(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        caso = tmp_path / "informes" / "S-E Prueba (Auditor Ensaut)"
        caso.mkdir(parents=True)
        _crear_pdf(caso / "informe.pdf")
        _crear_excel(caso / "Tabla de Hallazgos.xlsx")
        salida = entorno["runs"].parent / "evaluacion_B3.json"
        for campo in ("runs_a_retener", "runs_incompletos_a_retener"):
            monkeypatch.setattr(settings, campo, getattr(settings, campo))
        vistos = self._registrar(monkeypatch, entorno["llm"])
        monkeypatch.setattr(sys, "argv", [
            "evaluar_lote", "--modelos", "claude-test", "--salida", str(salida),
            "--base", str(tmp_path / "informes")])

        evaluar_lote.main()

        # La del lote (antes de crear el run) y la del pipeline (antes de Claude).
        assert [v for v in vistos if v[0] == [llm_embeddings.PALABRA_PRUEBA]] == [
            ([llm_embeddings.PALABRA_PRUEBA], 0), ([llm_embeddings.PALABRA_PRUEBA], 0)]
        assert _json(salida)["prueba_clave_openai"]["ok"] is True
        (fila,) = _json(salida)["resultados"]["claude-test"]
        assert "error" not in fila

    def test_solo_checklist_no_la_hace(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.pipeline import ejecutar_solo_checklist

        vistos = self._registrar(monkeypatch, entorno["llm"], AssertionError("no se llama"))
        asyncio.run(ejecutar_solo_checklist(_crear_pdf(entorno["insumos"] / "informe.pdf"),
                                            empresa_auditora=EmpresaAuditora.ENSAUT))
        caso = tmp_path / "informes" / "S-E Prueba (Auditor Ensaut)"
        caso.mkdir(parents=True)
        _crear_pdf(caso / "informe.pdf")
        _crear_excel(caso / "Tabla de Hallazgos.xlsx")
        salida = entorno["runs"].parent / "evaluacion_B3s.json"
        for campo in ("cita_checklist", "runs_a_retener", "runs_incompletos_a_retener"):
            monkeypatch.setattr(settings, campo, getattr(settings, campo))
        monkeypatch.setattr(sys, "argv", [
            "evaluar_lote", "--modelos", "claude-test", "--salida", str(salida),
            "--base", str(tmp_path / "informes"), "--solo-checklist", "--sin-cita-checklist"])

        evaluar_lote.main()

        assert vistos == []
        assert _json(salida)["prueba_clave_openai"] is None

    def test_reproducir_e4_la_hace_salvo_en_estricto(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        llamadas_antes = len(entorno["llm"].llamadas)
        vistos = self._registrar(monkeypatch, entorno["llm"])

        meta = reproducir_e4(run_dir, tmp_path / "mixto", cache="mixto")

        assert vistos == [([llm_embeddings.PALABRA_PRUEBA], llamadas_antes)]
        assert meta["prueba_clave_openai"]["ok"] is True
        assert _json(tmp_path / "mixto" / "metadata.json")["prueba_clave_openai"]["ok"] is True

        vistos.clear()
        meta = reproducir_e4(run_dir, tmp_path / "estricto", cache="estricto")
        assert vistos == [] and meta["prueba_clave_openai"] is None

    def test_reproducir_e4_no_gasta_si_la_clave_no_responde(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        llamadas_antes = len(entorno["llm"].llamadas)
        self._registrar(monkeypatch, entorno["llm"], _error_openai("sin_fondos"))

        with pytest.raises(llm_embeddings.ClaveOpenAIError, match="insufficient_quota"):
            reproducir_e4(run_dir, tmp_path / "replay")

        assert len(entorno["llm"].llamadas) == llamadas_antes
        assert not (tmp_path / "replay").exists()


class TestGuardia:
    def test_las_funciones_reales_chocan_con_el_guardia(self, entorno: dict[str, Any]) -> None:
        """Si un fake deja de interceptar, el código real cae en el guardia
        antes de crear un cliente: ningún camino llega a un proveedor."""
        llamadas: list[Callable[[], object]] = [
            lambda: llm_embeddings.calcular_embeddings(["texto"]),
            lambda: llm_client.contar_tokens_pdf(
                modelo="claude-test", system_prompt="s", user_prompt="u", pdf_base64=""),
            lambda: llm_client.llamar_llm_estructurado(
                modelo="claude-test", system_prompt="s", user_prompt="u",
                response_model=e1._ListaPanosResponse),
            lambda: llm_client.llamar_llm_estructurado(
                modelo="gpt-test", system_prompt="s", user_prompt="u",
                response_model=e1._ListaPanosResponse),
        ]
        for llamada in llamadas:
            with pytest.raises(AssertionError, match="cliente real"):
                llamada()


class TestSoloChecklist:
    """`pipeline.ejecutar_solo_checklist` (ronda B3): solo la Etapa 0, con la misma
    persistencia que el pipeline, sin gastar en E1-E4 ni exigir OpenAI."""

    def _correr(self, entorno: dict[str, Any], paginas_extra: int = 0) -> tuple[Any, Path]:
        from comparador.pipeline import ejecutar_solo_checklist

        pdf = _crear_pdf(entorno["insumos"] / "informe.pdf", paginas_extra)
        return asyncio.run(ejecutar_solo_checklist(pdf, empresa_auditora=EmpresaAuditora.ENSAUT))

    def test_persiste_el_checklist_y_nada_mas(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(settings, "openai_api_key", "")   # no hay embeddings
        checklist, run_dir = self._correr(entorno)

        assert checklist.total == 38
        meta = _json(run_dir / "metadata.json")
        assert meta["modo"] == "solo_checklist" and meta["cita_checklist"] is True
        assert meta["checklist_un_tramo"] == "marcar"
        assert meta["tokens_checklist"] == _tokens_raw(run_dir / "checklist_raw")
        assert (meta["tokens_e1"], meta["tokens_e3"], meta["tokens_e4"]) == (0, 0, 0)
        assert _json(run_dir / "checklist.json")["total"] == 38
        assert (run_dir / "version.json").exists() and (run_dir / "input_pdf.hash").exists()
        for no_corrio in ("etapa1_panos.json", "etapa2_excel_normalizado.json",
                          "etapa3_pdf_hallazgos.json", "reporte_final.json"):
            assert not (run_dir / no_corrio).exists(), no_corrio
        llm: LLMFalso = entorno["llm"]
        assert llm.de(e1._ListaPanosResponse) == [] and llm.de(e3._ListaHallazgosResponse) == []
        from comparador.persistence import es_solo_checklist

        assert es_solo_checklist(run_dir)

    def test_sin_cita(self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setattr(settings, "cita_checklist", False)
        checklist, run_dir = self._correr(entorno)
        assert checklist.bajados_sin_cita_verificada is None
        assert _json(run_dir / "metadata.json")["cita_checklist"] is False
        assert [k["response_model"] for k in entorno["llm"].llamadas] == [e0._ChecklistResponse]

    def test_con_troceo(self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch) -> None:
        entorno["contador"].por_pagina = 100_000
        monkeypatch.setattr(settings, "checklist_un_tramo", "bajar")
        _, run_dir = self._correr(entorno, paginas_extra=5)
        troceo = _json(run_dir / "troceo.json")
        assert troceo["descubrir_panos"] is None and len(troceo["checklist"]["tramos"]) == 2
        assert _json(run_dir / "metadata.json")["checklist_un_tramo"] == "bajar"
        guardado = _json(run_dir / "checklist.json")
        # El 1 y el 2 (fuera de la lista de PD44) se marcan; ninguno baja.
        assert (guardado["tramos_leidos"], guardado["marcados_un_tramo"],
                guardado["bajados_un_tramo"]) == ([[1, 5], [6, 10]], 2, 0)

    def test_si_falla_guarda_lo_pagado_y_no_escribe_metadata(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(LLMFalso, "_checklist", staticmethod(
            lambda kw: {"items": [{"numero": 1, "cita": None, "resultado": "cumple"}]}))
        with pytest.raises(e0.ChecklistIncompletoError):
            self._correr(entorno)
        (run_dir,) = list(entorno["runs"].iterdir())
        assert _json(run_dir / "checklist_raw" / "verificar_checklist.json")["llamada_fallida"]
        assert not (run_dir / "metadata.json").exists()

    def test_en_el_flujo_de_lote_con_planilla(self, entorno: dict[str, Any]) -> None:
        from tests.test_evaluar_checklist import _gt_sintetico

        pdf = _crear_pdf(entorno["insumos"] / "informe.pdf")
        excel = _crear_excel(entorno["insumos"] / "Tabla de Hallazgos.xlsx")
        caso = {"nombre": "prueba", "empresa": EmpresaAuditora.ENSAUT, "informe": pdf,
                "excel": excel, "gt": _gt_sintetico(entorno["insumos"])}

        resultado = asyncio.run(evaluar_lote._correr_caso(caso, solo_checklist=True))
        fila = evaluar_lote._resumen_caso(caso, resultado)

        assert "error" not in fila, fila.get("error")
        assert fila["modo"] == "solo_checklist" and "comparador" not in fila
        assert fila["checklist_gt"]["evaluados"] >= 1
        assert fila["citas"]["checklist"]["bajados_sin_cita_verificada"] == 1
        # La misma regla (etapa0.cita_respalda) baja el punto y lo cuenta (PD52).
        assert len(fila["citas"]["checklist"]["puntos_bajados"]) == 1
        run_dir = entorno["runs"] / resultado["run_id"]
        assert not (run_dir / "input.xlsx").exists()
        assert _json(run_dir / "inputs.json")["modo"] == "solo_checklist"


    def test_solo_checklist_que_falla_queda_marcado_en_la_fila(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Un run solo checklist que falla no es un error del comparador:
        la fila de la evaluación lleva el modo (metricas_lote la salta)."""
        monkeypatch.setattr(settings, "modelo_extraccion", "gpt-test")
        monkeypatch.setattr("comparador.convertir.MAX_PAGINAS_OTROS", 5)
        pdf = _crear_pdf(entorno["insumos"] / "informe.pdf", paginas_extra=1)  # 6 págs
        excel = _crear_excel(entorno["insumos"] / "Tabla de Hallazgos.xlsx")
        caso = {"nombre": "prueba", "empresa": EmpresaAuditora.ENSAUT, "informe": pdf,
                "excel": excel, "gt": None}

        resultado = asyncio.run(evaluar_lote._correr_caso(caso, solo_checklist=True))
        fila = evaluar_lote._resumen_caso(caso, resultado)

        assert "PdfExcedeLimitesError" in fila["error"]
        assert fila["modo"] == "solo_checklist"
        assert entorno["llm"].llamadas == []
        completo = evaluar_lote._resumen_caso(caso, {**resultado, "modo": "completo"})
        assert "modo" not in completo

    def test_la_cli_corre_solo_el_checklist_sin_cita(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        """De punta a punta por la CLI, como se correrá B3s: flags, correr_lote,
        configuración en la salida y la salida junto a sus runs (RUNS_DIR =
        <carpeta de la salida>/runs, como pide evaluar_lote)."""
        caso = tmp_path / "informes" / "S-E Prueba (Auditor Ensaut)"
        caso.mkdir(parents=True)
        _crear_pdf(caso / "informe.pdf")
        _crear_excel(caso / "Tabla de Hallazgos.xlsx")
        salida = entorno["runs"].parent / "evaluacion_B3s.json"
        # main() y correr_lote cambian estos valores globales: que vuelvan al final.
        for campo in ("cita_checklist", "runs_a_retener", "runs_incompletos_a_retener"):
            monkeypatch.setattr(settings, campo, getattr(settings, campo))
        argumentos = ["evaluar_lote", "--modelos", "claude-test", "--etiqueta", "B3s",
                      "--salida", str(salida), "--base", str(tmp_path / "informes"),
                      "--solo-checklist", "--sin-cita-checklist"]
        monkeypatch.setattr(sys, "argv", argumentos)

        evaluar_lote.main()

        evaluacion = _json(salida)
        assert evaluacion["config"]["solo_checklist"] is True
        assert evaluacion["config"]["cita_checklist"] is False
        assert evaluacion["config"]["checklist_un_tramo"] == "marcar"
        (fila,) = evaluacion["resultados"]["claude-test"]
        assert fila["modo"] == "solo_checklist" and "comparador" not in fila
        # Solo el checklist y sin la cita: ninguna otra etapa llamó al modelo.
        assert {k["response_model"] for k in entorno["llm"].llamadas} == {
            e0._ChecklistResponse}
        # Una segunda corrida no pisa la medición (y no gasta: corta antes).
        llamadas = len(entorno["llm"].llamadas)
        with pytest.raises(SystemExit, match="ya existe"):
            evaluar_lote.main()
        assert len(entorno["llm"].llamadas) == llamadas

    def test_verifica_las_citas_contra_el_pdf_con_texto(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Con un modelo que recibe el PDF rasterizado (sin capa de texto), las
        citas del checklist se verifican contra el PDF original con texto."""
        import fitz

        def _rasterizar(ruta: Path, modelo: str | None = None) -> tuple[Path, dict[str, Any]]:
            destino = ruta.with_name(ruta.stem + "_adecuado.pdf")
            with fitz.open(ruta) as src, fitz.open() as out:
                for pagina in src:
                    nueva = out.new_page(width=pagina.rect.width, height=pagina.rect.height)
                    nueva.insert_image(nueva.rect, stream=pagina.get_pixmap(dpi=60).tobytes("png"))
                out.save(destino)
            return destino, {"rasterizado": True, "truncado": False}

        monkeypatch.setattr(evaluar_lote, "adecuar_pdf_para_llm", _rasterizar)
        pdf = _crear_pdf(entorno["insumos"] / "informe.pdf")
        excel = _crear_excel(entorno["insumos"] / "Tabla de Hallazgos.xlsx")
        caso = {"nombre": "prueba", "empresa": EmpresaAuditora.ENSAUT, "informe": pdf,
                "excel": excel, "gt": None}

        resultado = asyncio.run(evaluar_lote._correr_caso(caso, solo_checklist=True))
        fila = evaluar_lote._resumen_caso(caso, resultado)

        run_dir = entorno["runs"] / resultado["run_id"]
        assert (run_dir / "input_pdf_texto.hash").read_text() == _sha256(run_dir / "input.pdf")
        # Igual que sin rasterizar: solo baja el punto cuya cita no está en el PDF.
        assert fila["citas"]["checklist"]["bajados_sin_cita_verificada"] == 1

    def test_la_cli_no_gasta_si_la_planilla_no_se_lee(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        caso = tmp_path / "informes" / "S-E Prueba (Auditor Ensaut)"
        caso.mkdir(parents=True)
        _crear_pdf(caso / "informe.pdf")
        _crear_excel(caso / "Tabla de Hallazgos.xlsx")
        (caso / "CheckList Revisión.xlsx").write_bytes(b"no es un xlsx")
        salida = entorno["runs"].parent / "evaluacion_B3.json"
        monkeypatch.setattr(sys, "argv", [
            "evaluar_lote", "--modelos", "claude-test", "--salida", str(salida),
            "--base", str(tmp_path / "informes")])

        with pytest.raises(SystemExit, match="No se puede leer la planilla"):
            evaluar_lote.main()
        assert entorno["llm"].llamadas == [] and not list(entorno["runs"].glob("*"))

    @pytest.mark.parametrize("salida", ["carpeta", "runs/evaluacion.json", "runs"])
    def test_la_cli_rechaza_una_salida_que_no_sirve(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
        salida: str,
    ) -> None:
        destino = {"carpeta": tmp_path / "mediciones" / "b3",
                   "runs/evaluacion.json": entorno["runs"] / "b3" / "evaluacion.json",
                   "runs": entorno["runs"]}[salida]
        monkeypatch.setattr(sys, "argv", [
            "evaluar_lote", "--modelos", "claude-test", "--salida", str(destino),
            "--base", str(tmp_path)])

        with pytest.raises(SystemExit):
            evaluar_lote.main()
        assert entorno["llm"].llamadas == []
        assert not (entorno["runs"] / "b3").exists()

    def test_una_falla_al_evaluar_no_bota_el_lote(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        pdf = _crear_pdf(entorno["insumos"] / "informe.pdf")
        excel = _crear_excel(entorno["insumos"] / "Tabla de Hallazgos.xlsx")
        casos = [{"nombre": n, "empresa": EmpresaAuditora.ENSAUT, "informe": pdf,
                  "excel": excel, "gt": None} for n in ("bien", "mal")]
        original = evaluar_lote._resumen_caso

        def _resumen(caso: dict[str, Any], resultado: dict[str, Any]) -> dict[str, Any]:
            if caso["nombre"] == "mal":
                raise KeyError("punto 39")
            return original(caso, resultado)

        monkeypatch.setattr(evaluar_lote, "_resumen_caso", _resumen)
        for campo in ("runs_a_retener", "runs_incompletos_a_retener"):
            monkeypatch.setattr(settings, campo, getattr(settings, campo))

        resultados = asyncio.run(evaluar_lote.correr_lote(["claude-test"], casos, 2, True))

        filas = {f["informe"]: f for f in resultados["claude-test"]}
        assert "error" not in filas["bien"]
        assert filas["mal"]["error"].startswith("evaluación: KeyError")
        # Solo checklist también en la fila de error: metricas_lote no la cuenta
        # como un comparador fallido.
        assert filas["mal"]["modo"] == "solo_checklist"
        from comparador import persistence
        assert not persistence._EN_CURSO

def _contenido(carpeta: Path) -> dict[Path, bytes]:
    return {p.relative_to(carpeta): p.read_bytes() for p in carpeta.rglob("*") if p.is_file()}


class TestAsignacionGlobal:
    def test_pipeline_con_el_interruptor(self, entorno: dict[str, Any],
                                          monkeypatch: pytest.MonkeyPatch) -> None:
        # Sin conflictos en el escenario: mismas consultas y resultados, y la
        # traza de cada consulta queda en su raw.
        monkeypatch.setattr(settings, "asignacion_global", True)
        reporte, run_dir, _ = _correr(entorno)

        assert _nombres(run_dir / "etapa4_llm_judge_raw") == [
            "embeddings.json", "judge_0.json", "judge_1.json", "judge_2.json"]
        meta = _json(run_dir / "metadata.json")
        assert (meta["asignacion_global"], meta["llamadas_juez_reintento"]) == (True, 0)
        assert meta["tokens_e4"] == 3 * 7
        assert _json(run_dir / "etapa4_llm_judge_raw" / "judge_0.json")["asignacion_global"][
            "resultado"] in {"tomo_fila_libre", "sin_match"}
        assert reporte.resumen.total == 11

    def test_por_defecto_encendido(self, entorno: dict[str, Any]) -> None:
        # Decisión del usuario 2026-10-26 (PD71): encendida por defecto, con la
        # traza en cada raw del juez de matching y registrada en metadata y config_e4.
        _, run_dir, _ = _correr(entorno)
        meta = _json(run_dir / "metadata.json")
        assert (meta["asignacion_global"], meta["modo"]) == (True, "completo")
        assert _json(run_dir / "config_e4.json")["asignacion_global"] is True
        assert "asignacion_global" in _json(run_dir / "etapa4_llm_judge_raw" / "judge_0.json")

    def test_apagado_comportamiento_previo(self, entorno: dict[str, Any],
                                           monkeypatch: pytest.MonkeyPatch) -> None:
        # ASIGNACION_GLOBAL=false = comportamiento previo exacto: en orden, sin traza.
        monkeypatch.setattr(settings, "asignacion_global", False)
        _, run_dir, _ = _correr(entorno)
        meta = _json(run_dir / "metadata.json")
        assert (meta["asignacion_global"], meta["modo"]) == (False, "completo")
        assert _json(run_dir / "config_e4.json")["asignacion_global"] is False
        assert "asignacion_global" not in _json(run_dir / "etapa4_llm_judge_raw" / "judge_0.json")

    @pytest.mark.parametrize(("en_el_run", "en_el_proceso"),
                             [(False, True), (True, False), (True, True)],
                             ids=["run-apagado-proceso-encendido", "run-encendido-proceso-apagado",
                                  "run-y-proceso-encendidos-no-mide"])
    def test_reproducir_e4_sobre_un_run_guardado(self, entorno: dict[str, Any],
                                                 monkeypatch: pytest.MonkeyPatch,
                                                 tmp_path: Path, en_el_run: bool,
                                                 en_el_proceso: bool) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        monkeypatch.setattr(settings, "asignacion_global", en_el_run)
        reporte, run_dir, _ = _correr(entorno)
        llamadas_antes = len(entorno["llm"].llamadas)
        # ASIGNACION_GLOBAL es lo que se mide: sale del proceso, no del run (el
        # control de PD26 repite con false los runs nuevos, que corren encendidos).
        monkeypatch.setattr(settings, "asignacion_global", en_el_proceso)
        # La configuración del proceso es otra: el replay usa la del run.
        otra = {"modelo_extraccion": "claude-otro", "modelo_juez_matching": "claude-juez-otro",
                "juez_redaccion": False,
                "threshold_descripcion": 0.5, "threshold_accion": 0.5, "top_k_candidatos": 1,
                "modelo_embeddings": "embeddings-otro"}
        for campo, valor in otra.items():
            monkeypatch.setattr(settings, campo, valor)
        salida = tmp_path / "replay"

        meta = reproducir_e4(run_dir, salida)

        assert meta["modo"] == "replay_e4" and meta["run_origen"] == run_dir.name
        assert meta["conteo_replay"] == meta["conteo_original"]
        assert (salida / "reporte_final.json").exists()
        assert (salida / "etapa4_matches.json").exists()
        # Solo se repitió E4: 3 jueces de matching + 1 de redacción, nada de E0-E3,
        # con el modelo y los interruptores del run.
        nuevas = entorno["llm"].llamadas[llamadas_antes:]
        assert {k["response_model"] for k in nuevas} == {
            e4._JuezResponse, e4._JuezRedaccionResponse}
        assert {k["modelo"] for k in nuevas} == {"claude-test"}
        meta_run = _json(run_dir / "metadata.json")
        for campo in ("modelo_extraccion", "juez_redaccion", "threshold_descripcion",
                      "threshold_accion", "top_k_candidatos", "modelo_embeddings"):
            assert meta[campo] == meta_run[campo], campo
        assert (meta_run["asignacion_global"], meta["asignacion_global"]) == (
            en_el_run, en_el_proceso)
        # La del run queda al lado y, si es la misma, el aviso de que no midió
        # el interruptor (el control de PD26 lanzado sin ASIGNACION_GLOBAL=false).
        assert meta["asignacion_global_run"] is en_el_run
        if en_el_run == en_el_proceso:
            assert meta["aviso_asignacion_global"] == (
                "AVISO: la repetición usa la misma asignación global que el run (True): no mide "
                "el interruptor (el control de PD26 repite con ASIGNACION_GLOBAL=false)")
        else:
            assert meta["aviso_asignacion_global"] is None
        # Al terminar vuelve la configuración del proceso.
        assert {campo: getattr(settings, campo) for campo in otra} == otra
        assert ReporteInforme.model_validate(_json(salida / "reporte_final.json")).checklist == (
            reporte.checklist)

    def test_el_control_de_pd26_sin_fijar_la_asignacion_avisa_que_no_mide(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
        capsys: pytest.CaptureFixture[str]
    ) -> None:
        """Un run nuevo corre con el default (encendida) y el control de PD26 se
        lanza olvidando ASIGNACION_GLOBAL=false: la repetición sale de la caché
        sin pagar nada y con el mismo conteo, en verde sin haber medido el
        interruptor. El preflight y la repetición lo avisan (no lo cortan: así
        se repite B2 con la caché estricta)."""
        from comparador import reproducir_e4 as modulo

        _, run_dir, _ = _correr(entorno)
        assert _json(run_dir / "config_e4.json")["asignacion_global"] is True
        sin_gt = tmp_path / "gt"
        sin_gt.mkdir()
        argv = ["reproducir_e4", str(run_dir), "--salida", str(tmp_path / "replay"),
                "--cache", "mixto", "--gt-comparador", str(sin_gt)]
        aviso = ("AVISO: la repetición usa la misma asignación global que el run (True): no "
                 "mide el interruptor (el control de PD26 repite con ASIGNACION_GLOBAL=false)")

        monkeypatch.setattr(sys, "argv", [*argv, "--preflight"])
        modulo.main()
        preflight = capsys.readouterr().out
        assert "asignación global: True (el run: True)" in preflight
        assert aviso in preflight
        monkeypatch.setattr(sys, "argv", argv)
        modulo.main()

        assert aviso in capsys.readouterr().out
        meta = _json(tmp_path / "replay" / "metadata.json")
        assert (meta["asignacion_global"], meta["asignacion_global_run"]) == (True, True)
        assert meta["cache_jueces"]["pagadas"] == 0
        assert meta["aviso_asignacion_global"] == aviso
        # Con ASIGNACION_GLOBAL=false, el control sí mide: sin aviso.
        monkeypatch.setattr(settings, "asignacion_global", False)
        monkeypatch.setattr(sys, "argv", [*argv[:3], str(tmp_path / "control"), *argv[4:]])
        modulo.main()
        assert "AVISO: la repetición usa la misma" not in capsys.readouterr().out
        assert _json(tmp_path / "control" / "metadata.json")["aviso_asignacion_global"] is None

    def test_reproducir_e4_no_toca_el_run_de_origen(self, entorno: dict[str, Any],
                                                    tmp_path: Path) -> None:
        """Los runs medidos viven fuera de runs/ (mediciones/): la salida tampoco
        puede quedar dentro del run de origen ni mezclarse con otra repetición."""
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = tmp_path / "mediciones" / run_dir.name
        shutil.copytree(run_dir, copia)
        antes = _contenido(copia)
        llamadas_antes = len(entorno["llm"].llamadas)

        for salida in (copia, copia / "replay"):
            with pytest.raises(ValueError, match="dentro del run de origen"):
                reproducir_e4(copia, salida)
        ocupada = tmp_path / "replay_anterior"
        ocupada.mkdir()
        (ocupada / "judge_9.json").write_text("{}")
        with pytest.raises(ValueError, match="ya existe con contenido"):
            reproducir_e4(copia, ocupada)

        assert _contenido(copia) == antes
        assert len(entorno["llm"].llamadas) == llamadas_antes

    def test_una_salida_que_git_no_ignora_falla_antes_de_escribir(
        self, entorno: dict[str, Any], tmp_path: Path
    ) -> None:
        """Dentro de un repositorio git, --salida y --solo-vectores (el .npz y su
        .json) van a una carpeta que git ignore: llevan texto de los informes."""
        from comparador.reproducir_e4 import calcular_vectores_del_run, reproducir_e4
        from tests.test_persistence import repo_git

        _, run_dir, _ = _correr(entorno)
        copia = tmp_path / "mediciones" / run_dir.name
        shutil.copytree(run_dir, copia)
        repo = repo_git(tmp_path, "mediciones/\n*.npz\n")
        llamadas_antes = len(entorno["llm"].llamadas)

        with pytest.raises(ValueError, match="git no la ignora"):
            reproducir_e4(copia, repo / "replay")
        with pytest.raises(ValueError, match="git no la ignora"):   # el .json de al lado
            calcular_vectores_del_run(copia, repo / "vectores" / "v.npz")

        assert sorted(p.name for p in repo.iterdir()) == [".git", ".gitignore"]
        assert len(entorno["llm"].llamadas) == llamadas_antes
        calcular_vectores_del_run(copia, repo / "mediciones" / "v.npz")
        reproducir_e4(copia, repo / "mediciones" / "replay")
        assert (repo / "mediciones" / "replay" / "reporte_final.json").is_file()

    def test_reproducir_e4_rechaza_la_salida_con_otras_mayusculas_o_si_es_archivo(
        self, entorno: dict[str, Any], tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = tmp_path / "mediciones" / run_dir.name
        shutil.copytree(run_dir, copia)
        archivo = tmp_path / "salida.json"
        archivo.write_text("{}")
        with pytest.raises(ValueError, match="ya existe con contenido"):
            reproducir_e4(copia, archivo)
        # En un disco que no distingue mayúsculas (el de macOS), RUNS/ es runs/.
        runs_mayusculas = Path(str(entorno["runs"]).upper())
        if not runs_mayusculas.exists():
            pytest.skip("este disco distingue mayúsculas")
        with pytest.raises(ValueError, match="fuera de"):
            reproducir_e4(copia, runs_mayusculas / "replay")
        with pytest.raises(ValueError, match="dentro del run de origen"):
            reproducir_e4(copia, Path(str(copia).upper()) / "replay")

    def test_reproducir_e4_que_falla_guarda_lo_pagado_y_restaura_la_configuracion(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        monkeypatch.setattr(settings, "modelo_extraccion", "claude-otro")
        llamadas = 0

        def _juez_que_se_equivoca(**kw: Any) -> Any:
            nonlocal llamadas
            llamadas += 1
            return (e4._JuezResponse(indice_candidato=9, confianza=90, justificacion="x"),
                    {"tokens_total": 5, "modelo": kw["modelo"]})

        monkeypatch.setattr(e4, "llamar_llm_estructurado", _juez_que_se_equivoca)
        salida = tmp_path / "replay"

        with pytest.raises(e4.JuezInvalidoError):
            reproducir_e4(run_dir, salida)

        # Lo pagado queda guardado (invariante 20) y vuelve la configuración del proceso.
        # Los vectores salieron del run (M70): no se pagaron de nuevo.
        assert llamadas >= 1
        assert _nombres(salida / "etapa4_llm_judge_raw") == ["judge_0.json"]
        assert not (salida / "etapa4_embeddings.npz").exists()
        assert not (salida / "reporte_final.json").exists()
        assert settings.modelo_extraccion == "claude-otro"

        # Un run sin vectores (como B2): los que pagó la repetición también quedan.
        sin_vectores = tmp_path / "mediciones" / run_dir.name
        shutil.copytree(run_dir, sin_vectores)
        (sin_vectores / "etapa4_embeddings.npz").unlink()
        with pytest.raises(e4.JuezInvalidoError):
            reproducir_e4(sin_vectores, tmp_path / "replay2")
        assert {"embeddings.json", "judge_0.json"} <= set(
            _nombres(tmp_path / "replay2" / "etapa4_llm_judge_raw"))
        assert (tmp_path / "replay2" / "etapa4_embeddings.npz").is_file()

    def test_reproducir_e4_desde_dentro_del_run_registra_su_nombre(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = tmp_path / "mediciones" / run_dir.name
        shutil.copytree(run_dir, copia)
        monkeypatch.chdir(copia)

        meta = reproducir_e4(Path("."), tmp_path / "replay")

        assert meta["run_origen"] == run_dir.name

    def test_reproducir_e4_run_anterior_a_4e_no_enciende_el_juez_de_redaccion(
        self, entorno: dict[str, Any], tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = tmp_path / "mediciones" / run_dir.name
        shutil.copytree(run_dir, copia)
        meta_run = _json(copia / "metadata.json")
        del meta_run["juez_redaccion"]   # como en B0 y B1, que tampoco tienen config_e4.json
        (copia / "metadata.json").write_text(json.dumps(meta_run), encoding="utf-8")
        (copia / "config_e4.json").unlink()
        llamadas_antes = len(entorno["llm"].llamadas)

        meta = reproducir_e4(copia, tmp_path / "replay")

        assert meta["juez_redaccion"] is False
        nuevas = entorno["llm"].llamadas[llamadas_antes:]
        assert {k["response_model"] for k in nuevas} == {e4._JuezResponse}

    def test_reproducir_e4_no_escribe_dentro_de_runs(self, entorno: dict[str, Any]) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        with pytest.raises(ValueError, match="fuera de"):
            reproducir_e4(run_dir, entorno["runs"] / "replay")

    @staticmethod
    def _gt_del_run(run_dir: Path, reporte: ReporteInforme, carpeta: Path,
                    items: list[tuple[str, str, str]], nombre: str = "gt.json") -> Path:
        """GT sintético atado a los insumos del run: un ítem por (descripción,
        clase, subclase), con las claves del resultado de esa descripción."""
        from comparador.adjudicacion import campos_que_difieren
        from comparador.gt_comparador import GtInforme, serializar
        from tests.test_evaluar_comparador import gt_de

        def item(i: int, descripcion: str, clase: str, subclase: str) -> dict[str, Any]:
            r = _por_descripcion(reporte, descripcion)
            p, e = r.hallazgo_pdf, r.hallazgo_excel
            return {
                "id": f"i{i}", "pano": r.pano, "rele": r.rele, "estado": r.estado.value,
                "campos": [c.value for c in campos_que_difieren(r)],
                "lado_excel": None if e is None else {"fila": e.evidencia.fila_excel},
                "lado_pdf": None if p is None else {
                    "pagina": p.evidencia.pagina_pdf, "clasificaciones": [p.clasificacion],
                    "citas": [{"texto": p.descripcion, "fuente": "adjudicacion",
                               "modelo": "claude-test"}]},
                "clase": clase, "subclase": subclase, "nivel_validacion": "verificado",
                "depende_de": ["P14"] if clase == "C" else [],
                "etapa_causa": "e4" if clase == "D" else None,
                "origenes": [{"modelo": "claude-test", "ronda": "B2", "run_id": run_dir.name,
                              "idx": i, "tipo_match": r.tipo_match.value, "subclase": subclase,
                              "nivel_validacion": "verificado"}],
            }

        gt = gt_de([item(i, *x) for i, x in enumerate(items)],
                   universos={"claude-test": "todos"}, run_id=run_dir.name)
        carpeta.mkdir(exist_ok=True)
        ruta = carpeta / nombre
        ruta.write_text(serializar(GtInforme.model_validate({
            **gt.model_dump(mode="json"),
            "sha256_informe_original": (run_dir / "input_informe_original.hash").read_text(),
            "sha256_excel": (run_dir / "input_excel.hash").read_text(),
        })), encoding="utf-8")
        return ruta

    def test_reproducir_e4_evalua_contra_el_gt_antes_y_despues(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.evaluar_comparador import evaluar_run, gts_de, resumen
        from comparador.reproducir_e4 import reproducir_e4

        reporte, run_dir, _ = _correr(entorno)
        # Como evaluar_lote: el sha256 del informe original (aquí, el mismo PDF).
        shutil.copyfile(run_dir / "input_pdf.hash", run_dir / "input_informe_original.hash")
        ruta_gt = self._gt_del_run(run_dir, reporte, tmp_path / "gt",
                                   [(P1, "OK", "OK"), (P2, "A", "A4")])
        # En la repetición el juez ya no empareja P1: el par OK se pierde.
        monkeypatch.delitem(JUEZ_MATCHING, P1)
        salida = tmp_path / "replay"

        meta = reproducir_e4(run_dir, salida, gt_dir=tmp_path / "gt")

        gt = meta["comparador_gt"]
        assert gt["sha256_gt"] == hashlib.sha256(ruta_gt.read_bytes()).hexdigest()
        antes, despues = gt["original"], gt["replay"]
        assert (antes["OK"], antes["A"]) == (1, 1)
        assert (despues["OK"], despues["A"]) == (0, 1)
        # P1 sale como falta y como extra, que el GT no tiene (el resto de los
        # resultados no está en este GT de dos ítems: sin adjudicar en ambos).
        assert despues["sin_adjudicar"] == antes["sin_adjudicar"] + 2
        assert gt["delta"]["OK"] == -1 and gt["delta"]["D"] == 0
        assert gt["delta"]["sin_adjudicar"] == 2
        assert gt["delta"]["a_perdidos_nuevos"] == [] and gt["delta"]["a_recuperados"] == []
        assert gt["delta"]["posibles_escondidos_nuevos"] == []
        assert _json(salida / "metadata.json")["comparador_gt"] == gt
        # La repetición queda atada a los mismos insumos: metricas_lote la
        # recalcula sin buscar el run de origen y obtiene lo mismo.
        for nombre in ("input_informe_original.hash", "input_excel.hash"):
            assert (salida / nombre).read_bytes() == (run_dir / nombre).read_bytes()
        (aplicable,) = gts_de(tmp_path / "gt")
        assert resumen(evaluar_run(salida, aplicable.gt)) == despues
        assert resumen(evaluar_run(run_dir, aplicable.gt)) == antes

    def test_reproducir_e4_sin_gt_aplicable_lo_deja_dicho(
        self, entorno: dict[str, Any], tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)   # sin input_informe_original.hash: no se ata
        (tmp_path / "gt").mkdir()

        meta = reproducir_e4(run_dir, tmp_path / "replay", gt_dir=tmp_path / "gt")

        assert meta["comparador_gt"] is None

    def test_reproducir_e4_con_gt_ambiguo_falla_antes_de_gastar(
        self, entorno: dict[str, Any], tmp_path: Path
    ) -> None:
        from comparador.evaluar_comparador import GtNoAplicaError
        from comparador.reproducir_e4 import reproducir_e4

        reporte, run_dir, _ = _correr(entorno)
        shutil.copyfile(run_dir / "input_pdf.hash", run_dir / "input_informe_original.hash")
        for nombre in ("a.json", "b.json"):
            self._gt_del_run(run_dir, reporte, tmp_path / "gt", [(P1, "OK", "OK")], nombre)
        llamadas_antes = len(entorno["llm"].llamadas)

        with pytest.raises(GtNoAplicaError, match="más de un GT"):
            reproducir_e4(run_dir, tmp_path / "replay", gt_dir=tmp_path / "gt")

        assert len(entorno["llm"].llamadas) == llamadas_antes
        assert not (tmp_path / "replay").exists()


class TestVectoresYCacheE4:
    """M70: cada run guarda los vectores de E4 (etapa4_embeddings.npz) y
    reproducir_e4 repite E4 con ellos y con la caché de los jueces del run:
    estricta (sin API ni credenciales) o mixta (paga solo lo que falta)."""

    @staticmethod
    def _hallazgos(run_dir: Path) -> tuple[list[Any], list[Any]]:
        from comparador.schemas import Hallazgo

        return ([Hallazgo.model_validate(h) for h in _json(run_dir / "etapa3_pdf_hallazgos.json")],
                [Hallazgo.model_validate(h)
                 for h in _json(run_dir / "etapa2_excel_normalizado.json")])

    @staticmethod
    def _copia(run_dir: Path, tmp_path: Path, sin: tuple[str, ...] = ()) -> Path:
        copia = tmp_path / "mediciones" / run_dir.name
        shutil.copytree(run_dir, copia)
        for nombre in sin:
            (copia / nombre).unlink()
        return copia

    @staticmethod
    def _sin_api(monkeypatch: pytest.MonkeyPatch) -> None:
        from comparador import reproducir_e4 as modulo

        def _prohibido(*_: Any, **__: Any) -> Any:
            raise AssertionError("la repetición estricta tocó la API o las credenciales")

        monkeypatch.setattr(e4, "llamar_llm_estructurado", _prohibido)
        monkeypatch.setattr(e4, "calcular_embeddings", _prohibido)
        monkeypatch.setattr(modulo, "_validar_credenciales", _prohibido)
        monkeypatch.setattr(settings, "openai_api_key", "")
        monkeypatch.setattr(settings, "anthropic_api_key", "")

    @staticmethod
    def _resultados(ruta: Path) -> list[dict[str, Any]]:
        return [r.model_dump(mode="json")
                for r in ReporteInforme.model_validate(_json(ruta)).resultados]

    def test_el_run_guarda_los_vectores_de_e4(self, entorno: dict[str, Any]) -> None:
        from comparador.persistence import cargar_vectores

        _, run_dir, _ = _correr(entorno)

        vectores, modelo = cargar_vectores(run_dir / "etapa4_embeddings.npz")
        textos = e4.textos_a_embeber(*self._hallazgos(run_dir))
        esperados, _ = embeddings_falsos(textos)
        assert list(vectores) == textos and modelo == settings.modelo_embeddings
        assert all(np.array_equal(vectores[t], esperados[i]) for i, t in enumerate(textos))
        # En la raíz del run, fuera de los *_raw/ que metricas_lote recorre como
        # llamadas (test_metricas_lote: los CSV no cambian con el archivo).
        assert [p.name for p in run_dir.rglob("*.npz")] == ["etapa4_embeddings.npz"]
        assert (run_dir / "etapa4_embeddings.npz").is_file()

    def test_si_e4_falla_despues_de_pagar_quedan_los_vectores(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from comparador.persistence import cargar_vectores

        TestFallosDelJuezPersistenLoPagado._falla_en(monkeypatch, "_juez_matching", 1)
        with pytest.raises(RuntimeError, match="Error de red simulado"):
            _correr(entorno)
        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "reporte_final.json").exists()
        vectores, _ = cargar_vectores(run_dir / "etapa4_embeddings.npz")
        assert list(vectores) == e4.textos_a_embeber(*self._hallazgos(run_dir))

    def test_si_tampoco_se_pueden_guardar_los_vectores_sale_el_error_de_e4(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from comparador import pipeline

        TestFallosDelJuezPersistenLoPagado._falla_en(monkeypatch, "_juez_matching", 1)

        def _no_guarda(*_: Any) -> Any:
            raise OSError("disco lleno")

        monkeypatch.setattr(pipeline, "guardar_vectores", _no_guarda)
        with pytest.raises(RuntimeError, match="Error de red simulado") as info:
            _correr(entorno)
        assert info.value.__notes__ == [
            "Tampoco se guardaron los vectores de E4: disco lleno"]
        (run_dir,) = list(entorno["runs"].iterdir())
        assert "judge_0.json" in _nombres(run_dir / "etapa4_llm_judge_raw")

    def test_si_fallan_los_embeddings_no_queda_un_archivo_vacio(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        def _falla(textos: list[str]) -> Any:
            if textos == [llm_embeddings.PALABRA_PRUEBA]:   # la prueba de la clave pasa
                return embeddings_falsos(textos)
            raise RuntimeError("OpenAI caído")

        monkeypatch.setattr(e4, "calcular_embeddings", _falla)
        with pytest.raises(RuntimeError, match="OpenAI caído"):
            _correr(entorno)
        (run_dir,) = list(entorno["runs"].iterdir())
        assert not (run_dir / "etapa4_embeddings.npz").exists()

    def test_repeticion_estricta_sin_api_ni_credenciales(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.metricas_lote import llamadas_del_run
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        self._sin_api(monkeypatch)
        salida = tmp_path / "replay"

        meta = reproducir_e4(run_dir, salida, cache="estricto")

        assert self._resultados(salida / "reporte_final.json") == self._resultados(
            run_dir / "reporte_final.json")
        cache = meta["cache_jueces"]
        assert (cache["modo"], cache["cargadas"], cache["desde_cache"], cache["pagadas"],
                cache["sin_usar"]) == ("estricto", 4, 4, 0, 0)
        textos = e4.textos_a_embeber(*self._hallazgos(run_dir))
        assert meta["vectores"] == {
            "origen": "run", "archivo": "etapa4_embeddings.npz",
            "sha256": _sha256(run_dir / "etapa4_embeddings.npz"),
            "cargados": len(textos), "calculados": 0,
        }
        assert (meta["tokens_e4"], meta["tokens_e4_redaccion"], meta["tokens_embeddings"]) == (
            0, 0, 0)
        assert _json(salida / "metadata.json")["cache_jueces"] == cache
        # Nada se pagó: ni vectores nuevos que guardar ni registro de embeddings.
        assert not (salida / "etapa4_embeddings.npz").exists()
        assert _nombres(salida / "etapa4_llm_judge_raw") == [
            "judge_0.json", "judge_1.json", "judge_2.json"]
        llamadas = llamadas_del_run(salida)
        assert len(llamadas) == 4
        assert all(ll.desde_cache and ll.tokens_input == ll.tokens_output == 0 for ll in llamadas)
        assert _json(salida / "etapa4_llm_judge_raw" / "judge_1.json")["raw_origen"] == str(
            run_dir.resolve() / "etapa4_llm_judge_raw" / "judge_1.json")

    def test_repeticion_estricta_falla_ante_una_llamada_que_no_esta(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.cache_jueces import CacheJuecesFaltanteError
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path, sin=("etapa4_llm_judge_raw/judge_1.json",))
        self._sin_api(monkeypatch)

        with pytest.raises(CacheJuecesFaltanteError, match="estricto"):
            reproducir_e4(copia, tmp_path / "replay", cache="estricto")
        assert not (tmp_path / "replay" / "reporte_final.json").exists()
        # La llamada que la caché no tenía no se hizo: no queda como llamada
        # fallida ni metricas_lote la cuenta como pagada.
        from comparador.metricas_lote import llamadas_del_run

        raws = [_json(p) for p in (tmp_path / "replay").rglob("*_raw/*.json")]
        assert raws and not any(r.get("llamada_fallida") for r in raws)
        llamadas = llamadas_del_run(tmp_path / "replay")
        assert all(ll.desde_cache for ll in llamadas) and not any(ll.fallida for ll in llamadas)

    def test_repeticion_estricta_sin_vectores_falla_antes_de_empezar(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path, sin=("etapa4_embeddings.npz",))   # como B2
        self._sin_api(monkeypatch)

        with pytest.raises(e4.VectoresFaltantesError, match="estricto"):
            reproducir_e4(copia, tmp_path / "replay", cache="estricto")
        assert not (tmp_path / "replay").exists()

    def test_repeticion_mixta_paga_solo_lo_que_falta(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path, sin=("etapa4_llm_judge_raw/judge_1.json",))
        monkeypatch.setattr(e4, "calcular_embeddings", _solo_la_prueba_de_la_clave)
        llamadas_antes = len(entorno["llm"].llamadas)
        salida = tmp_path / "replay"

        meta = reproducir_e4(copia, salida, cache="mixto")

        nuevas = entorno["llm"].llamadas[llamadas_antes:]
        assert [k["response_model"] for k in nuevas] == [e4._JuezResponse]
        assert (meta["cache_jueces"]["desde_cache"], meta["cache_jueces"]["pagadas"]) == (3, 1)
        assert meta["tokens_e4"] == 7 and meta["tokens_e4_redaccion"] == 0
        assert self._resultados(salida / "reporte_final.json") == self._resultados(
            run_dir / "reporte_final.json")
        pagado = _json(salida / "etapa4_llm_judge_raw" / "judge_1.json")
        assert "desde_cache" not in pagado and pagado["tokens_total"] == 7

    def test_una_repeticion_mixta_sirve_de_cache_a_una_estricta(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        """La base de B2 (D1): lo que pagó una repetición mixta queda como caché
        de las repeticiones estrictas siguientes, que leen la caché de esa salida
        (`cache_desde`) en vez de la del run de origen."""
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path, sin=("etapa4_llm_judge_raw/judge_1.json",))
        monkeypatch.setattr(e4, "calcular_embeddings", _solo_la_prueba_de_la_clave)
        base = tmp_path / "base"
        reproducir_e4(copia, base, cache="mixto")
        self._sin_api(monkeypatch)

        meta = reproducir_e4(copia, tmp_path / "replay", cache="estricto", cache_desde=base)

        cache = meta["cache_jueces"]
        assert (cache["desde_cache"], cache["pagadas"], cache["sin_usar"]) == (4, 0, 0)
        assert cache["desde"] == str(base.resolve())
        assert self._resultados(tmp_path / "replay" / "reporte_final.json") == self._resultados(
            base / "reporte_final.json")
        juez = tmp_path / "replay" / "etapa4_llm_judge_raw"
        # Lo pagado por la base apunta a la base; lo que la base sirvió de la caché
        # conserva su primer origen, el run.
        assert _json(juez / "judge_1.json")["raw_origen"] == str(
            base.resolve() / "etapa4_llm_judge_raw" / "judge_1.json")
        assert _json(juez / "judge_0.json")["raw_origen"] == str(
            copia.resolve() / "etapa4_llm_judge_raw" / "judge_0.json")
        # Sin la base, la misma repetición estricta no encuentra judge_1.
        from comparador.cache_jueces import CacheJuecesFaltanteError
        with pytest.raises(CacheJuecesFaltanteError):
            reproducir_e4(copia, tmp_path / "otra", cache="estricto")

    def test_la_repeticion_restaura_el_effort_del_run(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        """El effort va en la request (la caché lo compara): la repetición usa el
        de los raw del run, no el del entorno, lo registra y al terminar vuelve
        el del proceso."""
        from comparador.reproducir_e4 import reproducir_e4

        monkeypatch.setattr(settings, "effort_llm", "high")
        _, run_dir, _ = _correr(entorno)
        assert _json(run_dir / "etapa4_llm_judge_raw" / "judge_0.json")[
            "parametros_thinking"] == {"output_config": {"effort": "high"}}
        monkeypatch.setattr(settings, "effort_llm", "")
        self._sin_api(monkeypatch)

        meta = reproducir_e4(run_dir, tmp_path / "replay", cache="estricto")

        assert (meta["effort_llm"], meta["effort_llm_origen"]) == ("high", "raw")
        assert meta["cache_jueces"]["desde_cache"] == 4
        assert settings.effort_llm == ""

    def test_effort_distinto_entre_los_raw_falla_antes_de_crear_la_salida(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path)
        raw = copia / "etapa4_llm_judge_raw" / "judge_1.json"
        datos = _json(raw)
        datos["parametros_thinking"] = {"output_config": {"effort": "low"}}
        raw.write_text(json.dumps(datos), encoding="utf-8")
        self._sin_api(monkeypatch)

        with pytest.raises(ValueError, match="effort"):
            reproducir_e4(copia, tmp_path / "replay", cache="estricto")
        assert not (tmp_path / "replay").exists()

    def test_sin_effort_en_los_raw_vale_el_registrado_o_el_del_entorno(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        """Raw sin parametros_thinking (OpenAI, o Claude anterior al registro):
        el effort de config_e4.json o, en un run anterior a ese archivo, el del
        entorno, siempre registrado en la metadata de la repetición."""
        from comparador.reproducir_e4 import reproducir_e4

        monkeypatch.setattr(settings, "effort_llm", "high")
        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path)
        for raw in copia.glob("etapa4_*_raw/*_*.json"):
            datos = _json(raw)
            datos.pop("parametros_thinking", None)
            raw.write_text(json.dumps(datos), encoding="utf-8")
        monkeypatch.setattr(settings, "effort_llm", "medium")

        meta = reproducir_e4(copia, tmp_path / "a")
        assert (meta["effort_llm"], meta["effort_llm_origen"]) == ("high", "config_e4.json")
        (copia / "config_e4.json").unlink()
        meta = reproducir_e4(copia, tmp_path / "b")
        assert (meta["effort_llm"], meta["effort_llm_origen"]) == ("medium", "configuracion_actual")
        assert settings.effort_llm == "medium"

    def test_cache_desde_exige_una_cache_y_una_carpeta_con_raw(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        self._sin_api(monkeypatch)
        with pytest.raises(ValueError, match="cache_desde"):
            reproducir_e4(run_dir, tmp_path / "a", cache="no", cache_desde=run_dir)
        vacia = tmp_path / "vacia"
        vacia.mkdir()
        with pytest.raises(ValueError, match="no tiene raw de los jueces"):
            reproducir_e4(run_dir, tmp_path / "b", cache="estricto", cache_desde=vacia)
        # La base de otro run (revisión del Bloque 24): ninguna clave acertaría y se
        # pagarían todos los jueces, mezclando su variación con el efecto del cambio.
        otra = tmp_path / "base-de-otro-run"
        shutil.copytree(run_dir / "etapa4_llm_judge_raw", otra / "etapa4_llm_judge_raw")
        (otra / "metadata.json").write_text(
            json.dumps({"modo": "replay_e4", "run_origen": "otro-run"}), encoding="utf-8")
        with pytest.raises(ValueError, match=r"es de otro run \(otro-run, no "):
            reproducir_e4(run_dir, tmp_path / "c", cache="estricto", cache_desde=otra)
        assert not any((tmp_path / d).exists() for d in ("a", "b", "c"))

    def test_vectores_de_otro_modelo_fallan_antes_de_gastar(
        self, entorno: dict[str, Any], tmp_path: Path
    ) -> None:
        from comparador.persistence import cargar_vectores, escribir_vectores
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        vectores, _ = cargar_vectores(run_dir / "etapa4_embeddings.npz")
        otros = tmp_path / "otros.npz"
        escribir_vectores(otros, vectores, "text-embedding-3-small")
        llamadas_antes = len(entorno["llm"].llamadas)

        with pytest.raises(ValueError, match="text-embedding-3-small"):
            reproducir_e4(run_dir, tmp_path / "replay", vectores=otros)
        assert len(entorno["llm"].llamadas) == llamadas_antes
        assert not (tmp_path / "replay").exists()

    def test_vectores_indicados_se_usan_y_quedan_en_la_metadata(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path, sin=("etapa4_embeddings.npz",))
        aparte = tmp_path / "vectores_e4" / f"{run_dir.name}.npz"
        aparte.parent.mkdir()
        shutil.copyfile(run_dir / "etapa4_embeddings.npz", aparte)
        self._sin_api(monkeypatch)

        meta = reproducir_e4(copia, tmp_path / "replay", vectores=aparte, cache="estricto")

        assert (meta["vectores"]["origen"], meta["vectores"]["archivo"]) == (
            "indicado", aparte.name)
        assert meta["vectores"]["sha256"] == _sha256(aparte)

    def test_sin_vectores_en_el_run_los_paga_y_los_guarda_en_la_salida(
        self, entorno: dict[str, Any], tmp_path: Path
    ) -> None:
        from comparador.persistence import cargar_vectores
        from comparador.reproducir_e4 import reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path, sin=("etapa4_embeddings.npz",))
        salida = tmp_path / "replay"

        meta = reproducir_e4(copia, salida)

        textos = e4.textos_a_embeber(*self._hallazgos(run_dir))
        assert meta["vectores"] == {"origen": None, "archivo": None, "sha256": None,
                                    "cargados": 0, "calculados": len(textos)}
        assert meta["cache_jueces"] == {"modo": "no"}
        assert list(cargar_vectores(salida / "etapa4_embeddings.npz")[0]) == textos
        assert "embeddings.json" in _nombres(salida / "etapa4_llm_judge_raw")

    def test_solo_vectores_calcula_y_guarda_sin_tocar_el_run(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador.persistence import cargar_vectores
        from comparador.pipeline import CredencialesFaltantesError
        from comparador.reproducir_e4 import calcular_vectores_del_run, reproducir_e4

        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path, sin=("etapa4_embeddings.npz",))   # como B2
        antes = _contenido(copia)
        destino = tmp_path / "b2" / "vectores_e4" / f"{copia.name}.npz"

        resumen = calcular_vectores_del_run(copia, destino)

        assert _contenido(copia) == antes
        vectores, modelo = cargar_vectores(destino)
        textos = e4.textos_a_embeber(*self._hallazgos(copia))
        assert list(vectores) == textos
        assert modelo == _json(copia / "metadata.json")["modelo_embeddings"]
        guardado = _json(destino.with_suffix(".json"))
        assert guardado == resumen
        assert (guardado["run_origen"], guardado["textos"]) == (copia.name, len(textos))
        assert guardado["uso"]["tokens_total"] == TOKENS_EMBEDDINGS
        assert guardado["sha256_npz"] == _sha256(destino)
        assert guardado["similitudes"]["distintas"] == []
        assert guardado["similitudes"]["comparadas"] >= 2
        for malo, motivo in ((destino, "ya existe"), (copia / "v.npz", "dentro del run"),
                             (tmp_path / "v.json", ".npz"),
                             (entorno["runs"] / "v.npz", "fuera de")):
            with pytest.raises((ValueError, FileExistsError), match=motivo):
                calcular_vectores_del_run(copia, malo)
        monkeypatch.setattr(settings, "openai_api_key", "")
        with pytest.raises(CredencialesFaltantesError, match="OPENAI_API_KEY"):
            calcular_vectores_del_run(copia, tmp_path / "otro.npz")
        assert not (tmp_path / "otro.npz").exists()
        # Con esos vectores, la repetición estricta reproduce el run sin API.
        self._sin_api(monkeypatch)
        reproducir_e4(copia, tmp_path / "replay", vectores=destino, cache="estricto")
        assert self._resultados(tmp_path / "replay" / "reporte_final.json") == (
            self._resultados(copia / "reporte_final.json"))

    def test_verificar_similitudes_detecta_vectores_distintos_o_faltantes(
        self, entorno: dict[str, Any]
    ) -> None:
        from comparador.persistence import cargar_vectores
        from comparador.reproducir_e4 import verificar_similitudes

        _, run_dir, _ = _correr(entorno)
        vectores, _ = cargar_vectores(run_dir / "etapa4_embeddings.npz")
        lote = _json(run_dir / "etapa4_redaccion_raw" / "redaccion_0.json")["lote"]
        equivalencias = [e for r in _json(run_dir / "reporte_final.json")["resultados"]
                         for e in r["equivalencias_redaccion"]]

        detalles = [d for r in _json(run_dir / "reporte_final.json")["resultados"]
                    for d in r["diferencias"] if "imilitud semántica" in (d["detalle"] or "")]
        jueces = list((run_dir / "etapa4_llm_judge_raw").glob("judge_*.json"))

        bien = verificar_similitudes(run_dir, vectores)

        assert detalles and len(jueces) == 3
        assert bien == {"comparadas": len(lote) + len(equivalencias) + len(detalles),
                        "distintas": [],
                        "orden_top_k": {"consultas": 3, "fuera_de_orden": []}}
        corridos = {**vectores, SINCRO_EXCEL: vectores[SINCRO_EXCEL] + 0.05}
        mal = verificar_similitudes(run_dir, corridos)
        assert {d["origen"] for d in mal["distintas"]} == {
            "etapa4_redaccion_raw/redaccion_0.json", "reporte_final.json"}
        assert all(d["guardada"] == 0.8 and d["recalculada"] != 0.8 for d in mal["distintas"])
        sin_uno = {t: v for t, v in vectores.items() if t != SINCRO_PDF}
        assert {d["recalculada"] for d in verificar_similitudes(run_dir, sin_uno)["distintas"]} == {
            None}


    def test_verificar_similitudes_detecta_un_top_k_en_otro_orden(self, tmp_path: Path) -> None:
        """El juez vio A antes que B: con vectores donde B se parece más al
        hallazgo del PDF, el prompt de hoy sería otro (la caché no lo encontraría)."""
        from comparador.llm.prompts import user_juez_matching
        from comparador.reproducir_e4 import verificar_similitudes

        carpeta = tmp_path / "run" / "etapa4_llm_judge_raw"
        carpeta.mkdir(parents=True)
        (carpeta / "judge_0.json").write_text(json.dumps(
            {"user_prompt_text": user_juez_matching("X", ["A", "B"])}), encoding="utf-8")
        vectores = {"X": np.array([1.0, 0.0]), "A": np.array([0.9, 0.1]),
                    "B": np.array([0.0, 1.0])}

        bien = verificar_similitudes(tmp_path / "run", vectores)
        mal = verificar_similitudes(tmp_path / "run", {**vectores, "B": np.array([1.0, 0.01])})
        sin_b = verificar_similitudes(tmp_path / "run", {"X": vectores["X"], "A": vectores["A"]})

        assert bien["orden_top_k"] == {"consultas": 1, "fuera_de_orden": []}
        assert mal["orden_top_k"]["fuera_de_orden"] == [
            {"origen": "etapa4_llm_judge_raw/judge_0.json", "similitudes": [0.9939, 1.0]}]
        assert sin_b["orden_top_k"]["fuera_de_orden"][0]["similitudes"] is None

    def test_la_cli_repite_estricto_y_calcula_solo_vectores(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
        capsys: pytest.CaptureFixture[str]
    ) -> None:
        from comparador import reproducir_e4 as modulo

        _, run_dir, _ = _correr(entorno)
        copia = self._copia(run_dir, tmp_path, sin=("etapa4_embeddings.npz",))
        destino = tmp_path / "vectores" / f"{copia.name}.npz"
        monkeypatch.setattr(sys, "argv", ["reproducir_e4", str(copia), "--solo-vectores",
                                          str(destino)])
        modulo.main()
        assert destino.is_file() and "similitudes" in capsys.readouterr().out

        self._sin_api(monkeypatch)
        monkeypatch.setattr(sys, "argv", [
            "reproducir_e4", str(copia), "--salida", str(tmp_path / "replay"),
            "--vectores", str(destino), "--cache", "estricto"])
        modulo.main()
        salida = capsys.readouterr().out
        assert "caché estricto: 4 desde la caché, 0 pagadas" in salida

        monkeypatch.setattr(sys, "argv", [
            "reproducir_e4", str(copia), "--salida", str(tmp_path / "replay2"),
            "--vectores", str(destino), "--cache", "estricto",
            "--cache-desde", str(tmp_path / "replay")])
        modulo.main()
        assert (f"caché estricto (desde {(tmp_path / 'replay').resolve()}): 4 desde la caché, "
                "0 pagadas") in capsys.readouterr().out

        for argv in (["--salida", "x", "--solo-vectores", "y.npz"], ["--cache", "estricto"],
                     ["--salida", "x", "--cache", "talvez"],
                     ["--salida", "x", "--cache-desde", str(tmp_path / "no-existe"),
                      "--cache", "estricto"]):
            monkeypatch.setattr(sys, "argv", ["reproducir_e4", str(copia), *argv])
            with pytest.raises(SystemExit):
                modulo.main()

    @pytest.mark.parametrize("modo", ["--salida", "--solo-vectores"])
    def test_la_cli_con_codigo_congelado_distinto_sale_sin_llamar_a_nada(
        self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, modo: str
    ) -> None:
        """E10: `--codigo-congelado` antes de todo, en las dos formas de la CLI."""
        from comparador import reproducir_e4 as modulo
        from tests.test_persistence import copiar_src

        llamados: list[str] = []

        def _no_llamar(nombre: str) -> Callable[..., Any]:
            def _registrar(*_: Any, **__: Any) -> Any:
                llamados.append(nombre)
                raise AssertionError(f"{nombre} no debía llamarse")
            return _registrar

        monkeypatch.setattr(modulo, "reproducir_e4", _no_llamar("reproducir_e4"))
        monkeypatch.setattr(modulo, "calcular_vectores_del_run",
                            _no_llamar("calcular_vectores_del_run"))
        run = tmp_path / "run"
        run.mkdir()
        (run / "etapa3_pdf_hallazgos.json").write_text("[]", encoding="utf-8")
        copia = copiar_src(tmp_path / "src")
        (copia / "comparador" / "reproducir_e4.py").write_text("# otra\n", encoding="utf-8")
        destino = tmp_path / ("x.npz" if modo == "--solo-vectores" else "replay")
        monkeypatch.setattr(sys, "argv", ["reproducir_e4", str(run), modo, str(destino),
                                          "--codigo-congelado", str(copia)])

        with pytest.raises(SystemExit) as info:
            modulo.main()

        assert isinstance(info.value.code, str)
        assert "comparador/reproducir_e4.py" in info.value.code
        assert llamados == [] and not destino.exists()

    def test_la_cli_con_codigo_congelado_igual_sigue(
        self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        from comparador import reproducir_e4 as modulo
        from tests.test_persistence import copiar_src

        llamados: list[Path] = []

        def _calcular(run_dir: Path, destino: Path) -> dict[str, Any]:
            llamados.append(destino)
            return {"textos": 0, "run_origen": run_dir.name, "uso": None,
                    "modelo_embeddings": "m", "similitudes": {"comparadas": 0, "distintas": []}}

        monkeypatch.setattr(modulo, "calcular_vectores_del_run", _calcular)
        run = tmp_path / "run"
        run.mkdir()
        (run / "etapa3_pdf_hallazgos.json").write_text("[]", encoding="utf-8")
        monkeypatch.setattr(sys, "argv", [
            "reproducir_e4", str(run), "--solo-vectores", str(tmp_path / "x.npz"),
            "--codigo-congelado", str(copiar_src(tmp_path / "src"))])

        modulo.main()

        assert llamados == [tmp_path / "x.npz"]
        assert "Código congelado" in capsys.readouterr().out


class TestRepetirUnRunFallido:
    """Un run que falla en E4 o en E5 ya pagó el checklist, E1 y E3 (y quizás
    parte de E4): el pipeline guarda antes de E4 lo que la repetición necesita
    (config_e4.json: configuración de E4, empresa y metadatos del informe) y
    reproducir_e4 lo repite sin metadata.json ni reporte_final.json."""

    @staticmethod
    def _sin_resultado(run_dir: Path) -> None:
        for nombre in ("metadata.json", "reporte_final.json"):
            assert not (run_dir / nombre).exists(), nombre

    @staticmethod
    def _run_fallido(entorno: dict[str, Any]) -> Path:
        entorno["llm"].intentos_e3.clear()   # E3 falsa con estado: como en un run nuevo
        with pytest.raises(RuntimeError):
            _correr(entorno)
        (run_dir,) = [p for p in entorno["runs"].iterdir() if not (p / "metadata.json").exists()]
        return run_dir

    def test_el_pipeline_guarda_la_config_de_e4_antes_de_e4(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(settings, "effort_llm", "medium")
        TestFallosDelJuezPersistenLoPagado._falla_en(monkeypatch, "_juez_matching", 1)

        run_dir = self._run_fallido(entorno)

        self._sin_resultado(run_dir)
        config = _json(run_dir / "config_e4.json")
        assert config == {
            "empresa_auditora": "Ensaut",
            "modelo_extraccion": "claude-test", "modelo_juez_matching": None,
            "modelo_embeddings": settings.modelo_embeddings, "juez_redaccion": True,
            "asignacion_global": settings.asignacion_global,
            "descripcion_identica": settings.descripcion_identica,
            "control_codigos": settings.control_codigos,
            "unificar_copias": settings.unificar_copias,
            "juez_justifica_primero": settings.juez_justifica_primero, "top_k_candidatos": 5,
            "threshold_descripcion": 0.85, "threshold_accion": 0.85, "effort_llm": "medium",
            "metadatos_informe": config["metadatos_informe"],
        }
        assert config["metadatos_informe"]["coordinado"]

    def test_se_repite_e4_de_un_run_que_fallo_en_e4(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        """La caché mixta sirve la llamada que el run fallido sí pagó y paga el
        resto; el reporte es el de un run que terminó con los mismos insumos."""
        from comparador.reproducir_e4 import reproducir_e4

        completo, _, _ = _correr(entorno)
        TestFallosDelJuezPersistenLoPagado._falla_en(monkeypatch, "_juez_matching", 2)
        run_dir = self._run_fallido(entorno)
        self._sin_resultado(run_dir)
        monkeypatch.setattr(e4, "calcular_embeddings", _solo_la_prueba_de_la_clave)

        meta = reproducir_e4(run_dir, tmp_path / "replay", cache="mixto")

        repetido = ReporteInforme.model_validate(_json(tmp_path / "replay" / "reporte_final.json"))
        assert repetido.model_dump(exclude={"fecha_procesamiento"}) == completo.model_dump(
            exclude={"fecha_procesamiento"})
        assert meta["conteo_original"] is None
        assert meta["conteo_replay"]["coinciden"] == completo.resumen.coinciden
        assert (meta["cache_jueces"]["desde_cache"], meta["cache_jueces"]["pagadas"],
                meta["cache_jueces"]["excluidas_fallidas"]) == (1, 3, 1)
        assert meta["vectores"]["origen"] == "run" and meta["vectores"]["calculados"] == 0
        assert _json(tmp_path / "replay" / "metadata.json") == meta

    def test_se_repite_e4_sin_api_de_un_run_que_fallo_en_e5(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        from comparador import pipeline
        from comparador.reproducir_e4 import reproducir_e4

        completo, _, _ = _correr(entorno)

        def _e5_rota(*_: Any, **__: Any) -> Any:
            raise RuntimeError("E5 rota")

        monkeypatch.setattr(pipeline, "consolidar_reporte", _e5_rota)
        run_dir = self._run_fallido(entorno)
        self._sin_resultado(run_dir)
        TestVectoresYCacheE4._sin_api(monkeypatch)
        # Con GT del informe: sin reporte original no hay antes ni delta.
        shutil.copyfile(run_dir / "input_pdf.hash", run_dir / "input_informe_original.hash")
        TestAsignacionGlobal._gt_del_run(run_dir, completo, tmp_path / "gt", [(P1, "OK", "OK")])

        meta = reproducir_e4(run_dir, tmp_path / "replay", tmp_path / "gt", cache="estricto")

        repetido = ReporteInforme.model_validate(_json(tmp_path / "replay" / "reporte_final.json"))
        assert repetido.model_dump(exclude={"fecha_procesamiento"}) == completo.model_dump(
            exclude={"fecha_procesamiento"})
        assert (meta["cache_jueces"]["desde_cache"], meta["cache_jueces"]["pagadas"]) == (4, 0)
        gt = meta["comparador_gt"]
        assert (gt["original"], gt["delta"], gt["replay"]["OK"]) == (None, None, 1)

    def test_un_run_fallido_sin_config_e4_falla_antes_de_crear_la_salida(
        self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
    ) -> None:
        """Un run fallido anterior a config_e4.json: no se adivina la configuración."""
        from comparador.reproducir_e4 import reproducir_e4

        TestFallosDelJuezPersistenLoPagado._falla_en(monkeypatch, "_juez_matching", 1)
        run_dir = self._run_fallido(entorno)
        (run_dir / "config_e4.json").unlink()

        with pytest.raises(FileNotFoundError, match="config_e4.json"):
            reproducir_e4(run_dir, tmp_path / "replay", cache="mixto")
        assert not (tmp_path / "replay").exists()


def _arbol(raiz: Path) -> dict[Path, bytes | None]:
    """Cada ruta bajo `raiz` (carpetas en None) con su contenido: nada se escribió."""
    return {p.relative_to(raiz): (p.read_bytes() if p.is_file() else None)
            for p in raiz.rglob("*")}


class TestPreflightReproducirE4:
    """`reproducir_e4 --preflight` (PD23): todo lo que la repetición valida antes
    de gastar, sin llamar a ninguna API (tampoco la prueba de la clave de OpenAI,
    que es una llamada) y sin escribir nada. La medición tiene la forma de PD23:
    un run de B2 sin vectores (los de D1 aparte), la base de regresión hecha con
    la asignación global apagada, el GT del informe y la copia congelada del
    código fuera de replay_e4/."""

    @pytest.fixture
    def medicion(self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch,
                 tmp_path: Path) -> dict[str, Path]:
        from comparador.reproducir_e4 import reproducir_e4
        from tests.test_persistence import copiar_src

        # Como B2: el run y su base de regresión, con la asignación global apagada;
        # después vuelve la del proceso (cada test fija la suya o queda el default).
        proceso = settings.asignacion_global
        monkeypatch.setattr(settings, "asignacion_global", False)
        reporte, run_dir, _ = _correr(entorno)
        med = tmp_path / "mediciones"
        run = med / "b2" / "runs" / run_dir.name
        shutil.copytree(run_dir, run)
        shutil.copyfile(run / "input_pdf.hash", run / "input_informe_original.hash")
        npz = med / "b2" / "vectores_e4" / f"{run.name}.npz"
        npz.parent.mkdir(parents=True)
        (run / "etapa4_embeddings.npz").rename(npz)   # como B2: el run no guardó vectores
        gt = med / "gt" / "comparador"
        gt.parent.mkdir(parents=True)
        TestAsignacionGlobal._gt_del_run(run, reporte, gt, [(P1, "OK", "OK"), (P2, "A", "A4")])
        monkeypatch.setattr(e4, "calcular_embeddings", _solo_la_prueba_de_la_clave)
        base = med / "b2" / "replay_base" / run.name
        reproducir_e4(run, base, vectores=npz, cache="mixto")
        monkeypatch.setattr(settings, "asignacion_global", proceso)
        return {"raiz": tmp_path, "run": run, "npz": npz, "gt": gt, "base": base,
                "salida": med / "replay_e4" / "AG1" / run.name,
                "codigo": copiar_src(med / "replay_e4_codigo" / "AG1" / "src")}

    @staticmethod
    def _argv(m: dict[str, Path], **cambios: str) -> list[str]:
        opciones = {"--salida": str(m["salida"]), "--vectores": str(m["npz"]),
                    "--cache": "mixto", "--cache-desde": str(m["base"]),
                    "--gt-comparador": str(m["gt"]), "--codigo-congelado": str(m["codigo"]),
                    **cambios}
        return ["reproducir_e4", cambios.get("run", str(m["run"])),
                *[x for k, v in opciones.items() if k != "run" for x in (k, v)], "--preflight"]

    @staticmethod
    def _sin_llamadas(monkeypatch: pytest.MonkeyPatch) -> list[str]:
        llamadas: list[str] = []

        def _registrar(nombre: str) -> Callable[..., Any]:
            def _llamada(*_: Any, **__: Any) -> Any:
                llamadas.append(nombre)
                raise AssertionError(f"el preflight llamó a {nombre}")
            return _llamada

        monkeypatch.setattr(e4, "llamar_llm_estructurado", _registrar("juez"))
        monkeypatch.setattr(e4, "calcular_embeddings", _registrar("embeddings"))
        return llamadas

    def test_preflight_no_llama_a_nada_ni_escribe_nada(
        self, medicion: dict[str, Path], monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str]
    ) -> None:
        from comparador import reproducir_e4 as modulo

        monkeypatch.setattr(settings, "asignacion_global", True)
        llamadas = self._sin_llamadas(monkeypatch)
        antes = _arbol(medicion["raiz"])
        monkeypatch.setattr(sys, "argv", self._argv(medicion))

        modulo.main()

        assert llamadas == []
        assert _arbol(medicion["raiz"]) == antes
        assert not medicion["salida"].exists()
        salida = capsys.readouterr().out
        textos = len(e4.textos_a_embeber(*TestVectoresYCacheE4._hallazgos(medicion["run"])))
        for esperado in ("Código congelado", "asignación global: True",
                         f"vectores: {textos} de {textos} textos de E4",
                         f"caché mixto desde {medicion['base'].resolve()}",
                         "GT gt.json", "respaldo: (ninguno)", "Preflight OK"):
            assert esperado in salida, esperado
        assert "AVISO" not in salida

    def test_preflight_devuelve_lo_que_valido(
        self, medicion: dict[str, Path], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from comparador.reproducir_e4 import preflight

        monkeypatch.setattr(settings, "asignacion_global", True)
        monkeypatch.setattr(settings, "modelo_extraccion_fallback", "claude-respaldo")
        llamadas = self._sin_llamadas(monkeypatch)

        resumen = preflight(medicion["run"], medicion["salida"], medicion["gt"],
                            vectores=medicion["npz"], cache="mixto", cache_desde=medicion["base"])

        assert llamadas == [] and not medicion["salida"].exists()
        assert (resumen["run_origen"], resumen["modelo_extraccion"], resumen["asignacion_global"],
                resumen["modelo_extraccion_fallback"]) == (
            medicion["run"].name, "claude-test", True, "claude-respaldo")
        assert (resumen["asignacion_global_run"], resumen["aviso_asignacion_global"]) == (
            False, None)
        assert resumen["vectores"]["faltan"] == 0 and resumen["vectores"]["origen"] == "indicado"
        assert resumen["cache_jueces"]["modo"] == "mixto"
        assert resumen["cache_jueces"]["cargadas"] > 0
        assert resumen["credenciales"] == "presentes"
        assert resumen["comparador_gt"]["gt"] == "gt.json"
        # La configuración del run rige solo dentro del preflight.
        assert settings.modelo_extraccion == "claude-test"

    def test_preflight_estricto_no_exige_credenciales(
        self, medicion: dict[str, Path], monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str]
    ) -> None:
        from comparador import reproducir_e4 as modulo

        llamadas = self._sin_llamadas(monkeypatch)
        monkeypatch.setattr(settings, "openai_api_key", "")
        monkeypatch.setattr(settings, "anthropic_api_key", "")
        antes = _arbol(medicion["raiz"])
        monkeypatch.setattr(sys, "argv", self._argv(medicion, **{"--cache": "estricto"}))

        modulo.main()

        assert llamadas == [] and "Preflight OK" in capsys.readouterr().out
        # La repetición estricta tampoco llamaría a nada: el preflight no la corre.
        assert _arbol(medicion["raiz"]) == antes

    def test_preflight_avisa_si_los_vectores_no_cubren_el_run(
        self, medicion: dict[str, Path], monkeypatch: pytest.MonkeyPatch,
        capsys: pytest.CaptureFixture[str]
    ) -> None:
        """Un npz de otro run (mismo modelo) en mixto: E4 pagaría embeddings y, con
        otros vectores, la caché no encontraría los jueces. Se avisa, no se corta."""
        from comparador import reproducir_e4 as modulo
        from comparador.persistence import cargar_vectores, escribir_vectores

        vectores, modelo = cargar_vectores(medicion["npz"])
        parcial = medicion["raiz"] / "mediciones" / "b2" / "vectores_e4" / "parcial.npz"
        escribir_vectores(parcial, dict(list(vectores.items())[:2]), modelo)
        self._sin_llamadas(monkeypatch)
        monkeypatch.setattr(sys, "argv", self._argv(medicion, **{"--vectores": str(parcial)}))

        modulo.main()

        salida = capsys.readouterr().out
        assert f"vectores: 2 de {len(vectores)} textos de E4" in salida
        assert f"AVISO: faltan {len(vectores) - 2} vectores" in salida

    @pytest.mark.parametrize(("caso", "mensaje"), [
        ("npz-de-otro-modelo", "tiene vectores de otro-modelo"),
        ("npz-inexistente", "--vectores: no existe"),
        ("salida-con-contenido", "ya existe con contenido"),
        ("salida-no-ignorada-por-git", "git no la ignora"),
        ("codigo-distinto", "no es el de la copia congelada"),
        ("sin-credenciales", "ANTHROPIC_API_KEY"),
        ("gt-ambiguo", "más de un GT"),
        ("cache-desde-sin-raw", "no tiene raw de los jueces"),
        ("cache-desde-de-otro-run", "es de otro run (otro-run, no "),
        ("run-sin-e3", "no tiene etapa3_pdf_hallazgos.json"),
    ])
    def test_preflight_falla_sin_llamar_ni_escribir(
        self, medicion: dict[str, Path], monkeypatch: pytest.MonkeyPatch, caso: str,
        mensaje: str
    ) -> None:
        from comparador import reproducir_e4 as modulo
        from comparador.persistence import cargar_vectores, escribir_vectores
        from tests.test_persistence import repo_git

        cambios: dict[str, str] = {}
        raiz = medicion["raiz"]
        if caso == "npz-de-otro-modelo":
            otro = raiz / "otro.npz"
            escribir_vectores(otro, cargar_vectores(medicion["npz"])[0], "otro-modelo")
            cambios["--vectores"] = str(otro)
        elif caso == "npz-inexistente":
            cambios["--vectores"] = str(raiz / "no-existe.npz")
        elif caso == "salida-con-contenido":
            medicion["salida"].mkdir(parents=True)
            (medicion["salida"] / "judge_0.json").write_text("{}", encoding="utf-8")
        elif caso == "salida-no-ignorada-por-git":
            cambios["--salida"] = str(repo_git(raiz, "") / "replay")
        elif caso == "codigo-distinto":
            (medicion["codigo"] / "comparador" / "reproducir_e4.py").write_text("# otra\n")
        elif caso == "sin-credenciales":
            monkeypatch.setattr(settings, "anthropic_api_key", "")
        elif caso == "gt-ambiguo":
            shutil.copyfile(medicion["gt"] / "gt.json", medicion["gt"] / "otro.json")
        elif caso == "cache-desde-sin-raw":
            (raiz / "vacia").mkdir()
            cambios["--cache-desde"] = str(raiz / "vacia")
        elif caso == "cache-desde-de-otro-run":
            meta_base = _json(medicion["base"] / "metadata.json")
            (medicion["base"] / "metadata.json").write_text(
                json.dumps({**meta_base, "run_origen": "otro-run"}), encoding="utf-8")
        elif caso == "run-sin-e3":
            (medicion["run"] / "etapa3_pdf_hallazgos.json").unlink()
        llamadas = self._sin_llamadas(monkeypatch)
        antes = _arbol(raiz)
        monkeypatch.setattr(sys, "argv", self._argv(medicion, **cambios))

        with pytest.raises(SystemExit) as info:
            modulo.main()

        assert mensaje in str(info.value.code)
        assert llamadas == []
        assert _arbol(raiz) == antes

    def test_la_cli_guarda_el_codigo_congelado_y_el_respaldo_en_la_metadata(
        self, medicion: dict[str, Path], monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Lo que la repetición necesita para atribuir sus cifras queda en su
        metadata.json (PD38): la copia congelada contra la que se verificó el
        código y el modelo de respaldo que rigió."""
        from comparador import reproducir_e4 as modulo
        from comparador.persistence import hash_arbol

        monkeypatch.setattr(e4, "calcular_embeddings", _solo_la_prueba_de_la_clave)
        monkeypatch.setattr(sys, "argv", self._argv(medicion)[:-1])   # sin --preflight

        modulo.main()

        meta = _json(medicion["salida"] / "metadata.json")
        assert meta["codigo_congelado"] == {
            "carpeta": str(medicion["codigo"]),
            "sha256_codigo": hash_arbol(medicion["codigo"])[0][:16],
            "archivos_codigo": hash_arbol(medicion["codigo"])[1]}
        assert meta["codigo_congelado"]["sha256_codigo"] == meta["version_codigo"]["sha256_codigo"]
        assert meta["modelo_extraccion_fallback"] is None


@pytest.fixture(autouse=True)
def _sin_descripcion_identica(monkeypatch: pytest.MonkeyPatch) -> None:
    """Estos tests fijan el camino del juez con descripciones idénticas (y lo que
    reproducen de B2, AG1 y AG2, que corrieron sin M30). DESCRIPCION_IDENTICA,
    encendido por defecto desde el 2026-10-27, se prueba en
    test_etapa4_descripcion_identica.py."""
    monkeypatch.setattr(settings, "descripcion_identica", False)


class TestControlCodigos:
    def test_por_defecto_encendido_y_registrado_en_el_run(
            self, entorno: dict[str, Any], monkeypatch: pytest.MonkeyPatch) -> None:
        # M28 V2 (encendido por defecto desde el 2026-10-27): el interruptor queda en
        # la configuración de E4 y en la metadata, como los demás de E4.
        assert Settings.model_fields["control_codigos"].default is True
        monkeypatch.setattr(settings, "control_codigos", False)
        _, run_dir, _ = _correr(entorno)

        assert _json(run_dir / "metadata.json")["control_codigos"] is False
        assert _json(run_dir / "config_e4.json")["control_codigos"] is False


class TestDescripcionIdentica:
    def test_por_defecto_encendida_empareja_sin_juez(self, entorno: dict[str, Any],
                                                     monkeypatch: pytest.MonkeyPatch) -> None:
        # M30 (encendido por defecto desde el 2026-10-27): los dos hallazgos de BT1, con
        # la misma descripción que sus filas, salen sin juez; queda una sola consulta.
        assert Settings.model_fields["descripcion_identica"].default is True
        monkeypatch.setattr(settings, "descripcion_identica", True)
        reporte, run_dir, _ = _correr(entorno)

        assert _nombres(run_dir / "etapa4_llm_judge_raw") == ["embeddings.json", "judge_0.json"]
        bt1 = [r for r in reporte.resultados if r.rele == "SEL 551C"]
        assert [r.tipo_match for r in bt1] == [TipoMatch.DESCRIPCION_IDENTICA] * 2
        assert _json(run_dir / "metadata.json")["descripcion_identica"] is True
        assert _json(run_dir / "config_e4.json")["descripcion_identica"] is True
        assert reporte.resumen.total == 11
