"""Tests de la Etapa 0: checklist de aprobación (38 puntos). Sin LLM real."""

from datetime import datetime
from pathlib import Path
from typing import Any

import pytest

from comparador import etapa0_checklist
from comparador.config import settings
from comparador.etapa0_checklist import (
    TOTAL_PUNTOS,
    ChecklistConfigError,
    _ChecklistResponse,
    _ItemResponse,
    cargar_puntos,
    verificar_checklist,
)
from comparador.exportar import exportar_xlsx
from comparador.llm.prompts import user_checklist
from comparador.schemas import (
    ChecklistInforme,
    EmpresaAuditora,
    ItemChecklist,
    ReporteInforme,
    ResultadoChecklist,
    ResumenInforme,
)
from comparador.troceo import Tramo

_PLANILLAS_GT = sorted(
    (Path(__file__).resolve().parents[1] / "informes_reales").glob("*/*[Cc]heck[Ll]ist*.xlsx")
)


@pytest.fixture(autouse=True)
def sin_cita(monkeypatch: pytest.MonkeyPatch) -> None:
    """Estos tests fijan el checklist sin cita (CITA_CHECKLIST=false, el
    comportamiento previo exacto); el de la cita está en TestChecklistConCita."""
    from comparador.config import settings

    monkeypatch.setattr(settings, "cita_checklist", False)


def _respuesta_completa(resultado: ResultadoChecklist = ResultadoChecklist.CUMPLE):
    return _ChecklistResponse(items=[
        _ItemResponse(numero=n, resultado=resultado, paginas="1", referencia="portada")
        for n in range(1, TOTAL_PUNTOS + 1)
    ])


class TestCargarPuntos:
    def test_carga_exactamente_38_en_orden(self):
        puntos = cargar_puntos()
        assert len(puntos) == TOTAL_PUNTOS
        assert [p["numero"] for p in puntos] == list(range(1, 39))
        assert all(p["enunciado"].strip() for p in puntos)
        assert all(p["donde_buscar"].strip() for p in puntos)

    def test_enunciados_17_20_y_sin_espacios_sobrantes(self):
        por_n = {p["numero"]: str(p["enunciado"]) for p in cargar_puntos()}
        assert len(por_n) == 38
        assert "tablas" in por_n[17]
        assert "todos los casos en que" in por_n[20]
        assert "despeje de la falla es inferior a 400" in por_n[20]
        assert not [n for n, e in por_n.items() if "  " in e or e != e.strip()]

    @pytest.mark.skipif(not _PLANILLAS_GT, reason="planillas CEN privadas no disponibles")
    @pytest.mark.parametrize("planilla", _PLANILLAS_GT, ids=lambda p: p.parent.name[:20])
    def test_enunciados_literales_de_la_planilla_cen(self, planilla: Path):
        # 3, 4 y 5 desglosan el ítem 1.4: no son literales por diseño.
        import pandas as pd

        from comparador.evaluar_checklist import _COL_CODIGO, _COL_ENUNCIADO, cargar_mapeo

        df = pd.read_excel(planilla, sheet_name="Check List", header=None, keep_default_na=False)
        oficial: dict[str, str] = {}
        for codigo, texto in zip(df[_COL_CODIGO], df[_COL_ENUNCIADO], strict=True):
            codigo, texto = str(codigo).strip().rstrip("."), " ".join(str(texto).split())
            if codigo and texto:
                clave = codigo
                n = 2
                while clave in oficial:
                    clave = f"{codigo}#{n}"
                    n += 1
                oficial[clave] = texto
        mapeo = cargar_mapeo()
        distintos = [
            p["numero"] for p in cargar_puntos()
            if p["numero"] not in (3, 4, 5) and p["enunciado"] != oficial[mapeo[p["numero"]]]
        ]
        assert distintos == []

    def test_yaml_incompleto_falla_ruidosamente(self, tmp_path: Path):
        ruta = tmp_path / "malo.yaml"
        ruta.write_text("puntos:\n  - numero: 1\n    enunciado: x\n", encoding="utf-8")
        with pytest.raises(ChecklistConfigError, match="38"):
            cargar_puntos(ruta)


class TestPromptChecklist:
    def test_incluye_los_38_enunciados(self):
        puntos = cargar_puntos()
        prompt = user_checklist(puntos)
        for p in puntos:
            assert str(p["enunciado"])[:40] in prompt
        assert "38. " in prompt


@pytest.fixture
def un_tramo(monkeypatch: pytest.MonkeyPatch) -> None:
    """El informe cabe entero: un solo tramo, sin llamar a count_tokens."""
    monkeypatch.setattr(
        etapa0_checklist, "planificar_tramos",
        lambda ruta, **kw: [Tramo(1, 10, 10, "b64", 1000)],
    )


class TestVerificarChecklist:
    @pytest.mark.usefixtures("un_tramo")
    def test_consolida_y_cuenta(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        pdf = tmp_path / "x.pdf"
        pdf.write_bytes(b"%PDF-1.4 fake")

        resp = _respuesta_completa()
        resp.items[9].resultado = ResultadoChecklist.NO_CUMPLE
        resp.items[10].resultado = ResultadoChecklist.NO_VERIFICABLE

        monkeypatch.setattr(
            etapa0_checklist, "llamar_llm_estructurado",
            lambda **kw: (resp, {"tokens_total": 100}),
        )

        checklist, meta = verificar_checklist(pdf)
        assert checklist.total == 38
        assert checklist.cumplen == 36
        assert checklist.items[9].resultado == ResultadoChecklist.NO_CUMPLE
        assert checklist.items[0].enunciado.startswith(
            "El Informe de Auditoría identifica apropiadamente al Coordinado"
        )

    @pytest.mark.usefixtures("un_tramo")
    @pytest.mark.parametrize("batch", [False, True])
    def test_el_interruptor_de_batch_llega_al_cliente(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, batch: bool
    ):
        # M61: CHECKLIST_BATCH solo cambia por dónde viaja la llamada; el prompt no.
        pdf = tmp_path / "x.pdf"
        pdf.write_bytes(b"%PDF-1.4 fake")
        pedidos: list[dict[str, Any]] = []

        def _llm(**kw: Any) -> tuple[Any, dict[str, Any]]:
            pedidos.append(kw)
            return _respuesta_completa(), {}

        monkeypatch.setattr(etapa0_checklist, "llamar_llm_estructurado", _llm)
        monkeypatch.setattr(settings, "checklist_batch", batch)
        verificar_checklist(pdf)
        (pedido,) = pedidos
        assert pedido["batch"] is batch

    @pytest.mark.usefixtures("un_tramo")
    def test_respuesta_incompleta_falla_ruidosamente(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        pdf = tmp_path / "x.pdf"
        pdf.write_bytes(b"%PDF-1.4 fake")

        resp = _respuesta_completa()
        resp.items = resp.items[:35]  # faltan 36, 37, 38

        monkeypatch.setattr(
            etapa0_checklist, "llamar_llm_estructurado", lambda **kw: (resp, {}),
        )

        with pytest.raises(RuntimeError, match=r"\[36, 37, 38\]"):
            verificar_checklist(pdf)


class TestChecklistPorTramos:
    """Informe que no cabe en el contexto: tramos + fusión conservadora."""

    def _dos_tramos(self, monkeypatch: pytest.MonkeyPatch) -> list[dict[str, Any]]:
        monkeypatch.setattr(
            etapa0_checklist, "planificar_tramos",
            lambda ruta, **kw: [Tramo(1, 228, 456, "t1", 580_000),
                                Tramo(229, 456, 456, "t2", 460_000)],
        )
        llamadas: list[dict[str, Any]] = []

        def _llm(**kw: Any) -> tuple[_ChecklistResponse, dict[str, Any]]:
            llamadas.append(kw)
            primero = kw["pdf_base64"] == "t1"
            items = []
            for n in range(1, TOTAL_PUNTOS + 1):
                if n == 6:  # DUF en el anexo final: solo el tramo 2 lo ve
                    r = ResultadoChecklist.NO_VERIFICABLE if primero else ResultadoChecklist.CUMPLE
                elif n == 23:  # despeje > 400 ms visto en el tramo 1
                    r = ResultadoChecklist.NO_CUMPLE if primero else ResultadoChecklist.CUMPLE
                elif n == 30:  # en ningún tramo
                    r = ResultadoChecklist.NO_VERIFICABLE
                else:
                    r = ResultadoChecklist.CUMPLE
                items.append(_ItemResponse(numero=n, resultado=r, paginas="2, 5-6",
                                           referencia="t1" if primero else "t2"))
            return _ChecklistResponse(items=items), {"tokens_total": 10}

        monkeypatch.setattr(etapa0_checklist, "llamar_llm_estructurado", _llm)
        return llamadas

    def test_fusion_no_cumple_gana_luego_cumple(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        self._dos_tramos(monkeypatch)
        checklist, meta = verificar_checklist(tmp_path / "x.pdf")
        por_n = {i.numero: i for i in checklist.items}
        assert por_n[6].resultado == ResultadoChecklist.CUMPLE
        assert por_n[23].resultado == ResultadoChecklist.NO_CUMPLE
        assert por_n[30].resultado == ResultadoChecklist.NO_VERIFICABLE
        assert meta["troceo"]["tramos"][1]["paginas"] == [229, 456]
        assert meta["tokens_total"] == 20

    def test_paginas_del_tramo_se_vuelven_absolutas(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        self._dos_tramos(monkeypatch)
        checklist, _ = verificar_checklist(tmp_path / "x.pdf")
        punto6 = next(i for i in checklist.items if i.numero == 6)
        assert punto6.referencia == "t2"
        assert punto6.paginas == "230, 233-234"  # 2 y 5-6 del tramo 229-456

    def test_prompt_de_tramo_avisa_y_cambia_regla_de_ausencia(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        llamadas = self._dos_tramos(monkeypatch)
        verificar_checklist(tmp_path / "x.pdf")
        assert "páginas 229 a 456 de un informe de 456" in llamadas[1]["user_prompt"]
        assert "REGLAS PARA UN TRAMO" in llamadas[1]["user_prompt"]


def _reporte(checklist: ChecklistInforme | None) -> ReporteInforme:
    return ReporteInforme(
        empresa_auditora=EmpresaAuditora.ENSAUT,
        coordinado="COORDINADO C", subestacion="Epsilon", informe_id="EN-1",
        fecha_auditoria=None, fecha_procesamiento=datetime(2026, 7, 7),
        modelo_extraccion="claude-haiku-4-5", modelo_embeddings="emb",
        resumen=ResumenInforme(
            total=0, coinciden=0, difieren=0,
            faltan_en_excel=0, extra_en_excel=0, requieren_revision=0,
        ),
        resultados=[],
        checklist=checklist,
    )


class TestChecklistEnReporte:
    def test_roundtrip_con_checklist(self):
        cl = ChecklistInforme(
            items=[ItemChecklist(
                numero=1, enunciado="x", resultado=ResultadoChecklist.CUMPLE,
                paginas="1", referencia="portada",
            )],
            cumplen=1, total=38,
        )
        data: dict[str, Any] = _reporte(cl).model_dump(mode="json")
        restaurado = ReporteInforme.model_validate(data)
        assert restaurado.checklist is not None
        assert restaurado.checklist.cumplen == 1

    def test_retrocompatible_sin_checklist(self):
        data: dict[str, Any] = _reporte(None).model_dump(mode="json")
        del data["checklist"]  # JSON de un run previo a la feature
        restaurado = ReporteInforme.model_validate(data)
        assert restaurado.checklist is None

    def test_export_xlsx_incluye_hoja(self):
        import io

        from openpyxl import load_workbook

        cl = ChecklistInforme(
            items=[
                ItemChecklist(numero=1, enunciado="Punto uno",
                              resultado=ResultadoChecklist.CUMPLE, paginas="1"),
                ItemChecklist(numero=2, enunciado="Punto dos",
                              resultado=ResultadoChecklist.NO_CUMPLE),
            ],
            cumplen=1, total=38,
        )
        contenido = exportar_xlsx(_reporte(cl))
        wb = load_workbook(io.BytesIO(contenido))
        assert "Checklist" in wb.sheetnames
        ws = wb["Checklist"]
        assert ws.cell(row=2, column=3).value == "Cumple"
        assert ws.cell(row=3, column=3).value == "No cumple"

    def test_export_xlsx_sin_checklist_no_agrega_hoja(self):
        import io

        from openpyxl import load_workbook

        contenido = exportar_xlsx(_reporte(None))
        wb = load_workbook(io.BytesIO(contenido))
        assert "Checklist" not in wb.sheetnames


class TestFeedback:
    def test_export_feedback_md(self):
        from comparador.exportar import exportar_feedback_md
        from comparador.schemas import FeedbackRun

        cl = ChecklistInforme(
            items=[ItemChecklist(numero=1, enunciado="Punto uno",
                                 resultado=ResultadoChecklist.CUMPLE)],
            cumplen=1, total=38,
        )
        fb = FeedbackRun(
            veredictos_checklist={"1": "incorrecto"},
            comentarios="El punto 1 en realidad no cumple.",
        )
        md = exportar_feedback_md(_reporte(cl), fb, "run-x")
        assert "# Feedback de revisión — run-x" in md
        assert "✗ INCORRECTO" in md
        assert "checklist punto 1" in md
        assert "El punto 1 en realidad no cumple." in md


class TestFormaDelResultado:
    """Si el modelo varía la forma del valor, se tolera la forma, nunca el sentido."""

    @pytest.mark.parametrize(("crudo", "esperado"), [
        ("Cumple", ResultadoChecklist.CUMPLE),
        (" NO_VERIFICABLE ", ResultadoChecklist.NO_VERIFICABLE),
        ("No Cumple", ResultadoChecklist.NO_CUMPLE),
        ("no-cumple", ResultadoChecklist.NO_CUMPLE),
        ("no_ cumple", ResultadoChecklist.NO_CUMPLE),
        ("no_cumple", ResultadoChecklist.NO_CUMPLE),
    ])
    def test_normaliza_solo_la_forma(self, crudo: str, esperado: ResultadoChecklist) -> None:
        assert _ItemResponse(numero=1, resultado=crudo).resultado == esperado

    @pytest.mark.parametrize(
        "crudo", ["sí", "ok", "cumple parcialmente", "Cumple.", "nocumple", ""]
    )
    def test_un_valor_distinto_sigue_fallando(self, crudo: str) -> None:
        # Nunca por subcadena: "cumple parcialmente" o "nocumple" no se convierten en nada.
        with pytest.raises(ValueError):
            _ItemResponse(numero=1, resultado=crudo)

    def test_el_schema_enviado_al_llm_sigue_siendo_el_enum(self) -> None:
        schema = _ChecklistResponse.model_json_schema()
        enum = schema["$defs"]["ResultadoChecklist"]["enum"]
        assert sorted(enum) == ["cumple", "no_cumple", "no_verificable"]

    def test_avisa_cuando_normaliza(self, caplog: pytest.LogCaptureFixture) -> None:
        with caplog.at_level("WARNING", logger="comparador.etapa0_checklist"):
            _ItemResponse(numero=1, resultado="Cumple")
            _ItemResponse(numero=2, resultado="cumple")
        assert len(caplog.records) == 1 and "'Cumple'" in caplog.records[0].getMessage()


class TestChecklistConCita:
    """CITA_CHECKLIST=true (por defecto, decisión del usuario 2026-10-18): el modelo
    copia una cita ANTES del resultado; el código la verifica contra el PDF y un
    "cumple" sin cita verificada baja a no_verificable, con el motivo a la vista."""

    PORTADA = "Informe de auditoría técnica de protecciones eléctricas, subestación Prueba"
    TABLA = "Tabla 4-1: resultados de las pruebas de inyección secundaria del paño C1"

    @pytest.fixture(autouse=True)
    def con_cita(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from comparador.config import settings

        monkeypatch.setattr(settings, "cita_checklist", True)

    def _pdf(self, tmp_path: Path) -> Path:
        """Pág. 1 portada, pág. 2 tabla de pruebas, pág. 3 escaneada (sin texto)."""
        import fitz

        doc = fitz.open()
        for lineas in [[self.PORTADA], ["Capítulo 4", self.TABLA], []]:
            pagina = doc.new_page()
            for i, linea in enumerate(lineas):
                pagina.insert_text((40, 200 + 14 * i), linea, fontsize=9)
        ruta = tmp_path / "informe.pdf"
        doc.save(ruta)
        return ruta

    def _correr(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                citas: dict[int, tuple[str, dict[str, Any] | None]]):
        from comparador.etapa0_checklist import _ChecklistConCitaResponse, _ItemConCitaResponse

        monkeypatch.setattr(etapa0_checklist, "planificar_tramos",
                            lambda ruta, **kw: [Tramo(1, 3, 3, "b64", 1000)])
        llamadas: list[dict[str, Any]] = []

        def _llm(**kw: Any) -> tuple[Any, dict[str, Any]]:
            llamadas.append(kw)
            items = [
                _ItemConCitaResponse(
                    numero=n, cita=citas.get(n, ("cumple", None))[1],
                    resultado=citas.get(n, ("no_verificable", None))[0],
                    paginas="1", referencia="r",
                )
                for n in range(1, TOTAL_PUNTOS + 1)
            ]
            return _ChecklistConCitaResponse(items=items), {"tokens_total": 10}

        monkeypatch.setattr(etapa0_checklist, "llamar_llm_estructurado", _llm)
        checklist, meta = verificar_checklist(self._pdf(tmp_path))
        return checklist, meta, llamadas

    def test_prompt_y_schema_piden_la_cita_antes_del_resultado(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        from comparador.etapa0_checklist import _ChecklistConCitaResponse, _ItemConCitaResponse
        from comparador.llm.prompts import REGLA_CITA_CHECKLIST

        _, _, (llamada,) = self._correr(tmp_path, monkeypatch, {})
        assert llamada["response_model"] is _ChecklistConCitaResponse
        assert REGLA_CITA_CHECKLIST in llamada["system_prompt"]
        assert "numero, cita (texto literal" in llamada["user_prompt"]
        schema = _ItemConCitaResponse.model_json_schema()
        assert list(schema["properties"])[:3] == ["numero", "cita", "resultado"]
        assert "cita" in schema["required"]   # obligatoria, aunque puede ser null

    def test_verifica_las_citas_y_baja_los_cumple_sin_respaldo(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        checklist, meta, _ = self._correr(tmp_path, monkeypatch, {
            1: ("cumple", {"pagina": 1, "texto": self.PORTADA}),         # verificada
            2: ("cumple", {"pagina": 2, "texto": "Tabla 9-9: no existe"}),  # no aparece
            3: ("cumple", None),                                          # sin cita
            4: ("cumple", {"pagina": 1, "texto": self.TABLA}),           # está en la p. 2
            5: ("no_cumple", None),                                       # ausencia: no se toca
            6: ("cumple", {"pagina": 3, "texto": self.PORTADA}),         # p. 3 escaneada
        })
        por_n = {i.numero: i for i in checklist.items}

        assert por_n[1].resultado == ResultadoChecklist.CUMPLE
        assert por_n[1].resultado_modelo is None
        assert por_n[1].citas[0].verificacion.estado == "verificada"  # type: ignore[union-attr]
        for n in (2, 3, 6):
            assert por_n[n].resultado == ResultadoChecklist.NO_VERIFICABLE
            assert por_n[n].resultado_modelo == ResultadoChecklist.CUMPLE
        assert "no aparece en la p. 2" in (por_n[2].motivo_ajuste or "")
        assert "sin una cita literal" in (por_n[3].motivo_ajuste or "")
        assert "la p. 3 no tiene texto" in (por_n[6].motivo_ajuste or "")
        # En otra página: se corrige la página (la citada queda) y sigue cumpliendo.
        assert por_n[4].resultado == ResultadoChecklist.CUMPLE
        (cita_4,) = por_n[4].citas
        assert cita_4.pagina == 2 and cita_4.verificacion.pagina_citada == 1  # type: ignore[union-attr]
        assert por_n[5].resultado == ResultadoChecklist.NO_CUMPLE
        assert por_n[5].resultado_modelo is None and por_n[5].citas == []
        assert (checklist.cumplen, checklist.bajados_sin_cita_verificada) == (2, 3)
        assert "errores_capa_texto" not in meta
        # 0.8.0: la causa de cada bajada; un informe leído entero no lleva nada de M08.
        from comparador.schemas import CausaAjuste

        assert {n for n, i in por_n.items() if i.causa_ajuste} == {2, 3, 6}
        assert {i.causa_ajuste for i in checklist.items} == {None, CausaAjuste.SIN_CITA_VERIFICADA}
        assert all(i.tramos_respaldo is None and i.aviso_tramos is None for i in checklist.items)
        assert (checklist.tramos_leidos, checklist.marcados_un_tramo,
                checklist.bajados_un_tramo) == (None, None, None)

    def test_cita_corta_o_en_varias_paginas_no_respalda(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        # Revisión 2026-10-18: «Correcta» o «5» aparecen en cualquier informe; una
        # cita que está en varias páginas (no en la citada) no tiene página que mostrar.
        checklist, _, _ = self._correr(tmp_path, monkeypatch, {
            1: ("cumple", {"pagina": 2, "texto": "Capítulo 4"}),          # 2 palabras
            2: ("cumple", {"pagina": 2, "texto": "Informe de auditoría técnica"}),
        })
        por_n = {i.numero: i for i in checklist.items}
        assert por_n[1].resultado == ResultadoChecklist.NO_VERIFICABLE
        assert "demasiado corta" in (por_n[1].motivo_ajuste or "")
        # "Informe de auditoría técnica" está en la p. 1 (portada) y no en la 2: en otra
        # página, una sola → se corrige y respalda.
        (cita_2,) = por_n[2].citas
        assert (cita_2.pagina, por_n[2].resultado) == (1, ResultadoChecklist.CUMPLE)

    def test_verifica_contra_el_pdf_con_texto(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        # El modelo leyó el adecuado (sin capa de texto): se verifica contra el original.
        from comparador.etapa0_checklist import _ChecklistConCitaResponse, _ItemConCitaResponse

        original = self._pdf(tmp_path)
        monkeypatch.setattr(etapa0_checklist, "planificar_tramos",
                            lambda ruta, **kw: [Tramo(1, 3, 3, "b64", 1000)])
        monkeypatch.setattr(etapa0_checklist, "llamar_llm_estructurado", lambda **kw: (
            _ChecklistConCitaResponse(items=[
                _ItemConCitaResponse(numero=n, resultado="cumple", cita={
                    "pagina": 1, "texto": self.PORTADA}) for n in range(1, TOTAL_PUNTOS + 1)
            ]), {}))
        checklist, _ = verificar_checklist(tmp_path / "adecuado.pdf", original)
        assert checklist.cumplen == 38 and checklist.bajados_sin_cita_verificada == 0

    def test_pdf_que_no_abre_falla_con_la_respuesta_pagada(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ):
        from comparador.etapa0_checklist import _ChecklistConCitaResponse, _ItemConCitaResponse

        monkeypatch.setattr(etapa0_checklist, "planificar_tramos",
                            lambda ruta, **kw: [Tramo(1, 3, 3, "b64", 1000)])
        monkeypatch.setattr(etapa0_checklist, "llamar_llm_estructurado", lambda **kw: (
            _ChecklistConCitaResponse(items=[
                _ItemConCitaResponse(numero=n, resultado="cumple", cita=None)
                for n in range(1, TOTAL_PUNTOS + 1)
            ]), {"tokens_total": 7}))
        with pytest.raises(Exception) as info:  # noqa: B017 — el error de fitz al abrir
            verificar_checklist(tmp_path / "no_existe.pdf")
        assert info.value.llm_metadata == {"tokens_total": 7}  # type: ignore[attr-defined]


class TestFusionDeTramos:
    """La fusión no pierde evidencia: une páginas, referencias y citas de los
    tramos con el resultado ganador (antes quedaban solo las del primero)."""

    @staticmethod
    def _item(resultado: str, paginas: str, referencia: str, cita: str | None):
        from comparador.etapa0_checklist import _ItemConCitaResponse

        return _ItemConCitaResponse(
            numero=6, resultado=resultado, paginas=paginas, referencia=referencia,
            cita={"pagina": int(paginas.split(",")[0]), "texto": cita} if cita else None,
        )

    def test_empate_une_la_evidencia(self):
        from comparador.etapa0_checklist import _fusionar

        punto = {"numero": 6, "enunciado": "DUF"}
        item = _fusionar([
            {6: self._item("cumple", "12", "Figura 2-1", "Diagrama unilineal general")},
            {6: self._item("cumple", "240, 12", "Anexo A", "Diagrama unilineal detallado")},
        ], punto)
        assert item.resultado == ResultadoChecklist.CUMPLE
        assert item.paginas == "12, 240, 12"
        assert item.referencia == "Figura 2-1 · Anexo A"
        assert [c.texto for c in item.citas] == [
            "Diagrama unilineal general", "Diagrama unilineal detallado"]

    def test_gana_no_cumple_con_su_evidencia(self):
        from comparador.etapa0_checklist import _fusionar

        item = _fusionar([
            {6: self._item("cumple", "12", "Figura 2-1", "Diagrama unilineal general")},
            {6: self._item("no_cumple", "300", "Tabla 9", "Prueba pendiente")},
        ], {"numero": 6, "enunciado": "DUF"})
        assert (item.resultado, item.paginas, item.referencia) == (
            ResultadoChecklist.NO_CUMPLE, "300", "Tabla 9")
        assert [c.texto for c in item.citas] == ["Prueba pendiente"]


class TestCitaApagada:
    def test_prompts_identicos_a_los_previos(self):
        from comparador.llm.prompts import SYSTEM_CHECKLIST, system_checklist

        puntos = cargar_puntos()
        assert system_checklist(False) == SYSTEM_CHECKLIST
        assert user_checklist(puntos, False) == user_checklist(puntos)
        assert user_checklist(puntos).endswith(
            "Para cada punto entrega: numero, resultado (cumple / no_cumple / "
            "no_verificable), paginas (si aplica), referencia breve (si aplica).\n")
        assert "cita (texto literal" not in user_checklist(puntos)

    @pytest.mark.usefixtures("un_tramo")
    def test_salida_sin_citas(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        llamadas: list[dict[str, Any]] = []

        def _llm(**kw: Any) -> tuple[Any, dict[str, Any]]:
            llamadas.append(kw)
            return _respuesta_completa(), {}

        monkeypatch.setattr(etapa0_checklist, "llamar_llm_estructurado", _llm)
        checklist, _ = verificar_checklist(tmp_path / "x.pdf")
        assert llamadas[0]["response_model"] is _ChecklistResponse
        assert checklist.bajados_sin_cita_verificada is None
        assert all(i.citas == [] and i.resultado_modelo is None for i in checklist.items)


@pytest.mark.parametrize(("texto", "estado", "respalda"), [
    ("Tabla 9: resultados de las pruebas", "verificada", True),
    ("Figura 6-3: DUF general", "verificada", True),             # justo 4 palabras
    ("Capítulo 4", "verificada", False),                         # demasiado corta
    ("Tabla 9: resultados de las pruebas", "no_encontrada", False),
    ("Tabla 9: resultados de las pruebas", None, False),         # sin verificar
])
def test_cita_respalda_es_la_regla_unica_de_la_bajada(
    texto: str, estado: str | None, respalda: bool
):
    """Una sola regla (etapa0.cita_respalda) para bajar un «cumple» y para contarlo
    después en las métricas (evaluar_lote.resumen_citas_checklist, PD52)."""
    from comparador.etapa0_checklist import cita_respalda
    from comparador.schemas import CitaChecklist, EstadoCita, VerificacionCita

    verificacion = None if estado is None else VerificacionCita(
        estado=EstadoCita(estado), pagina=1)
    assert cita_respalda(CitaChecklist(pagina=1, texto=texto,
                                       verificacion=verificacion)) is respalda


# --- M08: informes leídos por partes (schema 0.8.0, decisión del usuario 2026-10-24) ---

DOS_PARTES = [(1, 228), (229, 456)]
TRES_PARTES = [(1, 150), (151, 300), (301, 456)]
AVISO_FIN = " no lo respaldan: revísalas antes de darlo por cumplido."
MOTIVO_PD44 = ("El informe se leyó en partes por su longitud; este punto exige revisar todo el "
               "informe: revisa las páginas ")
PORTADA = "Informe de auditoría técnica de protecciones eléctricas, subestación Prueba"
TABLA = "Tabla 4-1: resultados de las pruebas de inyección secundaria del paño C1"

# Resultado de un punto en un tramo: "cumple", o ("cumple", cita) con la cita en
# la pág. 1 del tramo (solo con cita).
_Valor = str | tuple[str, str | None]


def _pdf_con_texto(ruta: Path, paginas: int, textos: dict[int, str]) -> None:
    import fitz

    doc = fitz.open()
    for p in range(1, paginas + 1):
        pagina = doc.new_page()
        if p in textos:
            pagina.insert_text((40, 200), textos[p], fontsize=9)
    doc.save(ruta)


def _por_partes(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, rangos: list[tuple[int, int]],
    resultados: dict[int, tuple[_Valor, ...]], *, modo: str = "marcar", con_cita: bool = False,
    textos: dict[int, str] | None = None,
) -> tuple[ChecklistInforme, dict[str, Any], list[dict[str, Any]]]:
    """verificar_checklist con un tramo por rango. resultados[n] = el valor del punto n
    en cada tramo; los puntos que no están, no_verificable en todos. La metadata de
    cada llamada lleva su response_raw, como la real (páginas relativas al tramo)."""
    from comparador.config import settings
    from comparador.etapa0_checklist import _ChecklistConCitaResponse, _ItemConCitaResponse

    monkeypatch.setattr(settings, "checklist_un_tramo", modo)
    monkeypatch.setattr(settings, "cita_checklist", con_cita)
    total = rangos[-1][1]
    monkeypatch.setattr(etapa0_checklist, "planificar_tramos", lambda ruta, **kw: [
        Tramo(i, f, total, f"t{k}", 1000) for k, (i, f) in enumerate(rangos)])
    llamadas: list[dict[str, Any]] = []

    def _item(n: int, valor: _Valor, k: int) -> Any:
        resultado, cita = (valor, None) if isinstance(valor, str) else valor
        comunes: dict[str, Any] = {"numero": n, "resultado": resultado, "paginas": "1",
                                   "referencia": f"t{k}"}
        if con_cita:
            return _ItemConCitaResponse(
                **comunes, cita={"pagina": 1, "texto": cita} if cita else None)
        return _ItemResponse(**comunes)

    def _llm(**kw: Any) -> tuple[Any, dict[str, Any]]:
        llamadas.append(kw)
        k = int(kw["pdf_base64"].removeprefix("t"))
        items = [_item(n, resultados.get(n, ("no_verificable",) * len(rangos))[k], k)
                 for n in range(1, TOTAL_PUNTOS + 1)]
        resp = (_ChecklistConCitaResponse if con_cita else _ChecklistResponse)(items=items)
        return resp, {"tokens_total": 10, "response_raw": resp.model_dump(mode="json")}

    monkeypatch.setattr(etapa0_checklist, "llamar_llm_estructurado", _llm)
    pdf = tmp_path / "x.pdf"
    if con_cita:
        _pdf_con_texto(pdf, total, textos or {})
    checklist, meta = verificar_checklist(pdf)
    return checklist, meta, llamadas


class TestLecturaPorPartes:
    """M08: en un informe leído por partes, un «cumple» que respaldan solo algunas
    partes lleva el aviso con las páginas que no lo respaldan; con
    CHECKLIST_UN_TRAMO=bajar, los puntos que exigen revisar todo el informe (PD44)
    bajan además a no verificable. Va después de la cita y no toca el prompt."""

    def test_lista_fija_de_14(self) -> None:
        lista = etapa0_checklist.PUNTOS_EXIGEN_TODO_EL_INFORME
        assert lista == frozenset({7, 8, 15, 17, 18, 19, 20, 21, 28, 30, 31, 32, 33, 34})

    def test_dos_partes_marca_sin_cambiar_el_resultado(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        cl, _, _ = _por_partes(monkeypatch, tmp_path, DOS_PARTES, {
            1: ("cumple", "cumple"),            # las dos partes lo respaldan
            7: ("cumple", "no_verificable"),     # en la lista, solo la primera
            6: ("no_verificable", "cumple"),     # fuera de la lista, solo la segunda
            8: ("no_cumple", "cumple"),          # gana no_cumple: no es un «cumple»
        })
        por_n = {i.numero: i for i in cl.items}
        assert por_n[7].aviso_tramos == (
            "Visto en 1 de 2 partes del informe (págs. 1-228); las págs. 229-456" + AVISO_FIN)
        assert por_n[6].aviso_tramos == (
            "Visto en 1 de 2 partes del informe (págs. 229-456); las págs. 1-228" + AVISO_FIN)
        for n in (6, 7):
            assert (por_n[n].resultado, por_n[n].resultado_modelo, por_n[n].motivo_ajuste,
                    por_n[n].causa_ajuste) == (ResultadoChecklist.CUMPLE, None, None, None)
        assert (por_n[7].tramos_respaldo, por_n[6].tramos_respaldo) == ([1], [2])
        assert (por_n[1].tramos_respaldo, por_n[1].aviso_tramos) == ([1, 2], None)
        assert (por_n[8].resultado, por_n[8].tramos_respaldo, por_n[8].aviso_tramos) == (
            ResultadoChecklist.NO_CUMPLE, [1], None)
        assert (por_n[30].tramos_respaldo, por_n[30].aviso_tramos) == ([1, 2], None)
        assert cl.tramos_leidos == DOS_PARTES
        assert (cl.cumplen, cl.marcados_un_tramo, cl.bajados_un_tramo) == (3, 2, None)
        assert cl.bajados_sin_cita_verificada is None

    def test_dos_partes_bajar_solo_los_puntos_de_la_lista(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from comparador.schemas import CausaAjuste

        cl, _, _ = _por_partes(monkeypatch, tmp_path, DOS_PARTES, {
            1: ("cumple", "cumple"),
            7: ("cumple", "no_verificable"),
            30: ("no_verificable", "cumple"),
            6: ("no_verificable", "cumple"),     # fuera de la lista: se marca, no baja
            8: ("no_cumple", "cumple"),          # en la lista, gana no_cumple: no se toca
            15: ("no_verificable", "no_verificable"),
        }, modo="bajar")
        por_n = {i.numero: i for i in cl.items}
        for n, revisar in ((7, "229-456"), (30, "1-228")):
            item = por_n[n]
            assert (item.resultado, item.resultado_modelo, item.causa_ajuste) == (
                ResultadoChecklist.NO_VERIFICABLE, ResultadoChecklist.CUMPLE, CausaAjuste.UN_TRAMO)
            assert item.motivo_ajuste == MOTIVO_PD44 + f"{revisar}."
            assert item.aviso_tramos is not None and f"las págs. {revisar} no" in item.aviso_tramos
        assert (por_n[6].resultado, por_n[6].resultado_modelo, por_n[6].causa_ajuste) == (
            ResultadoChecklist.CUMPLE, None, None)
        assert por_n[6].aviso_tramos is not None
        for n in (8, 15):
            assert (por_n[n].resultado_modelo, por_n[n].motivo_ajuste, por_n[n].causa_ajuste,
                    por_n[n].aviso_tramos) == (None, None, None, None), n
        assert por_n[8].resultado == ResultadoChecklist.NO_CUMPLE
        assert por_n[15].resultado == ResultadoChecklist.NO_VERIFICABLE
        assert (por_n[1].resultado, por_n[1].resultado_modelo) == (ResultadoChecklist.CUMPLE, None)
        assert (cl.cumplen, cl.marcados_un_tramo, cl.bajados_un_tramo) == (2, 1, 2)

    @pytest.mark.parametrize("modo", ["marcar", "bajar"])
    def test_tres_partes(
        self, modo: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        cl, _, _ = _por_partes(monkeypatch, tmp_path, TRES_PARTES, {
            1: ("cumple", "cumple", "cumple"),
            15: ("cumple", "no_verificable", "cumple"),           # 2 de 3: sigue siendo parte
            17: ("no_verificable", "cumple", "no_verificable"),   # 1 de 3, la del medio
            18: ("cumple", "no_verificable", "no_verificable"),   # las otras dos, contiguas
        }, modo=modo)
        por_n = {i.numero: i for i in cl.items}
        assert por_n[15].aviso_tramos == ("Visto en 2 de 3 partes del informe (págs. 1-150 y "
                                          "301-456); las págs. 151-300" + AVISO_FIN)
        assert por_n[17].aviso_tramos == ("Visto en 1 de 3 partes del informe (págs. 151-300); "
                                          "las págs. 1-150 y 301-456" + AVISO_FIN)
        assert por_n[18].aviso_tramos == ("Visto en 1 de 3 partes del informe (págs. 1-150); "
                                          "las págs. 151-456" + AVISO_FIN)
        assert [por_n[n].tramos_respaldo for n in (1, 15, 17, 18)] == [
            [1, 2, 3], [1, 3], [2], [1]]
        assert por_n[1].aviso_tramos is None
        assert cl.tramos_leidos == TRES_PARTES
        if modo == "marcar":
            assert all(por_n[n].resultado == ResultadoChecklist.CUMPLE for n in (15, 17, 18))
            assert (cl.cumplen, cl.marcados_un_tramo, cl.bajados_un_tramo) == (4, 3, None)
        else:
            assert [por_n[n].motivo_ajuste for n in (15, 17, 18)] == [
                MOTIVO_PD44 + "151-300.", MOTIVO_PD44 + "1-150 y 301-456.",
                MOTIVO_PD44 + "151-456."]
            assert (cl.cumplen, cl.marcados_un_tramo, cl.bajados_un_tramo) == (1, 0, 3)

    @pytest.mark.parametrize("modo", ["marcar", "bajar"])
    def test_un_tramo_no_lleva_nada(
        self, modo: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        cl, _, _ = _por_partes(monkeypatch, tmp_path, [(1, 10)], {
            n: ("cumple",) for n in range(1, TOTAL_PUNTOS + 1)}, modo=modo)
        assert all((i.tramos_respaldo, i.aviso_tramos, i.causa_ajuste, i.resultado_modelo,
                    i.motivo_ajuste) == (None,) * 5 for i in cl.items)
        assert (cl.tramos_leidos, cl.marcados_un_tramo, cl.bajados_un_tramo) == (None,) * 3
        assert cl.cumplen == TOTAL_PUNTOS

    @pytest.mark.parametrize("modo", ["marcar", "bajar"])
    def test_con_cita_va_despues_de_la_cita(
        self, modo: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """La bajada por cita conserva su motivo y su causa y lleva el aviso, sin
        contarse dos veces; lo que sigue en «cumple» tras la cita es lo que M08 marca
        o baja, y nunca entra al contador de la cita (control de PD52)."""
        from comparador.evaluar_lote import resumen_citas_checklist
        from comparador.schemas import CausaAjuste

        cl, _, _ = _por_partes(monkeypatch, tmp_path, [(1, 2), (3, 4)], {
            6: (("cumple", PORTADA), ("cumple", TABLA)),                      # las dos partes
            7: (("cumple", PORTADA), "no_verificable"),                       # cita ok, 1 de 2
            15: ("no_verificable", ("cumple", TABLA)),                        # cita ok, 1 de 2
            8: (("cumple", "Texto que el informe no tiene"), "no_verificable"),  # sin respaldo
        }, modo=modo, con_cita=True, textos={1: PORTADA, 3: TABLA})
        por_n = {i.numero: i for i in cl.items}

        bajado = por_n[8]
        assert (bajado.resultado, bajado.resultado_modelo, bajado.causa_ajuste) == (
            ResultadoChecklist.NO_VERIFICABLE, ResultadoChecklist.CUMPLE,
            CausaAjuste.SIN_CITA_VERIFICADA)
        assert "no se pudo verificar" in (bajado.motivo_ajuste or "")
        assert bajado.aviso_tramos == (
            "Visto en 1 de 2 partes del informe (págs. 1-2); las págs. 3-4" + AVISO_FIN)
        assert (por_n[6].resultado, por_n[6].aviso_tramos) == (ResultadoChecklist.CUMPLE, None)
        assert por_n[15].citas[0].pagina == 3 and por_n[15].citas[0].verificacion is not None
        assert por_n[15].citas[0].verificacion.respaldada
        assert cl.bajados_sin_cita_verificada == 1
        resumen = resumen_citas_checklist(cl)
        assert resumen is not None and resumen["puntos_bajados"] == [8]
        assert len(resumen["puntos_bajados"]) == cl.bajados_sin_cita_verificada
        if modo == "marcar":
            assert [(por_n[n].resultado, por_n[n].causa_ajuste) for n in (7, 15)] == [
                (ResultadoChecklist.CUMPLE, None)] * 2
            assert (cl.cumplen, cl.marcados_un_tramo, cl.bajados_un_tramo) == (3, 2, None)
        else:
            for n, revisar in ((7, "3-4"), (15, "1-2")):
                assert (por_n[n].resultado, por_n[n].resultado_modelo, por_n[n].causa_ajuste,
                        por_n[n].motivo_ajuste) == (
                    ResultadoChecklist.NO_VERIFICABLE, ResultadoChecklist.CUMPLE,
                    CausaAjuste.UN_TRAMO, MOTIVO_PD44 + f"{revisar}."), n
                assert por_n[n].aviso_tramos is not None
            assert (cl.cumplen, cl.marcados_un_tramo, cl.bajados_un_tramo) == (1, 0, 2)

    @pytest.mark.parametrize("modo", ["marcar", "bajar"])
    def test_con_cita_una_parte_sin_cita_que_respalde_no_cuenta_como_respaldo(
        self, modo: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """Revisión del Bloque 24: con la cita, una parte respalda un «cumple» solo si
        su propia cita respalda (cita_respalda). Una parte que dice «cumple» sin cita,
        o con una que no se verifica, no aporta evidencia: no puede quitar el aviso
        ni la bajada que tendría si hubiera dicho «no verificable». Si ninguna parte
        respalda, el punto ya bajó por la cita y conserva las partes del modelo."""
        from comparador.schemas import CausaAjuste

        sin_respaldo = "Texto que el informe no tiene en ninguna página"
        cl, _, _ = _por_partes(monkeypatch, tmp_path, [(1, 2), (3, 4)], {
            30: (("cumple", PORTADA), ("cumple", sin_respaldo)),   # cita no encontrada
            31: (("cumple", PORTADA), ("cumple", None)),           # sin cita
            6: (("cumple", None), ("cumple", TABLA)),              # fuera de la lista
            32: (("cumple", sin_respaldo), ("cumple", None)),      # ninguna respalda
            1: (("cumple", PORTADA), ("cumple", TABLA)),           # las dos respaldan
        }, modo=modo, con_cita=True, textos={1: PORTADA, 3: TABLA})
        por_n = {i.numero: i for i in cl.items}

        for n, partes, revisar in ((30, [1], "3-4"), (31, [1], "3-4"), (6, [2], "1-2")):
            item = por_n[n]
            assert item.tramos_respaldo == partes, n
            assert item.aviso_tramos is not None and f"las págs. {revisar} no" in item.aviso_tramos
        assert (por_n[1].tramos_respaldo, por_n[1].aviso_tramos) == ([1, 2], None)
        ninguna = por_n[32]
        assert (ninguna.resultado, ninguna.causa_ajuste, ninguna.tramos_respaldo,
                ninguna.aviso_tramos) == (ResultadoChecklist.NO_VERIFICABLE,
                                          CausaAjuste.SIN_CITA_VERIFICADA, [1, 2], None)
        assert (por_n[6].resultado, por_n[6].causa_ajuste) == (ResultadoChecklist.CUMPLE, None)
        assert cl.bajados_sin_cita_verificada == 1
        if modo == "marcar":
            for n in (30, 31):
                assert (por_n[n].resultado, por_n[n].causa_ajuste) == (
                    ResultadoChecklist.CUMPLE, None), n
            assert (cl.cumplen, cl.marcados_un_tramo, cl.bajados_un_tramo) == (4, 3, None)
        else:
            for n in (30, 31):
                assert (por_n[n].resultado, por_n[n].resultado_modelo, por_n[n].causa_ajuste,
                        por_n[n].motivo_ajuste) == (
                    ResultadoChecklist.NO_VERIFICABLE, ResultadoChecklist.CUMPLE,
                    CausaAjuste.UN_TRAMO, MOTIVO_PD44 + "3-4."), n
            assert (cl.cumplen, cl.marcados_un_tramo, cl.bajados_un_tramo) == (2, 1, 2)

    @pytest.mark.parametrize("con_cita", [False, True])
    def test_el_prompt_no_cambia_con_el_modo(
        self, con_cita: bool, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        enviados: dict[str, list[tuple[Any, ...]]] = {}
        for modo in ("marcar", "bajar"):
            _, _, llamadas = _por_partes(monkeypatch, tmp_path, [(1, 2), (3, 4)], {
                7: ("cumple", "no_verificable")}, modo=modo, con_cita=con_cita)
            enviados[modo] = [(k["system_prompt"], k["user_prompt"], k["response_model"])
                              for k in llamadas]
        assert len(enviados["marcar"]) == 2
        assert enviados["marcar"] == enviados["bajar"]

    @pytest.mark.parametrize("modo", ["marcar", "bajar"])
    @pytest.mark.parametrize("rangos", [DOS_PARTES, TRES_PARTES, [(1, 10)]],
                             ids=["dos", "tres", "uno"])
    def test_desde_raw_reproduce_la_etapa(
        self, modo: str, rangos: list[tuple[int, int]], tmp_path: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        import json

        from comparador.etapa0_checklist import checklist_desde_raw

        valores = ("cumple", "no_verificable", "cumple", "no_cumple", "cumple")
        resultados: dict[int, tuple[_Valor, ...]] = {
            n: tuple(valores[(n * (k + 2)) % len(valores)] for k in range(len(rangos)))
            for n in range(1, TOTAL_PUNTOS + 1)
        }
        cl, meta, _ = _por_partes(monkeypatch, tmp_path, rangos, resultados, modo=modo)
        crudo = json.loads(json.dumps(meta))   # como queda en checklist_raw/
        assert checklist_desde_raw(crudo, cargar_puntos(), modo=modo) == cl  # type: ignore[arg-type]

    def test_desde_raw_con_cita_falla(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from comparador.etapa0_checklist import checklist_desde_raw

        _, meta, _ = _por_partes(monkeypatch, tmp_path, [(1, 2), (3, 4)], {
            7: (("cumple", PORTADA), "no_verificable")}, con_cita=True, textos={1: PORTADA})
        with pytest.raises(ValueError, match="checklist.json"):
            checklist_desde_raw(meta, cargar_puntos(), modo="marcar")

    def test_desde_raw_incompleto_o_modo_invalido_falla(self) -> None:
        from comparador.etapa0_checklist import checklist_desde_raw

        raw = {"response_raw": {"items": [{"numero": 1, "resultado": "cumple"}]}}
        with pytest.raises(ValueError, match=r"\[2, 3, "):
            checklist_desde_raw(raw, cargar_puntos(), modo="marcar")
        completo = {"response_raw": _respuesta_completa().model_dump(mode="json")}
        with pytest.raises(ValueError, match="borrar"):
            checklist_desde_raw(completo, cargar_puntos(), modo="borrar")  # type: ignore[arg-type]


class TestConfigUnTramo:
    def test_por_defecto_marcar(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from comparador.config import Settings

        monkeypatch.delenv("CHECKLIST_UN_TRAMO", raising=False)
        assert Settings(_env_file=None).checklist_un_tramo == "marcar"  # type: ignore[call-arg]

    @pytest.mark.parametrize("valor", ["marcar", "bajar"])
    def test_lee_la_variable(self, valor: str, monkeypatch: pytest.MonkeyPatch) -> None:
        from comparador.config import Settings

        monkeypatch.setenv("CHECKLIST_UN_TRAMO", valor)
        assert Settings(_env_file=None).checklist_un_tramo == valor  # type: ignore[call-arg]

    @pytest.mark.parametrize("valor", ["borrar", "Bajar", "", "true"])
    def test_valor_invalido_falla_al_arrancar(
        self, valor: str, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from pydantic import ValidationError

        from comparador.config import Settings

        monkeypatch.setenv("CHECKLIST_UN_TRAMO", valor)
        with pytest.raises(ValidationError, match="checklist_un_tramo"):
            Settings(_env_file=None)  # type: ignore[call-arg]
