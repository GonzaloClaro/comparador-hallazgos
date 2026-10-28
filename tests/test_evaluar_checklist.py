"""Tests del evaluador de checklist contra ground truth humano."""

from pathlib import Path

import openpyxl
import pytest

from comparador.evaluar_checklist import cargar_gt, cargar_mapeo, clasificar
from comparador.schemas import ResultadoChecklist


class TestClasificar:
    """La función pura que define las cuatro categorías."""

    def test_cumple_con_gt_si_es_acierto(self) -> None:
        assert clasificar(ResultadoChecklist.CUMPLE, "SI") == "acierto"

    def test_no_cumple_con_gt_no_es_acierto(self) -> None:
        assert clasificar(ResultadoChecklist.NO_CUMPLE, "NO") == "acierto"

    def test_no_verificable_con_gt_si(self) -> None:
        assert clasificar(ResultadoChecklist.NO_VERIFICABLE, "SI") == "no_verificable"

    def test_no_verificable_con_gt_no(self) -> None:
        assert clasificar(ResultadoChecklist.NO_VERIFICABLE, "NO") == "no_verificable"

    def test_sistema_no_cumple_humano_si_es_falso_no_cumple(self) -> None:
        assert clasificar(ResultadoChecklist.NO_CUMPLE, "SI") == "falso_no_cumple"

    def test_sistema_cumple_humano_no_es_falso_cumple(self) -> None:
        assert clasificar(ResultadoChecklist.CUMPLE, "NO") == "falso_cumple"


class TestMapeo:
    def test_cubre_exactamente_los_38_puntos(self) -> None:
        mapeo = cargar_mapeo()
        assert sorted(mapeo) == list(range(1, 39))

    def test_codigos_tienen_formato_seccion_item(self) -> None:
        import re
        for codigo in cargar_mapeo().values():
            assert re.fullmatch(r"\d+\.\d+(#\d+)?", codigo), codigo

    def test_desgloses_de_vineta_comparten_item_gt(self) -> None:
        mapeo = cargar_mapeo()
        assert mapeo[3] == mapeo[4] == mapeo[5] == "1.4"

    def test_punto_20_es_el_item_de_explicitacion(self) -> None:
        # 5.8 es "determinar el despeje por escenario"; el 20 (indicar en
        # forma explícita los casos < 400 ms) es el 5.9 (bug corregido).
        assert cargar_mapeo()[20] == "5.9"


def _gt_sintetico(tmp_path: Path) -> Path:
    """Arma un xlsx con el formato del CheckList de Revisión del CEN."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Check List"
    # columnas 6..9 de openpyxl = índices 5..8 de pandas (código, enunciado,
    # veredicto, observación); incluye un código duplicado como el 7.7 real.
    filas = [
        ("1.2", "Identifica al Coordinado", "SI", None),
        ("4.4.", "Ajustes levantados en terreno", "SI", None),   # punto final espurio
        ("7.7", "Prueba 87T", "NO", "no se hace en el paño CT"),
        ("7.7", "Estabilidad 87T", "NO", None),
        ("9.9", "Ítem sin veredicto", None, None),
        ("3.7", "Esquemas adaptativos", "N/A", None),   # "no aplica" explícito
    ]
    for r, (cod, enun, ver, obs) in enumerate(filas, start=3):
        ws.cell(row=r, column=6, value=cod)
        ws.cell(row=r, column=7, value=enun)
        ws.cell(row=r, column=8, value=ver)
        ws.cell(row=r, column=9, value=obs)
    ruta = tmp_path / "gt.xlsx"
    wb.save(ruta)
    return ruta


class TestCargarGt:
    def test_lee_veredictos_y_observaciones(self, tmp_path: Path) -> None:
        gt = cargar_gt(_gt_sintetico(tmp_path))
        assert gt["1.2"] == ("SI", "")
        assert gt["7.7"] == ("NO", "no se hace en el paño CT")

    def test_normaliza_punto_final_espurio(self, tmp_path: Path) -> None:
        gt = cargar_gt(_gt_sintetico(tmp_path))
        assert "4.4" in gt  # el GT real trae "4.4." con punto final

    def test_codigo_duplicado_se_desambigua_con_sufijo(self, tmp_path: Path) -> None:
        gt = cargar_gt(_gt_sintetico(tmp_path))
        assert gt["7.7#2"] == ("NO", "")

    def test_item_sin_veredicto_queda_vacio(self, tmp_path: Path) -> None:
        gt = cargar_gt(_gt_sintetico(tmp_path))
        assert gt["9.9"][0] == ""

    def test_na_explicito_se_conserva_y_no_es_veredicto(self, tmp_path: Path) -> None:
        # pandas convertía "N/A" en NaN: se perdía la diferencia con una celda vacía.
        gt = cargar_gt(_gt_sintetico(tmp_path))
        assert gt["3.7"] == ("N/A", "")
        assert gt["3.7"][0] not in ("SI", "NO")

    def test_hoja_vacia_falla_ruidosamente(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Check List"
        ruta = tmp_path / "vacio.xlsx"
        wb.save(ruta)
        with pytest.raises(ValueError, match="No se encontraron"):
            cargar_gt(ruta)


class TestChecklistDelRun:
    """`cargar_checklist_run`: el del reporte final o, en un run que corrió solo el
    checklist (ronda B3), su checklist.json."""

    def _checklist(self, cumplen: int) -> dict[str, object]:
        from comparador.schemas import ChecklistInforme, ItemChecklist, ResultadoChecklist

        return ChecklistInforme(items=[
            ItemChecklist(numero=1, enunciado="e", resultado=ResultadoChecklist.CUMPLE),
        ], cumplen=cumplen, total=38).model_dump(mode="json")

    def test_sin_reporte_usa_checklist_json(self, tmp_path: Path,
                                            monkeypatch: pytest.MonkeyPatch):
        import json

        from comparador.config import settings
        from comparador.evaluar_checklist import cargar_checklist_run

        monkeypatch.setattr(settings, "runs_dir", tmp_path)
        (tmp_path / "run").mkdir()
        (tmp_path / "run" / "checklist.json").write_text(json.dumps(self._checklist(7)))
        assert cargar_checklist_run("run").cumplen == 7

    def test_sin_nada_falla_claro(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        from comparador.config import settings
        from comparador.evaluar_checklist import cargar_checklist_run

        monkeypatch.setattr(settings, "runs_dir", tmp_path)
        (tmp_path / "run").mkdir()
        with pytest.raises(ValueError, match="ni checklist.json"):
            cargar_checklist_run("run")
